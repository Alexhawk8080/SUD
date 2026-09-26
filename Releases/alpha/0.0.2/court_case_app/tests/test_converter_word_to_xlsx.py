# -*- coding: utf-8 -*-
"""
Тесты парсера Word-акта (stage13a, под-этап C1).

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tests\\test_converter_word_to_xlsx.py
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document

from core.converter.exceptions import NoCasesFound
from core.converter.word_to_xlsx import (
    extract_refs_from_doc,
    extract_table_records,
    find_cases_table,
    read_act_docx,
)
from core.converter.markers import mark_word_document

FAILED = 0
PASSED = 0


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def expect_error(name, fn, exc_type):
    global FAILED, PASSED
    try:
        fn()
    except exc_type:
        PASSED += 1
        print(f"  OK   {name}")
        return
    except Exception as exc:  # noqa: BLE001
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось {exc_type.__name__}, "
              f"получено {type(exc).__name__}: {exc}")
        return
    FAILED += 1
    print(f"  FAIL {name}\n       исключение не поднято")


HEADERS = ["№ п/п", "Заголовок дела", "Даты дела", "№ описи",
           "№ ед.хр.", "Кол-во ед.хр.", "Срок хранения", "Примечание"]


def make_act_with_cases(cases, save_path, marker=True):
    """
    cases — список кортежей:
      (seq, title, dates, opis, unit, count, retention, note).
    """
    doc = Document()
    doc.add_paragraph("Акт о выделении к уничтожению документов")
    doc.add_paragraph("Судебный участок № 9")
    doc.add_paragraph("Мировой судья: О. А. Левошина")
    doc.add_paragraph("Секретарь: А.М. Намаюшка")

    table = doc.add_table(rows=1, cols=len(HEADERS))
    for i, h in enumerate(HEADERS):
        table.cell(0, i).text = h
    for c in cases:
        row = table.add_row()
        for i, val in enumerate(c):
            row.cells[i].text = str(val)

    if marker:
        mark_word_document(doc)
    doc.save(save_path)
    return save_path


def main():
    tmp = tempfile.mkdtemp(prefix="word_parser_")

    print("[1] find_cases_table: единственная таблица")
    p1 = os.path.join(tmp, "act1.docx")
    make_act_with_cases([
        (1, "Гражданское дело №2-1/2020 по заявлению Тестова Т. Т. "
            "к Тестову Т. Т. о расторжении брака",
         "01.01.2020\n01.02.2020", 1, 1, 1, "3 года ЭК Ст. 137", ""),
    ], p1)
    doc = Document(p1)
    t, idx = find_cases_table(doc)
    check("Индекс таблицы 0", idx, 0)
    check("Таблица найдена", t is not None, True)

    print("\n[2] find_cases_table: две таблицы — берётся с большим числом дел")
    p2 = os.path.join(tmp, "two_tables.docx")
    doc2 = Document()
    doc2.add_paragraph("АКТ")
    # Маленькая таблица в шапке (не о делах)
    t_small = doc2.add_table(rows=1, cols=2)
    t_small.cell(0, 0).text = "Поле"
    t_small.cell(0, 1).text = "Значение"
    # Большая таблица — дела
    t_big = doc2.add_table(rows=1, cols=len(HEADERS))
    for i, h in enumerate(HEADERS):
        t_big.cell(0, i).text = h
    for i in range(1, 4):
        row = t_big.add_row()
        row.cells[0].text = str(i)
        row.cells[1].text = (f"Гражданское дело №2-{i}/2020 "
                             f"о защите прав")
    mark_word_document(doc2)
    doc2.save(p2)
    t3, idx3 = find_cases_table(Document(p2))
    check("Выбрана таблица дел (индекс 1)", idx3, 1)

    print("\n[3] Нет таблиц с делами -> NoCasesFound")
    p3 = os.path.join(tmp, "empty.docx")
    doc3 = Document()
    doc3.add_paragraph("Только текст, без таблиц")
    doc3.save(p3)
    expect_error("Нет таблицы дел",
                 lambda: find_cases_table(Document(p3)), NoCasesFound)

    print("\n[4] extract_table_records: 3 дела, все валидные")
    p4 = os.path.join(tmp, "act4.docx")
    make_act_with_cases([
        (1, "Гражданское дело №2-1/2020 о расторжении брака",
         "01.01.2020\n01.02.2020", 1, 1, 1, "3 года ЭК Ст. 137", ""),
        (2, "Гражданское дело №2-2/2020 о защите прав",
         "02.01.2020\n02.02.2020", 1, 2, 1, "3 года ЭПК Ст. 208", ""),
        (3, "Гражданское дело №2-3/2020 о взыскании",
         "03.01.2020\n03.02.2020", 1, 3, 1, "3 года Ст. 227", "архив"),
    ], p4)
    result = extract_table_records(find_cases_table(Document(p4))[0])
    check("3 записи", len(result["records"]), 3)
    check("Attention пуст", result["attention"], [])
    check("sequential[0]='1'", result["records"][0]["sequential"], "1")
    check("note[2]='архив'", result["records"][2]["note"], "архив")
    check("dates[0] с \\n", result["records"][0]["dates"], "01.01.2020\n01.02.2020")

    print("\n[5] Умная нумерация: последовательность сломана (1,3,4)")
    p5 = os.path.join(tmp, "seq_bad.docx")
    make_act_with_cases([
        (1, "Гражданское дело №2-1/2020 о расторжении брака",
         "01.01.2020", 1, 1, 1, "3 года ЭК Ст. 137", ""),
        (3, "Гражданское дело №2-3/2020 о защите прав",
         "03.01.2020", 1, 3, 1, "3 года ЭПК Ст. 208", ""),
        (4, "Гражданское дело №2-4/2020 о защите прав",
         "04.01.2020", 1, 4, 1, "3 года ЭПК Ст. 208", ""),
    ], p5)
    result = extract_table_records(find_cases_table(Document(p5))[0])
    seqs = [r["sequential"] for r in result["records"]]
    check("Переиндексация: 1,2,3", seqs, ["1", "2", "3"])
    check("Attention содержит замечание про нумерацию",
          any("нумерация" in a["reason"] for a in result["attention"]), True)

    print("\n[6] Плохие даты -> attention")
    p6 = os.path.join(tmp, "bad_dates.docx")
    make_act_with_cases([
        (1, "Гражданское дело №2-1/2020 о расторжении брака",
         "не дата", 1, 1, 1, "3 года ЭК Ст. 137", ""),
    ], p6)
    result = extract_table_records(find_cases_table(Document(p6))[0])
    check("Дата не распознана",
          any("даты" in a["reason"] for a in result["attention"]), True)
    check("Даты остались как есть",
          result["records"][0]["dates"], "не дата")

    print("\n[7] Нет номера в заголовке -> attention")
    p7 = os.path.join(tmp, "no_num.docx")
    make_act_with_cases([
        (1, "Гражданское дело №б/н о защите прав потребителей", "01.01.2020", 1, 1, 1, "", ""),
    ], p7)
    result = extract_table_records(find_cases_table(Document(p7))[0])
    check("Attention про номер",
          any("номер" in a["reason"].lower() for a in result["attention"]), True)

    print("\n[8] extract_refs_from_doc: реквизиты из шапки")
    p8 = os.path.join(tmp, "refs.docx")
    make_act_with_cases([
        (1, "Гражданское дело №2-1/2020 о защите прав",
         "01.01.2020", 1, 1, 1, "", ""),
    ], p8)
    doc8 = Document(p8)
    refs = extract_refs_from_doc(doc8)
    check("номер = '9'", refs.get("номер"), "9")
    check("судья содержит 'Левошина'", "Левошина" in refs.get("судья", ""), True)
    check("секретарь содержит 'Намаюшка'", "Намаюшка" in refs.get("секретарь", ""), True)

    print("\n[9] read_act_docx: полный API")
    result = read_act_docx(p4)
    check("Есть records", len(result["records"]), 3)
    check("Есть attention", isinstance(result["attention"], list), True)
    check("Есть refs", "номер" in result["refs"], True)
    check("marker.found = True", result["marker"]["found"], True)
    check("marker.version = 'V1'", result["marker"]["version"], "V1")
    check("table_index = 0", result["table_index"], 0)

    print("\n[10] read_act_docx: без метки (чужой акт)")
    p10 = os.path.join(tmp, "act_no_marker.docx")
    make_act_with_cases([
        (1, "Гражданское дело №2-1/2020 о защите прав",
         "01.01.2020", 1, 1, 1, "", ""),
    ], p10, marker=False)
    result10 = read_act_docx(p10)
    check("marker.found = False", result10["marker"]["found"], False)
    check("Records всё равно извлечены", len(result10["records"]), 1)

    print("\n[11] read_act_docx: не-docx файл -> BadFileError")
    from core.converter.exceptions import BadFileError
    fake = os.path.join(tmp, "fake.docx")
    with open(fake, "w", encoding="utf-8") as f:
        f.write("это не docx")
    expect_error("Не-docx -> BadFileError",
                 lambda: read_act_docx(fake), BadFileError)

    print("\n[12] Одна дата -> нормализуется в одну строку")
    p12 = os.path.join(tmp, "one_date.docx")
    make_act_with_cases([
        (1, "Гражданское дело №2-1/2020 о защите прав",
         "15.03.2020", 1, 1, 1, "", ""),
    ], p12)
    r12 = extract_table_records(find_cases_table(Document(p12))[0])
    check("Одна дата сохранена", r12["records"][0]["dates"], "15.03.2020")

    print("\n[13] Дубликат даты не дублируется")
    p13 = os.path.join(tmp, "dup_date.docx")
    make_act_with_cases([
        (1, "Гражданское дело №2-1/2020 о защите прав",
         "01.01.2020 01.01.2020", 1, 1, 1, "", ""),
    ], p13)
    r13 = extract_table_records(find_cases_table(Document(p13))[0])
    check("Одна дата", r13["records"][0]["dates"], "01.01.2020")

    print("\n[14] Регрессия fix_10: _x000D_ в датах нормализуется")
    from core.converter.word_to_xlsx import normalize_cell_text
    check("_x000D_ -> \n",
          normalize_cell_text("08.12.2017_x000D_08.01.2018"),
          "08.12.2017\n08.01.2018")
    check("\r -> \n", normalize_cell_text("a\rb"), "a\nb")
    check("\v -> \n", normalize_cell_text("a\vb"), "a\nb")
    check("Обычная строка без изменений",
          normalize_cell_text("обычный текст"), "обычный текст")
    check("None -> None", normalize_cell_text(None), None)

    # Сквозная проверка через extract_table_records на акте с _x000D_
    p14 = os.path.join(tmp, "with_x000d.docx")
    doc14 = Document()
    doc14.add_paragraph("Акт")
    t14 = doc14.add_table(rows=1, cols=len(HEADERS))
    for i, h in enumerate(HEADERS):
        t14.cell(0, i).text = h
    row14 = t14.add_row()
    row14.cells[0].text = "1"
    row14.cells[1].text = "Гражданское дело №2-1/2020 о защите прав"
    # Симулируем _x000D_ как литерал в ячейке (что и происходит в Word)
    row14.cells[2].text = "08.12.2017_x000D_08.01.2018"
    row14.cells[3].text = "1"
    row14.cells[4].text = "1"
    row14.cells[5].text = "1"
    row14.cells[6].text = "3 года Ст. 208"
    row14.cells[7].text = ""
    mark_word_document(doc14)
    doc14.save(p14)

    r14 = extract_table_records(find_cases_table(Document(p14))[0])
    check("_x000D_ в дате нормализован в \n",
          r14["records"][0]["dates"], "08.12.2017\n08.01.2018")
    check("Attention пуст (даты распознаны)",
          r14["attention"], [])

    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()
