# -*- coding: utf-8 -*-
"""
Веб-приложение «Акт уничтожения гражданских дел».

Flask-сервер: главная страница (обработка файлов), страница базы данных,
API для управления БД, скачивание результата, предпросмотр исходного файла.
Автозапуск браузера; остановка сервера при закрытии вкладки интерфейса
(heartbeat + сигнал при закрытии).
"""

import json
import os
import threading
import time
import webbrowser

from flask import Flask, abort, jsonify, render_template, request, send_file

from core.pipeline import process_rows, process_upload
from core.excel_reader import (
    ROLES,
    ROLE_TITLES,
    auto_detect_roles,
    detect_header_row,
    read_source_table,
)
from core.excel_writer import DEFAULT_COLUMNS
from core.category_check import analyze
from core.legacy_reader import parse_legacy_file
from core.case_processor import contains_alimony, is_valid_case_number
from database import db

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "database", "app.db")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
TEMPLATE_DOCX = os.path.join(
    BASE_DIR, "docs", "АКТ уничтожения гражданских дел.docx")

# Heartbeat: если вкладка интерфейса закрыта и heartbeat пропал — остановить сервер.
# Web Worker шлёт heartbeat каждые 4 с; таймаут 20 с = запас на 4 пропуска
# (переходы между страницами интерфейса, кратковременные паузы браузера).
HEARTBEAT_TIMEOUT = 20.0  # секунд без heartbeat до остановки
_last_beat = time.time()


def _beat_watchdog():
    """Фоновый поток: останавливает сервер, если heartbeat пропал."""
    while True:
        time.sleep(2.0)
        if time.time() - _last_beat > HEARTBEAT_TIMEOUT:
            os._exit(0)


def start_beat_watchdog():
    threading.Thread(target=_beat_watchdog, daemon=True).start()


app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024  # 25 МБ


def get_db():
    """Соединение с БД (создание/наполнение при первом обращении)."""
    return db.init_db(DB_PATH)


def _json_error(message, code=400):
    return jsonify({"error": message}), code


# ---------------------------------------------------------------------------
# Страницы
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    return render_template(
        "index.html",
        roles=ROLES,
        role_titles=ROLE_TITLES,
        columns=DEFAULT_COLUMNS,
    )


@app.route("/db_page")
def db_page():
    return render_template("db.html")


@app.route("/check_page")
def check_page():
    """Страница «Проверка категорий» (Доработка 6)."""
    return render_template("check.html")


@app.route("/heartbeat", methods=["GET"])
def heartbeat():
    """Сигнал «вкладка жива» от клиента."""
    global _last_beat
    _last_beat = time.time()
    return jsonify({"ok": True})


@app.route("/shutdown", methods=["POST"])
def shutdown():
    """Остановка сервера по сигналу клиента (закрытие вкладки)."""
    os._exit(0)


@app.route("/preview", methods=["POST"])
def preview():
    """
    Предпросмотр загруженного исходного файла.

    Возвращает первые строки листа, автоопределённые номер строки
    заголовков и маппинг ролей.
    """
    if "file" not in request.files:
        return _json_error("Файл не загружен")
    uploaded = request.files["file"]
    if not uploaded.filename:
        return _json_error("Файл не выбран")

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    path = os.path.join(UPLOAD_DIR, os.path.basename(uploaded.filename))
    uploaded.save(path)

    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        rows = []
        for values in ws.iter_rows(values_only=True):
            cells = []
            for v in values[:14]:
                if v is None:
                    cells.append("")
                else:
                    s = str(v)
                    cells.append(s[:80] + ("…" if len(s) > 80 else ""))
            rows.append(cells)
            if len(rows) >= 12:
                break
        wb.close()
    except Exception as exc:  # noqa: BLE001
        return _json_error(f"Не удалось прочитать файл: {exc}", 400)

    header_row = detect_header_row(rows)
    mapping = {}
    if header_row > 0 and header_row <= len(rows):
        mapping = auto_detect_roles(rows[header_row - 1])
    return jsonify({
        "rows": rows,
        "header_row": header_row,
        "mapping": mapping,
    })


# ---------------------------------------------------------------------------
# Проверка категорий дел (Доработка 6)
# ---------------------------------------------------------------------------

def _filter_rows_for_year(rows, year, alimony: bool) -> list:
    """
    Оставляет дела, попадающие в результат обработки:
    валидный номер за выбранный год + фильтр алиментных дел.
    """
    year = str(year)
    result = []
    for row in rows:
        if not is_valid_case_number(row.get("case_number"), year):
            continue
        if alimony or not contains_alimony(row.get("category")):
            result.append(row)
    return result


@app.route("/check_categories", methods=["POST"])
def check_categories():
    """
    Анализ загруженного файла: поиск дел с отсутствующей категорией,
    аналогов по истцу, групп «Без аналогов».

    Параметры формы:
        file       — файл (исходная таблица .xlsx/.xlsm или готовый результат);
        file_type  — 'source' (исходная таблица) или 'legacy' (готовый результат);
        year       — год дел;
        alimony    — '1' — включать алиментные дела;
        header_row — номер строки заголовков (для исходной таблицы).
    """
    if "file" not in request.files:
        return _json_error("Файл не загружен")
    uploaded = request.files["file"]
    if not uploaded.filename:
        return _json_error("Файл не выбран")

    year = request.form.get("year", "").strip()
    if not year.isdigit():
        return _json_error("Укажите год (например, 2023)")
    alimony = request.form.get("alimony") == "1"
    file_type = request.form.get("file_type", "source")

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    path = os.path.join(UPLOAD_DIR, os.path.basename(uploaded.filename))
    uploaded.save(path)

    try:
        if file_type == "legacy":
            # Готовый результат: записи извлекаются из заголовков
            raw_rows = parse_legacy_file(path)
        else:
            header_raw = request.form.get("header_row", "").strip()
            if header_raw == "":
                header_row = None
            else:
                try:
                    header_row = int(header_raw)
                except ValueError:
                    header_row = None
            _headers, raw_rows, _mapping = read_source_table(
                path, mapping=None, header_row=header_row)
    except Exception as exc:  # noqa: BLE001
        return _json_error(f"Не удалось прочитать файл: {exc}", 400)

    rows = _filter_rows_for_year(raw_rows, year, alimony)

    conn = get_db()
    try:
        org_names = db.organization_names(conn)
        categories_map = db.case_categories_map(conn)
        auto_fix_enabled = db.get_setting(conn, "auto_fix_categories", "1") == "1"
    finally:
        conn.close()

    result = analyze(rows, org_names, categories_map)
    result["cases_total"] = len(rows)
    result["auto_fix_enabled"] = auto_fix_enabled
    return jsonify(result)


@app.route("/apply_corrections", methods=["POST"])
def apply_corrections():
    """
    Сохранение исправлений в словарь «истец -> категория» (case_categories).

    Тело JSON: {"saves": [{"key": ..., "label": ..., "category": ...,
                           "is_organization": bool}, ...]}
    """
    data = request.get_json(force=True)
    saves = data.get("saves", []) or []
    conn = get_db()
    saved = 0
    try:
        for item in saves:
            key = str(item.get("key", "")).strip()
            label = str(item.get("label", "")).strip()
            category = str(item.get("category", "")).strip()
            if not key or not category:
                continue
            db.add_case_category(
                conn, key, label or key, category,
                bool(item.get("is_organization", False)))
            saved += 1
        return jsonify({"ok": True, "saved": saved})
    finally:
        conn.close()


@app.route("/download_no_analog_table", methods=["POST"])
def download_no_analog_table():
    """
    Выгрузка таблицы «Без аналогов» в Excel (.xlsx) для ручной работы.

    Тело JSON: {"groups": [...], "year": "2023"}
    """
    data = request.get_json(force=True)
    groups = data.get("groups", []) or []
    year = data.get("year", "")

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Без аналогов"
    ws.append(["Истец", "№ дела", "Категория (заполнить)"])
    for group in groups:
        label = str(group.get("label", ""))
        for case in group.get("cases", []):
            ws.append([label, str(case.get("case_number", "")), ""])
        # Подсказка из словаря БД
        opts = group.get("dictionary_options") or []
        if opts:
            ws.append(["", "", f"Варианты из словаря: {'; '.join(opts)}"])

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"без_аналогов_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = os.path.join(OUTPUT_DIR, filename)
    wb.save(path)
    return send_file(path, as_attachment=True, download_name=filename)


# ---------------------------------------------------------------------------
# Обработка и скачивание
# ---------------------------------------------------------------------------

@app.route("/process", methods=["POST"])
def process():
    if "file" not in request.files:
        return _json_error("Файл не загружен")
    uploaded = request.files["file"]
    if not uploaded.filename:
        return _json_error("Файл не выбран")

    year = request.form.get("year", "").strip()
    if not year.isdigit():
        return _json_error("Укажите год (например, 2023)")

    alimony = request.form.get("alimony") == "1"
    output_format = request.form.get("format", "excel")
    if output_format not in ("excel", "word"):
        output_format = "excel"

    # Ручной маппинг ролей (если передан)
    mapping = {}
    for role in ROLES:
        value = request.form.get(f"map_{role}", "").strip()
        if value.isdigit() and int(value) > 0:
            mapping[role] = int(value)
    header_raw = request.form.get("header_row", "").strip()
    if header_raw == "":
        header_row = None  # автоопределение
    else:
        try:
            header_row = int(header_raw)
        except ValueError:
            header_row = None
        if header_row is not None and header_row < 0:
            header_row = 0

    # Выбранные столбцы для Excel (пусто = все)
    columns = request.form.getlist("columns")

    # Доработка 6: исправления категорий со страницы проверки + автоисправление
    fixes_raw = request.form.get("fixes", "").strip()
    category_fixes = {}
    if fixes_raw:
        try:
            raw = json.loads(fixes_raw) or {}
            # fix_02: ключи номеров дел — без хвостовых/неразрывных пробелов
            category_fixes = {
                str(k).strip(): v for k, v in raw.items() if str(k).strip()
            }
        except ValueError:
            category_fixes = {}
    auto_fix_enabled = request.form.get("auto_fix", "1") != "0"
    file_type = request.form.get("file_type", "source")

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    upload_path = os.path.join(UPLOAD_DIR, os.path.basename(uploaded.filename))
    uploaded.save(upload_path)

    try:
        if file_type == "legacy":
            # Готовый результат: записи собраны разбором заголовков,
            # даты/описи недоступны — предупреждаем о неполноте данных
            raw_rows = parse_legacy_file(upload_path)
            if not raw_rows:
                raise ValueError(
                    "В файле не найдено дел (проверьте, что это результат "
                    "обработки или акт уничтожения)")
            result = process_rows(
                raw_rows, int(year), alimony, output_format,
                DB_PATH, TEMPLATE_DOCX, OUTPUT_DIR,
                columns=columns or None,
                category_fixes=category_fixes,
                auto_fix_enabled=auto_fix_enabled)
            result["legacy_note"] = (
                "ВНИМАНИЕ: обработан готовый результат — даты дела, "
                "№ описи и № ед.хр. не восстановлены и останутся пустыми.")
        else:
            result = process_upload(
                upload_path, int(year), alimony, output_format,
                DB_PATH, TEMPLATE_DOCX, OUTPUT_DIR,
                columns=columns or None,
                mapping=mapping or None,
                header_row=header_row,
                category_fixes=category_fixes,
                auto_fix_enabled=auto_fix_enabled,
            )
    except ValueError as exc:
        return _json_error(str(exc))
    except Exception as exc:  # noqa: BLE001
        return _json_error(f"Ошибка обработки: {exc}", 500)

    # Доработка 7: прямое скачивание результата (кнопка «Скачать результат»)
    if request.form.get("download") == "1":
        return send_file(result["path"], as_attachment=True,
                         download_name=result["filename"])

    # Подготовка данных для предпросмотра
    preview = []
    for record in result["records"]:
        preview.append({
            "sequential": record.get("sequential"),
            "title": record.get("title") or "",
            "dates": (record.get("dates") or "").replace("\n", "<br>"),
            "opis": record.get("opis") or "",
            "unit": record.get("unit") or "",
            "count": record.get("count") or "",
            "retention": record.get("retention") or "",
            "note": record.get("note") or "",
        })

    stats = result["stats"]
    stats_message = (
        f"Всего номеров в диапазоне: {stats['total_in_range']}\n"
        f"Найдено подходящих дел: {stats['valid_count']}\n"
        f"Пропущенных номеров: {stats['skipped_numbers']}\n"
        f"Алиментных дел исключено: {stats['alimony_excluded']}\n"
        f"Категорий автоподставлено из словаря: {stats.get('auto_fixed', 0)}"
    )
    if result.get("legacy_note"):
        stats_message += "\n" + result["legacy_note"]
    return jsonify({
        "records": preview,
        "stats": stats,
        "stats_message": stats_message,
        "filename": result["filename"],
    })


@app.route("/download/<path:filename>")
def download(filename):
    safe = os.path.basename(filename)
    path = os.path.join(OUTPUT_DIR, safe)
    if not os.path.isfile(path):
        abort(404)
    return send_file(path, as_attachment=True)


# ---------------------------------------------------------------------------
# API: организации
# ---------------------------------------------------------------------------

@app.route("/api/organizations", methods=["GET"])
def api_orgs_list():
    conn = get_db()
    try:
        return jsonify(db.list_organizations(conn))
    finally:
        conn.close()


@app.route("/api/organizations", methods=["POST"])
def api_orgs_add():
    data = request.get_json(force=True)
    name = str(data.get("name", "")).strip()
    if not name:
        return _json_error("Название организации не может быть пустым")
    claim_type = str(data.get("claim_type", "по иску")).strip()
    conn = get_db()
    try:
        db.add_organization(conn, name, claim_type)
        return jsonify({"ok": True})
    finally:
        conn.close()


@app.route("/api/organizations/<int:org_id>", methods=["PUT"])
def api_orgs_update(org_id):
    data = request.get_json(force=True)
    name = str(data.get("name", "")).strip()
    if not name:
        return _json_error("Название организации не может быть пустым")
    claim_type = str(data.get("claim_type", "по иску")).strip()
    conn = get_db()
    try:
        db.update_organization(conn, org_id, name, claim_type)
        return jsonify({"ok": True})
    finally:
        conn.close()


@app.route("/api/organizations/<int:org_id>", methods=["DELETE"])
def api_orgs_delete(org_id):
    conn = get_db()
    try:
        db.delete_organization(conn, org_id)
        return jsonify({"ok": True})
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# API: импорт/экспорт базы данных
# ---------------------------------------------------------------------------

@app.route("/api/db/import", methods=["POST"])
def api_db_import():
    """Импорт организаций из Excel-файла (лист «База данных»)."""
    if "file" not in request.files:
        return _json_error("Файл не загружен")
    uploaded = request.files["file"]
    if not uploaded.filename:
        return _json_error("Файл не выбран")

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    path = os.path.join(UPLOAD_DIR, os.path.basename(uploaded.filename))
    uploaded.save(path)
    conn = get_db()
    try:
        added = db.import_database_from_excel(conn, path)
        return jsonify({"ok": True, "added": added})
    except Exception as exc:  # noqa: BLE001
        return _json_error(f"Ошибка импорта: {exc}", 400)
    finally:
        conn.close()


@app.route("/api/db/export")
def api_db_export():
    """Экспорт базы данных в Excel-файл."""
    conn = get_db()
    try:
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        filename = f"база_данных_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        path = os.path.join(OUTPUT_DIR, filename)
        db.export_database_to_excel(conn, path)
        return send_file(path, as_attachment=True,
                         download_name=filename)
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# API: правила сроков хранения
# ---------------------------------------------------------------------------

@app.route("/api/retention_rules", methods=["GET"])
def api_rules_list():
    conn = get_db()
    try:
        return jsonify(db.list_retention_rules(conn))
    finally:
        conn.close()


@app.route("/api/retention_rules", methods=["POST"])
def api_rules_add():
    data = request.get_json(force=True)
    keyword = str(data.get("keyword", "")).strip()
    code = str(data.get("code", "")).strip()
    if not keyword or not code:
        return _json_error("Укажите корень и код типа")
    conn = get_db()
    try:
        db.add_retention_rule(conn, keyword, code,
                              data.get("result_text"))
        return jsonify({"ok": True})
    finally:
        conn.close()


@app.route("/api/retention_rules/<int:rule_id>", methods=["PUT"])
def api_rules_update(rule_id):
    data = request.get_json(force=True)
    keyword = str(data.get("keyword", "")).strip()
    code = str(data.get("code", "")).strip()
    if not keyword or not code:
        return _json_error("Укажите корень и код типа")
    conn = get_db()
    try:
        db.update_retention_rule(conn, rule_id, keyword, code,
                                 data.get("result_text"))
        return jsonify({"ok": True})
    finally:
        conn.close()


@app.route("/api/retention_rules/<int:rule_id>", methods=["DELETE"])
def api_rules_delete(rule_id):
    conn = get_db()
    try:
        db.delete_retention_rule(conn, rule_id)
        return jsonify({"ok": True})
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# API: настройки
# ---------------------------------------------------------------------------

@app.route("/api/settings", methods=["GET"])
def api_settings_get():
    conn = get_db()
    try:
        return jsonify(db.get_settings(conn))
    finally:
        conn.close()


@app.route("/api/settings", methods=["POST"])
def api_settings_save():
    data = request.get_json(force=True)
    conn = get_db()
    try:
        for key, value in data.items():
            db.set_setting(conn, str(key), str(value))
        return jsonify({"ok": True})
    finally:
        conn.close()


if __name__ == "__main__":
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    # Инициализация БД при старте
    get_db().close()
    # Автозапуск браузера и контроль закрытия вкладки (остановка сервера)
    start_beat_watchdog()
    threading.Timer(1.0, lambda: webbrowser.open("http://127.0.0.1:5000")).start()
    app.run(host="127.0.0.1", port=5000, debug=False)