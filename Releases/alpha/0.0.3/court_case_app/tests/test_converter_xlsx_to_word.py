# -*- coding: utf-8 -*-
"""
Тесты конвертера Excel -> Word (stage12b, под-этап B2).

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tests\\test_converter_xlsx_to_word.py
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import Workbook
from docx import Document

from core.converter.exceptions import ColumnsNotFound, SheetNotFound
from core.converter.xlsx_to_word import (
    COLUMNS,
    convert_xlsx_to_word,
    read_result_xlsx,
)
from core.converter.markers import read_word_marker  # stage12b_v2

FAILED = 0
PASSED = 0

TEMPLATE = os.path.join(os.path.dirname(__file__), "..", "docs",
                        "АКТ уничтожения гражданских дел.docx")


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def expect_error(name, fn, exc_type):
    global FAILED, PASSED
    try:
        fn()
    except exc_type:
        PASSED += 1
        print(f"  OK   {name}")
        return
    except Exception as exc:  # noqa: BLE001
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось {exc_type.__name__}, "
              f"получено {type(exc).__name__}: {exc}")
        return
    FAILED += 1
    print(f"  FAIL {name}\n       исключение не поднято")


def make_result_xlsx(path, sheet_name="Результат обработки"):
    """Минимальный Excel-результат: шапка + 3 строки."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append([t for _, t in COLUMNS])
    ws.append([1,
               "Гражданское дело №2-1/2020 по заявлению Тестова Т. Т. "
               "к Тестову Т. Т. о расторжении брака",
               "01.01.2020\n01.02.2020", 1, 1, 1,
               "3 года ЭК Ст. 137", ""])
    # Пустая строка (пропуск)
    ws.append([2, "", "", "", "", "", "", ""])
    ws.append([3,
               "Гражданское дело №2-3/2020 по иску ООО \"Сеть\" "
               "к Тестову Т. Т. о защите прав потребителей",
               "15.02.2020\n20.03.2020", 1, 3, 1,
               "3 года ЭПК Ст. 208", "архив"])
    wb.save(path)


def make_bad_xlsx(path):
    """Файл без столбца 'Заголовок дела'."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Колонка1", "Колонка2", "Колонка3"])
    ws.append(["a", "b", "c"])
    ws.append(["d", "e", "f"])
    wb.save(path)


def make_common():
    return {
        "судебный_участок": "9",
        "судья": "О. А. Левошина",
        "дата_утверждения": "«___» ______________ 20__ года",
        "дата_акта": "«__» ____________ 20__ г.",
        "номер_акта": "1",
        "секретарь": "А.М. Намаюшка",
        "дата_подписи": "«__» ____________ 20__ г.",
        "протокол_эк_дата": "«__» ____________ 20__ г.",
        "протокол_эк_номер": "1",
    }


def docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def main():
    tmp = tempfile.mkdtemp(prefix="xlsx2word_")
    src = os.path.join(tmp, "результат.xlsx")
    make_result_xlsx(src)

    print("[1] read_result_xlsx: mapping и записи")
    records, mapping = read_result_xlsx(src)
    check("Все 8 ролей определены",
          sorted(mapping.keys()),
          sorted([k for k, _ in COLUMNS]))
    check("3 записи (включая пустую)", len(records), 3)
    check("sequential[0]=1", records[0]["sequential"], 1)
    check("title[0] начинается с 'Гражданское дело'",
          records[0]["title"].startswith("Гражданское дело"), True)
    check("retention[0]", records[0]["retention"], "3 года ЭК Ст. 137")
    check("Пустая строка title=''", records[1]["title"], "")
    check("note[2] = 'архив'", records[2]["note"], "архив")

    print("\n[2] read_result_xlsx: fallback на «Результат обработки»")
    # Оба листа, но active — «Другой»
    src2 = os.path.join(tmp, "two_sheets.xlsx")
    wb = Workbook()
    wb.active.title = "Другой"
    ws = wb.create_sheet("Результат обработки")
    ws.append([t for _, t in COLUMNS])
    ws.append([1, "Гражданское дело №2-9/2020 о защите прав", "", "", "", "", "", ""])
    wb.save(src2)
    records2, _ = read_result_xlsx(src2)
    check("Извлекли из «Результат обработки»", len(records2), 1)
    check("Заголовок есть", records2[0]["title"].startswith("Гражданское дело"), True)

    print("\n[3] Пустой лист -> ColumnsNotFound")
    expect_error("Пустой лист -> ColumnsNotFound",
                 lambda: read_result_xlsx(src2, sheet_name="Другой"),
                 ColumnsNotFound)

    print("\n[4] SheetNotFound при неверном имени")
    expect_error("Несуществующий лист -> SheetNotFound",
                 lambda: read_result_xlsx(src2, sheet_name="НетТакого"),
                 SheetNotFound)

    print("\n[5] ColumnsNotFound при отсутствии 'Заголовок дела'")
    bad = os.path.join(tmp, "bad.xlsx")
    make_bad_xlsx(bad)
    expect_error("Нет шапки -> ColumnsNotFound",
                 lambda: read_result_xlsx(bad), ColumnsNotFound)

    print("\n[6] convert_xlsx_to_word: полный цикл")
    if not os.path.exists(TEMPLATE):
        print("  SKIP шаблон не найден:", TEMPLATE)
    else:
        out = os.path.join(tmp, "акт.docx")
        res = convert_xlsx_to_word(src, TEMPLATE, out,
                                   common_values=make_common(), year="2020")
        check("Файл создан", os.path.exists(out), True)
        check("count = 2 (без пустой строки)", res["count"], 2)
        text = docx_text(out)
        check("Реквизит 'О. А. Левошина'", "О. А. Левошина" in text, True)
        check("Нет плейсхолдеров {{...}}", "{{" in text, False)
        check("Есть №2-1/2020", "2-1/2020" in text, True)
        check("Есть №2-3/2020", "2-3/2020" in text, True)
        check("Итого 2 (два)", "Итого 2 (два)" in text, True)
        # Метка присутствует
        marker = read_word_marker(out)
        check("Метка ставится", marker["found"], True)

    print("\n[7] convert_xlsx_to_word: year через None и пустые реквизиты")
    if os.path.exists(TEMPLATE):
        out2 = os.path.join(tmp, "акт_пустой.docx")
        convert_xlsx_to_word(src, TEMPLATE, out2,
                             common_values={}, year=None)
        text2 = docx_text(out2)
        check("Нет 'None' в тексте", "None" in text2, False)

    print("\n[8] read_result_xlsx: эвристика по содержимому (без шапки)")
    noheader = os.path.join(tmp, "noheader.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append([1, "Гражданское дело №2-5/2020 о защите прав", "01.01.2020",
               1, 1, 1, "3 года ЭК Ст. 208", ""])
    ws.append([2, "Гражданское дело №2-6/2020 о защите прав", "02.01.2020",
               1, 2, 1, "3 года ЭК Ст. 208", ""])
    wb.save(noheader)
    # Первая строка — уже данные, но эвристика найдёт title по содержимому.
    # sequential будет 0 (первый столбец), title — найден.
    recs_nh, mp_nh = read_result_xlsx(noheader, sheet_name=None)
    # Первая строка воспринимается как шапка -> в records её не будет,
    # в mapping['title'] укажет на столбец 1 (индекс)
    check("Эвристика нашла title", "title" in mp_nh, True)
    check("Индекс title = 1 (второй столбец)", mp_nh["title"], 1)

    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()
