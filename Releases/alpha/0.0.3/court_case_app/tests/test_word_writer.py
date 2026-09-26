# -*- coding: utf-8 -*-
"""
Тесты формирования Word-документа по шаблону акта (Этап 7).

Использует копию шаблона court_case_app/docs/АКТ уничтожения гражданских дел.docx.

Запуск: .venv\\Scripts\\python.exe court_case_app\\tests\\test_word_writer.py
"""

import sys
import os
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document
from core.word_writer import (
    write_word_result,
    number_to_words,
    _find_sample_row,
)

FAILED = 0
PASSED = 0

TEMPLATE = os.path.join(os.path.dirname(__file__), "..", "docs",
                        "АКТ уничтожения гражданских дел.docx")


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def make_records():
    return [
        {"sequential": 1,
         "title": "Гражданское дело №2-3/2019 по заявлению Саркисян Юлии "
                  "Геннадьевны к Саркисяну Давиду Сергеевичу о расторжении брака "
                  "супругов - имеющих детей",
         "dates": "26.11.2018\n26.12.2018",
         "opis": 1, "unit": 1, "count": 1,
         "retention": "3 года ЭК Ст. 137", "note": ""},
        # Пустая запись (пропущенный номер) — не должна попасть в Word
        {"sequential": 2, "title": "", "dates": "", "opis": "", "unit": "",
         "count": "", "retention": "", "note": ""},
        {"sequential": 3,
         "title": 'Гражданское дело №2-12/2019 по иску АО "МАКС" к '
                  "Гринчевой Оксане Геннадьевне о возмещении имущественного вреда",
         "dates": "21.12.2018\n21.01.2019",
         "opis": 1, "unit": 9, "count": 1,
         "retention": "3 года Ст. 227", "note": "архив"},
    ]


def make_common_values():
    return {
        "судебный_участок": "3",
        "судья": "О. А. Левошина",
        "дата_утверждения": "«___» ______________ 2026 года",
        "дата_акта": "«__» ____________ 2026 г.",
        "номер_акта": "1",
        "секретарь": "А.М. Намаюшка",
        "дата_подписи": "«__» ____________ 2026 г.",
        "протокол_эк_дата": "«__» ____________ 2026 г.",
        "протокол_эк_номер": "1",
        "год_дел": "2019",
    }


def doc_full_text(doc) -> str:
    """Весь текст документа (абзацы + ячейки таблиц)."""
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    parts.append(p.text)
    return "\n".join(parts)


def main():
    check("Шаблон существует", os.path.exists(TEMPLATE), True)

    # Предварительная проверка шаблона: наличие строки-образца
    doc_tpl = Document(TEMPLATE)
    found_any = any(_find_sample_row(t._tbl) is not None for t in doc_tpl.tables)
    check("В шаблоне найдена строка-образец", found_any, True)
    tpl_text = doc_full_text(doc_tpl)
    check("Шаблон содержит {{дата_подписи}} (правки внесены)",
          "{{дата_подписи}}" in tpl_text, True)

    print("\n[1] Числа прописью")
    check("0", number_to_words(0), "ноль")
    check("1", number_to_words(1), "один")
    check("21", number_to_words(21), "двадцать один")
    check("374", number_to_words(374), "триста семьдесят четыре")
    check("6374 (пример из акта)", number_to_words(6374),
          "шесть тысяч триста семьдесят четыре")
    check("21000", number_to_words(21000), "двадцать одна тысяча")
    check("1000000", number_to_words(1000000), "один миллион")

    print("\n[2] Формирование документа")
    tmp_dir = tempfile.mkdtemp(prefix="word_out_")
    out_path = os.path.join(tmp_dir, "акт_уничтожения.docx")
    write_word_result(make_records(), TEMPLATE, out_path, make_common_values())
    check("Файл создан", os.path.exists(out_path), True)

    doc = Document(out_path)
    text = doc_full_text(doc)

    # Нет оставшихся плейсхолдеров
    import re
    remaining = re.findall(r"\{\{[^}]+\}\}", text)
    check("Нет оставшихся плейсхолдеров {{...}}", remaining, [])

    # Реквизиты
    check("Итоговая строка",
          "Итого 2 (два) гражданских дел за 2019 год." in text, True)
    check("Судья подставлен", "О. А. Левошина" in text, True)
    check("Судебный участок", "Судебный участок № 3" in text, True)
    check("Секретарь", "А.М. Намаюшка" in text, True)
    check("Номер акта", "№ 1" in text, True)

    # Таблица данных дел (в документе несколько таблиц — ищем по заголовкам дел)
    tables_with_data = [
        t for t in doc.tables
        if any("Гражданское дело" in (c.text or "")
               for r in t.rows for c in r.cells)
    ]
    check("Найдена таблица со строками дел", len(tables_with_data), 1)
    table = tables_with_data[0]
    rows = len(table.rows)
    # Структура шаблона: Word разбил таблицу акта на секции.
    # Таблица данных: строка нумерации 1..8 + строки дел.
    # После обработки: нумерация + 2 строки данных (пустая запись исключена).
    check("Строк в таблице данных: нумерация + 2 дела", rows, 3)

    row_texts = [" ".join(c.text for c in r.cells) for r in table.rows]
    any_title = any("Гражданское дело №2-3/2019" in t for t in row_texts)
    check("Строка №2-3 присутствует", any_title, True)
    cases_12 = [t for t in row_texts if "Гражданское дело №2-12/2019" in t]
    check("Строка №2-12 присутствует", len(cases_12) == 1, True)

    # Проверка формата дат в ячейке таблицы
    dates_found = [t for t in row_texts if "26.11.2018" in t]
    check("Даты дела в таблице", len(dates_found) == 1, True)

    print("\n[3] Регрессия fix_03: None в плейсхолдерах")
    # Реквизит None -> пустая строка, а не литерал «None»
    tmp_dir2 = tempfile.mkdtemp(prefix="word_none_")
    out2 = os.path.join(tmp_dir2, "акт_none.docx")
    common_none = dict(make_common_values())
    common_none["судья"] = None
    common_none["номер_акта"] = None
    write_word_result(make_records(), TEMPLATE, out2, common_none)
    doc2 = Document(out2)
    text2 = doc_full_text(doc2)
    check("fix_03: в тексте нет литерала «None»", "None" in text2, False)
    check("fix_03: плейсхолдеры не остались",
          "{{судья}}" in text2 or "{{номер_акта}}" in text2, False)

    # Отсутствующий год не ломает итоговую строку
    tmp_dir3 = tempfile.mkdtemp(prefix="word_noyear_")
    out3 = os.path.join(tmp_dir3, "акт_noyear.docx")
    common_no_year = dict(make_common_values())
    common_no_year.pop("год_дел", None)
    write_word_result(make_records(), TEMPLATE, out3, common_no_year)
    text3 = doc_full_text(Document(out3))
    check("fix_03: нет «None» при отсутствии года", "None" in text3, False)

    print("\n[4] Регрессия fix_10: переносы строк через <w:br/>")
    from docx import Document as _D
    from core.word_writer import write_word_result as _wwr
    tmp_dir_br = tempfile.mkdtemp(prefix="word_br_")
    out_br = os.path.join(tmp_dir_br, "акт.docx")
    records_br = [{
        "sequential": 1,
        "title": "Гражданское дело №2-1/2020 о защите прав",
        "dates": "08.12.2017\n08.01.2018",
        "opis": 1, "unit": 1, "count": 1,
        "retention": "3 года Ст. 208", "note": "",
    }]
    _wwr(records_br, TEMPLATE, out_br, make_common_values())
    _doc_br = _D(out_br)
    dates_cells = []
    for _t in _doc_br.tables:
        for _r in _t.rows:
            for _c in _r.cells:
                if "08.12.2017" in (_c.text or ""):
                    dates_cells.append(_c.text)
    check("Найдена ячейка с датой", len(dates_cells) >= 1, True)
    check("Дата содержит \n между двумя датами",
          any("08.12.2017\n08.01.2018" in (t or "") for t in dates_cells), True)
    check("Нет артефакта _x000D_",
          any("_x000D_" in (t or "") for t in dates_cells), False)

    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()