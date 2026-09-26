# -*- coding: utf-8 -*-
"""
Тесты шаблона «Таблица» и записи результата (этап 14d).

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tests\\test_inventory_sheets_io.py
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import Workbook, load_workbook

from core.inventory.exceptions import TemplateSheetNotFound
from core.inventory.sheets_io import (
    detect_last_template_row,
    extend_formulas,
    extend_vlookup_ranges,
    find_template_sheet,
    shift_formula,
)
from core.inventory.writer import page_count, prepare_rows, write_result

FAILED = 0
PASSED = 0


def check(name, actual, expected):
    global FAILED, PASSED
    if actual == expected:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n"
              f"       получено:  {actual!r}")


def expect_error(name, fn, exc_type):
    global FAILED, PASSED
    try:
        fn()
    except exc_type:
        PASSED += 1
        print(f"  OK   {name}")
        return
    except Exception as exc:  # noqa: BLE001
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось {exc_type.__name__}, "
              f"получено {type(exc).__name__}: {exc}")
        return
    FAILED += 1
    print(f"  FAIL {name}\n       исключение не поднято")


def make_template(path, last_row=6):
    """Шаблон с листом «Таблица»: строки 4..last_row — данные."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Таблица"
    for row in range(4, last_row + 1):
        ws.cell(row=row, column=1).value = f"=ROW()"
        ws.cell(row=row, column=2).value = f"=IF(F{row}=F{row-1},B{row-1},1)"
        ws.cell(row=row, column=3).value = None
        ws.cell(row=row, column=4).value = "шаблонD"
        ws.cell(row=row, column=6).value = f"=F{row-1}+E{row}"
        ws.cell(row=row, column=7).value = 2018
        ws.cell(row=row, column=9).value = f"=I{row-1}+1"
        ws.cell(row=row, column=10).value = None
        ws.cell(row=row, column=11).value = f"=ROW()"
        ws.cell(row=row, column=12).value = f"=A{row}"
        ws.cell(row=row, column=14).value = "Том"
        ws.cell(row=row, column=15).value = \
            f"=VLOOKUP(O{row}, $A$4:$F$10003, 6, 0)"
    wb.save(path)
    return path


def main():
    tmp = tempfile.mkdtemp(prefix="inventory_xlsx_")

    print("[1] shift_formula")
    check("относительные сдвигаются",
          shift_formula("=IF(F5=F4,A4,A4+1)", 2), "=IF(F7=F6,A6,A6+1)")
    check("абсолютные не сдвигаются",
          shift_formula("=VLOOKUP(O4,$A$4:$F$10003,6,0)", 3),
          "=VLOOKUP(O7,$A$4:$F$10003,6,0)")
    check("не формула", shift_formula("текст", 2), "текст")

    print("\n[2] find_template_sheet")
    tpl = make_template(os.path.join(tmp, "tpl.xlsx"), last_row=6)
    wb_ok = load_workbook(tpl)
    check("лист найден", find_template_sheet(wb_ok).title, "Таблица")
    wb_bad = Workbook()
    wb_bad.active.title = "Иное"
    expect_error("нет листа «Таблица»",
                 lambda: find_template_sheet(wb_bad),
                 TemplateSheetNotFound)

    print("\n[3] detect_last_template_row / clean_template")
    check("последняя строка",
          detect_last_template_row(wb_ok["Таблица"]), 6)
    from core.inventory.sheets_io import clean_template
    clean_template(wb_ok["Таблица"], 6)
    check("C пусто (C4=None) -> None",
          wb_ok["Таблица"].cell(row=5, column=3).value, None)
    check("D реплика", wb_ok["Таблица"].cell(row=5, column=4).value,
          "шаблонD")
    check("E = 1", wb_ok["Таблица"].cell(row=6, column=5).value, 1)
    check("G реплика", wb_ok["Таблица"].cell(row=4, column=7).value, 2018)
    check("H = 1", wb_ok["Таблица"].cell(row=4, column=8).value, 1)

    print("\n[4] extend_formulas: только A,B,F,I,J,K,L")
    wb2 = load_workbook(tpl)
    ws2 = wb2["Таблица"]
    extend_formulas(ws2, 6, 8)
    check("A7 протянута со сдвигом",
          ws2.cell(row=7, column=1).value, "=ROW()")
    check("F8 со сдвигом формулы",
          ws2.cell(row=8, column=6).value, "=F7+E8")
    check("O8 НЕ протянута (None)",
          ws2.cell(row=8, column=15).value, None)
    check("E8 не заполнена формулой",
          ws2.cell(row=8, column=5).value, None)

    print("\n[5] extend_vlookup_ranges")
    wb3 = load_workbook(tpl)
    ws3 = wb3["Таблица"]
    new_end = extend_vlookup_ranges(ws3, 8)
    check("конец диапазона", new_end, 108)
    check("диапазон расширен",
          ws3.cell(row=4, column=15).value,
          "=VLOOKUP(O4, $A$4:$F$108, 6, 0)")

    print("\n[6] page_count")
    check("один лист", page_count("8"), (1, ""))
    check("диапазон", page_count("14-15"), (2, ""))
    check("диапазон 2", page_count("220-221"), (2, ""))
    check("битый 13135", page_count("13135"),
          (1, "битый диапазон листов"))
    check("битый 133-13135", page_count("133-13135"),
          (1, "битый диапазон листов"))
    check("пусто", page_count(""), (1, "пустые номера листов"))

    print("\n[7] prepare_rows")
    flow = {"attention": [], "rows": [
        {"file": "f.docx", "word_row": 2, "number_raw": "1",
         "name": "Решение по гражданскому делу № 2-14/2018",
         "sheets_raw": "8"},
        {"file": "f.docx", "word_row": 3, "number_raw": "2",
         "name": "Просто текст", "sheets_raw": "1"},
        {"file": "f.docx", "word_row": 4, "number_raw": "3",
         "name": "Определение по гражданскому делу № 2-15/2018",
         "sheets_raw": "14-15"},
    ]}
    prep = prepare_rows(flow)
    check("строк записано (без типа пропущена)", len(prep["rows"]), 2)
    check("f4", prep["f4"], 14)
    check("C первой", prep["rows"][0]["C"],
          "Решение по гражданскому делу")
    check("E первой None (F4 из Word)", prep["rows"][0]["E"], None)
    check("E второй = 1 (сквозной)", prep["rows"][1]["E"], 1)
    check("G", prep["rows"][0]["G"], 2018)
    check("H второй = 2", prep["rows"][1]["H"], 2)
    check("attention о пропуске",
          any("не распознан" in a["reason"] for a in prep["attention"]),
          True)

    print("\n[8] write_result: сквозная запись")
    tpl2 = make_template(os.path.join(tmp, "tpl2.xlsx"), last_row=4)
    out = os.path.join(tmp, "out.xlsx")
    write_result(tpl2, flow, out)
    wb_out = load_workbook(out)
    ws_out = wb_out["Таблица"]
    check("C4", ws_out.cell(row=4, column=3).value,
          "Решение по гражданскому делу")
    check("D4", ws_out.cell(row=4, column=4).value, "2")
    check("F4", ws_out.cell(row=4, column=6).value, 14)
    check("G4", ws_out.cell(row=4, column=7).value, 2018)
    check("H4", ws_out.cell(row=4, column=8).value, 1)
    check("C5", ws_out.cell(row=5, column=3).value,
          "Определение по гражданскому делу")
    check("E5 = 1", ws_out.cell(row=5, column=5).value, 1)
    check("H5 = 2", ws_out.cell(row=5, column=8).value, 2)
    check("лист внимания добавлен",
          "Требует внимания" in wb_out.sheetnames, True)
    att = wb_out["Требует внимания"]
    check("заголовки внимания", att.cell(row=1, column=1).value, "Файл")
    check("есть записи внимания", att.cell(row=2, column=4).value is not None,
          True)

    print("\n" + "=" * 60)
    print(f"ИТОГО: PASSED = {PASSED}, FAILED = {FAILED}")
    return 1 if FAILED else 0


if __name__ == "__main__":
    sys.exit(main())