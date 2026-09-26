# -*- coding: utf-8 -*-
"""
Тесты формирования Excel-результата (Этап 6).

Запуск: .venv\\Scripts\\python.exe court_case_app\\tests\\test_excel_writer.py
"""

import sys
import os
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import load_workbook
from core.excel_writer import write_excel_result, DEFAULT_COLUMNS, DEFAULT_KEYS

FAILED = 0
PASSED = 0


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def make_records():
    return [
        {"sequential": 1,
         "title": "Гражданское дело №2-3/2019 по заявлению Саркисян Юлии "
                  "Геннадьевны к Саркисяну Давиду Сергеевичу о расторжении брака",
         "dates": "26.11.2018\n26.12.2018",
         "opis": 1, "unit": 1, "count": 1,
         "retention": "3 года ЭК Ст. 137", "note": ""},
        {"sequential": 2, "title": "", "dates": "", "opis": "", "unit": "",
         "count": "", "retention": "", "note": ""},
        {"sequential": 3,
         "title": 'Гражданское дело №2-12/2019 по иску АО "МАКС" к '
                  "Гринчевой Оксане Геннадьевне о возмещении имущественного вреда",
         "dates": "21.12.2018\n21.01.2019",
         "opis": 1, "unit": 9, "count": 1,
         "retention": "3 года Ст. 227", "note": "архив"},
    ]


def main():
    tmp_dir = tempfile.mkdtemp(prefix="excel_out_")
    full_path = os.path.join(tmp_dir, "result_full.xlsx")
    subset_path = os.path.join(tmp_dir, "result_subset.xlsx")

    print("[1] Запись файла со всеми 8 столбцами")
    write_excel_result(make_records(), full_path)
    check("Файл создан", os.path.exists(full_path), True)

    wb = load_workbook(full_path)
    ws = wb.active
    check("Имя листа", ws.title, "Результат обработки")
    headers = [ws.cell(row=1, column=c).value for c in range(1, 9)]
    check("Заголовки",
          headers,
          ["№ п/п", "Заголовок дела", "Даты дела", "№ описи", "№ ед.хр.",
           "Кол-во ед.хр.", "Срок хранения", "Примечание"])
    check("Количество строк (3 записи + заголовок)", ws.max_row, 4)
    check("№ п/п первой строки", ws.cell(row=2, column=1).value, 1)
    check("Заголовок дела первой строки содержит №2-3",
          str(ws.cell(row=2, column=2).value).startswith("Гражданское дело №2-3/2019"),
          True)
    check("Даты с переносом строки",
          ws.cell(row=2, column=3).value, "26.11.2018\n26.12.2018")
    check("Срок хранения", ws.cell(row=2, column=7).value, "3 года ЭК Ст. 137")
    check("Пустая строка (пропуск): заголовок пуст",
          ws.cell(row=3, column=2).value in (None, ""), True)
    check("Пустая строка: срок пуст",
          ws.cell(row=3, column=7).value in (None, ""), True)

    # Форматирование
    check("Заголовок жирный", ws.cell(row=1, column=1).font.bold, True)
    check("Заголовок с заливкой",
          ws.cell(row=1, column=1).fill.start_color.rgb is not None, True)
    check("Границы у ячейки данных",
          ws.cell(row=2, column=2).border.left.style is not None, True)
    check("Перенос текста в столбце B",
          ws.cell(row=2, column=2).alignment.wrap_text, True)
    check("Ширина столбца B ограничена (~60)",
          ws.column_dimensions["B"].width <= 62, True)
    wb.close()

    print("\n[2] Выбор подмножества столбцов")
    subset = ["sequential", "title", "retention"]
    write_excel_result(make_records(), subset_path, columns=subset)
    wb2 = load_workbook(subset_path)
    ws2 = wb2.active
    headers2 = [ws2.cell(row=1, column=c).value for c in range(1, 4)]
    check("Заголовки подмножества",
          headers2, ["№ п/п", "Заголовок дела", "Срок хранения"])
    check("Нет лишнего столбца (max_column=3)", ws2.max_column, 3)
    check("Данные подмножества",
          ws2.cell(row=2, column=3).value, "3 года ЭК Ст. 137")
    wb2.close()

    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()