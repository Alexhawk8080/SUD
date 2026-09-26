# -*- coding: utf-8 -*-
"""
Запись результата «Внутренней описи» (этап 14d, решения 2, 8, 9, 19–21).

Конвертер заполняет ТОЛЬКО:
    C — нормализованное наименование;
    D — префикс выбранного номера;
    E — приращение номера (E4 не заполняется);
    G — год;
    H — количество страниц;
    F4 — числовая часть первого номера потока.

Остальные столбцы — формулы шаблона.

Лист «Требует внимания» (решение 20):
    Файл | Строка в Word | № п/п | Причина | Что сделали.
"""

import re

from .case_number import build_from_rows
from .normalizer import normalize_name
from .sheets_io import (
    FIRST_DATA_ROW,
    clean_template,
    detect_last_template_row,
    extend_formulas,
    extend_vlookup_ranges,
    find_template_sheet,
)

ATTENTION_COLUMNS = ("Файл", "Строка в Word", "№ п/п", "Причина",
                     "Что сделали")

MAX_PAGES = 999

_SINGLE_RE = re.compile(r"^(\d+)$")
_RANGE_RE = re.compile(r"^(\d+)\s*-\s*(\d+)$")


def page_count(sheets_raw):
    """
    H = J − I + 1 из «Номера листов» (решение 8).

        «8»         -> (1, "")
        «14-15»     -> (2, "")
        «220-221»   -> (2, "")
        «13135»     -> (1, "битый диапазон листов")
        «133-13135» -> (1, "битый диапазон листов")
    """
    if sheets_raw is None:
        return 1, "пустые номера листов"
    text = " ".join(str(sheets_raw).split())
    if not text:
        return 1, "пустые номера листов"
    m = _RANGE_RE.match(text)
    if m:
        start, end = int(m.group(1)), int(m.group(2))
        if end >= start and (end - start + 1) <= MAX_PAGES:
            return end - start + 1, ""
        return 1, "битый диапазон листов"
    m = _SINGLE_RE.match(text)
    if m:
        if int(m.group(1)) <= MAX_PAGES:
            return 1, ""
        return 1, "битый диапазон листов"
    return 1, "битый диапазон листов"


def prepare_rows(flow) -> dict:
    """
    Формирует данные для записи в «Таблицу» (C/D/E/G/H) и attention.

    flow — результат flow.build_flow. Строки без распознанного типа
    пропускаются + «Требует внимания» (решение 19). Нумерация сквозная:
    E = F_i − F_{i−1} по записанным строкам.
    """
    records = []
    attention = []
    for row in flow.get("rows", []):
        doc_type = normalize_name(row.get("name"))
        if not doc_type:
            attention.append({
                "file": row.get("file", ""),
                "word_row": row.get("word_row", ""),
                "number": row.get("number_raw", ""),
                "reason": "тип дела не распознан",
                "action": "строка пропущена",
            })
            continue
        records.append({**row, "C": doc_type})

    stream = build_from_rows(records)
    out_rows = []
    for rec, srow in zip(records, stream["rows"]):
        h, reason = page_count(rec.get("sheets_raw"))
        if reason:
            attention.append({
                "file": rec.get("file", ""),
                "word_row": rec.get("word_row", ""),
                "number": rec.get("number_raw", ""),
                "reason": reason, "action": "H=1",
            })
        if srow.get("attention"):
            attention.append({
                "file": rec.get("file", ""),
                "word_row": rec.get("word_row", ""),
                "number": rec.get("number_raw", ""),
                "reason": srow["attention"], "action": "год заменён",
            })
        out_rows.append({
            "C": rec["C"], "D": srow["D"], "E": srow["E"],
            "F": srow.get("F"), "G": srow["G"], "H": h,
            "word_row": rec.get("word_row"),
            "file": rec.get("file"), "number": rec.get("number_raw"),
            "case": f"{srow['D']}-{srow.get('F')}/{srow['G']}",
        })

    return {"rows": out_rows, "f4": stream.get("f4"),
            "year": stream.get("year"),
            "attention": attention + list(flow.get("attention", []))}


def _write_attention_sheet(ws, attention) -> None:
    """Лист «Требует внимания» (решение 20)."""
    for col, title in enumerate(ATTENTION_COLUMNS, start=1):
        ws.cell(row=1, column=col).value = title
    for i, item in enumerate(attention, start=2):
        ws.cell(row=i, column=1).value = item.get("file", "")
        ws.cell(row=i, column=2).value = item.get("word_row", "")
        ws.cell(row=i, column=3).value = item.get("number", "")
        ws.cell(row=i, column=4).value = item.get("reason", "")
        ws.cell(row=i, column=5).value = item.get("action", "")


def write_result(template_path: str, flow: dict, out_path: str) -> dict:
    """
    Заполняет шаблон и сохраняет результат (решения 10–13).

    Возвращает результат prepare_rows (rows, f4, year, attention).
    """
    from openpyxl import load_workbook

    prepared = prepare_rows(flow)
    workbook = load_workbook(template_path)   # формулы сохраняются
    ws = find_template_sheet(workbook)

    last_row = detect_last_template_row(ws)
    n = len(prepared["rows"])
    needed_last = FIRST_DATA_ROW + max(n, 1) - 1
    target_last = max(last_row, needed_last)

    clean_template(ws, max(last_row, needed_last))
    extend_formulas(ws, last_row, target_last)

    if prepared["f4"] is not None:
        ws.cell(row=FIRST_DATA_ROW, column=6).value = prepared["f4"]

    for i, row in enumerate(prepared["rows"]):
        r = FIRST_DATA_ROW + i
        ws.cell(row=r, column=3).value = row["C"]
        ws.cell(row=r, column=4).value = row["D"]
        # E4 не заполняется (решение 5): None очищает значение очистки.
        ws.cell(row=r, column=5).value = row["E"]
        ws.cell(row=r, column=7).value = row["G"]
        ws.cell(row=r, column=8).value = row["H"]

    extend_vlookup_ranges(ws, target_last)

    if "Требует внимания" in workbook.sheetnames:
        del workbook["Требует внимания"]
    att_ws = workbook.create_sheet("Требует внимания")
    _write_attention_sheet(att_ws, prepared["attention"])

    workbook.save(out_path)
    return prepared
