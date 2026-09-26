# -*- coding: utf-8 -*-
"""
Чтение ранее обработанных файлов («старой версии») для проверки категорий
(Доработка 6).

Поддерживаются:
    - результат обработки .xlsx (лист «Результат обработки», колонка
      «Заголовок дела»);
    - акт уничтожения .docx (ячейки таблиц с текстом «Гражданское дело №...»).

Из заголовка извлекаются номер дела, заявители и категория «о ...».
Категория — хвост заголовка, начинающийся с «о/об»; если его нет,
категория считается пустой (дело проблемное).
"""

import re

from .category_check import extract_phrase

# Заголовок дела: «Гражданское дело №2-3/2023 по иску Истец к Ответчик о ...»
_TITLE_RE = re.compile(
    r"Гражданское дело\s+№(?P<num>[0-9]+-[0-9]+/[0-9]{4})"
    r"\s+(?P<prefix>по иску|по заявлению)"
    r"(?:\s+(?P<applicants>.+?))?\s+к\s+(?P<rest>.+)",
    re.IGNORECASE)

# Фраза «о ...» / «об ...» для отделения ответчика от категории
_CATEGORY_START_RE = re.compile(r"(?:^|\s)(о|об)\s")


def parse_title(title) -> dict | None:
    """
    Разбор заголовка дела из готового результата.

    Возвращает словарь {case_number, applicants, respondents, category,
    prefix, title} или None, если текст не похож на заголовок дела.
    """
    if not title:
        return None
    text = " ".join(str(title).split())
    m = _TITLE_RE.match(text)
    if not m:
        return None
    applicants = (m.group("applicants") or "").strip()
    rest = m.group("rest").strip()

    # Ответчик — часть rest до фразы «о/об»; категория — с этой фразы
    cm = _CATEGORY_START_RE.search(rest.lower())
    if cm:
        respondents = rest[:cm.start()].strip()
        phrase = rest[cm.start(1):].strip()
    else:
        respondents = rest.strip()
        phrase = ""

    return {
        "case_number": m.group("num"),
        "prefix": m.group("prefix"),
        "applicants": applicants,
        "respondents": respondents,
        "category": phrase,          # «о ...» или "" (дело проблемное)
        "title": text,
    }


def parse_result_xlsx(filepath) -> list:
    """
    Чтение файла-результата .xlsx (лист «Результат обработки»).

    Колонки (как в excel_writer): № п/п, Заголовок дела, Даты дела,
    № описи, № ед.хр., Кол-во ед.хр., Срок хранения, Примечание.

    Возвращает список записей parse_title (пустые строки пропускаются).
    """
    from openpyxl import load_workbook

    wb = load_workbook(filepath, data_only=True, read_only=True)
    try:
        ws = (wb["Результат обработки"]
              if "Результат обработки" in wb.sheetnames else wb.active)
        records = []
        for row in ws.iter_rows(values_only=True):
            if not row or len(row) < 2 or not row[1]:
                continue
            parsed = parse_title(row[1])
            if parsed:
                records.append(parsed)
        return records
    finally:
        wb.close()


def parse_act_docx(filepath) -> list:
    """
    Чтение акта уничтожения .docx: поиск ячеек таблиц с текстом
    «Гражданское дело №...» (заголовки дел).

    Возвращает список записей parse_title.

    fix_05_v2: обходим XML-узлы w:tc напрямую (table._tbl.iter).
    Каждый узел w:tc уникален в дереве, поэтому объединённые ячейки
    (merge) не дают дубликатов, а разные ячейки не схлопываются.
    """
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(filepath)
    records = []
    for table in doc.tables:
        for tc in table._tbl.iter(qn("w:tc")):
            text = " ".join(
                (t.text or "") for t in tc.iter(qn("w:t"))
            )
            parsed = parse_title(text)
            if parsed:
                records.append(parsed)
    return records


def parse_legacy_file(filepath) -> list:
    """
    Универсальное чтение старого результата по расширению файла.

    Возвращает список записей (пустой, если ничего не найдено).
    """
    lower = str(filepath).lower()
    if lower.endswith(".docx"):
        return parse_act_docx(filepath)
    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        return parse_result_xlsx(filepath)
    return []