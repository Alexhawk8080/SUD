# -*- coding: utf-8 -*-
"""
Работа с листами .xlsx: проверка содержимого, список, default-лист.

Проверка содержимого (решение 34а): открываем как zip и убеждаемся,
что внутри есть xl/workbook.xml — это гарантирует, что файл действительно
Excel, даже если расширение подменили.
"""

import zipfile
from pathlib import Path

from .exceptions import BadFileError, NoSheetsFound

# Имя листа с результатом обработки (наш канонический Excel-результат)
RESULT_SHEET_NAME = "Результат обработки"


def verify_xlsx_content(filepath) -> None:
    """
    Проверяет, что файл — валидный xlsx (zip с xl/workbook.xml).

    Бросает BadFileError с понятным сообщением.
    """
    path = Path(filepath)
    if not path.exists():
        raise BadFileError(f"Файл не найден: {path.name}")
    if not zipfile.is_zipfile(str(path)):
        raise BadFileError(
            "Файл не является Excel-документом (не zip-архив). "
            "Возможно, расширение подменено или файл повреждён.")
    try:
        with zipfile.ZipFile(str(path)) as zf:
            names = set(zf.namelist())
    except zipfile.BadZipFile as exc:
        raise BadFileError(f"Файл повреждён: {exc}") from exc

    if "xl/workbook.xml" not in names:
        raise BadFileError(
            "Файл не содержит xl/workbook.xml — это не xlsx-книга.")


def inspect_sheets(filepath) -> dict:
    """
    Возвращает словарь:
        {
          "sheets":  ["Лист1", "Результат обработки", ...],
          "active":  "Лист1",
          "default": "Результат обработки" | <первый активный>
        }

    «default» — предпочтительный лист для конвертации: «Результат обработки»,
    если он есть; иначе активный (или первый).

    Бросает BadFileError при проблемах с файлом.
    """
    verify_xlsx_content(filepath)

    from openpyxl import load_workbook
    try:
        wb = load_workbook(filepath, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise BadFileError(f"Не удалось открыть книгу: {exc}") from exc

    try:
        sheets = list(wb.sheetnames)
        if not sheets:
            raise NoSheetsFound("В файле нет ни одного листа.")
        try:
            active = wb.active.title if wb.active else sheets[0]
        except Exception:  # noqa: BLE001
            active = sheets[0]
        default = RESULT_SHEET_NAME if RESULT_SHEET_NAME in sheets else active
        return {"sheets": sheets, "active": active, "default": default}
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass


def sheet_exists(filepath, sheet_name: str) -> bool:
    """Проверка, что лист с указанным именем есть в книге."""
    return sheet_name in inspect_sheets(filepath)["sheets"]


def verify_docx_content(filepath) -> None:
    """
    Проверяет, что файл — валидный docx (zip с word/document.xml).
    """
    path = Path(filepath)
    if not path.exists():
        raise BadFileError(f"Файл не найден: {path.name}")
    if not zipfile.is_zipfile(str(path)):
        raise BadFileError(
            "Файл не является Word-документом (не zip-архив).")
    try:
        with zipfile.ZipFile(str(path)) as zf:
            names = set(zf.namelist())
    except zipfile.BadZipFile as exc:
        raise BadFileError(f"Файл повреждён: {exc}") from exc
    if "word/document.xml" not in names:
        raise BadFileError(
            "Файл не содержит word/document.xml — это не docx-документ.")
