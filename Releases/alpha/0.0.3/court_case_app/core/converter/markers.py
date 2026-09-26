# -*- coding: utf-8 -*-
"""
Невидимые метки «наш файл» для Word-акта и Excel-результата.

Метка — гибрид: zero-width space (U+200B) в качестве обёртки + читаемый
ASCII-код. В Word хранится в core_properties.keywords и core_properties.comments,
в Excel — в wb.properties.keywords и wb.properties.description (дублирование
на случай, если пользователь затрёт одно поле).

Формат метки:
    \u200BACT-V{n}-COURT-CASE-APP\u200B
где n — версия формата.

Публичный API:
    WORD_MARKER, EXCEL_MARKER   — константы текущей версии
    CURRENT_MARKER_VERSION      — "V1"
    make_marker(version)        — сформировать строку метки
    parse_marker(text)          — извлечь версию из строки метки
    mark_word_document(doc)     — проставить метку в Document (python-docx)
    read_word_marker(filepath)  — прочитать метку из .docx
    mark_excel_workbook(wb)     — проставить метку в Workbook (openpyxl)
    read_excel_marker(filepath) — прочитать метку из .xlsx/.xlsm
"""

import re

CURRENT_MARKER_VERSION = "V1"

# Zero-width space: невидим, но стабильно сохраняется в XML.
_ZW = "\u200B"

# Регулярка для извлечения версии: ACT-(V\d+)-COURT-CASE-APP
_MARKER_RE = re.compile(r"ACT-(V\d+)-COURT-CASE-APP")


def make_marker(version: str = None) -> str:
    """Сформировать строку метки заданной версии."""
    v = (version or CURRENT_MARKER_VERSION).strip().upper()
    if not v.startswith("V"):
        v = "V" + v
    return f"{_ZW}ACT-{v}-COURT-CASE-APP{_ZW}"


# Готовые константы текущей версии
WORD_MARKER = make_marker(CURRENT_MARKER_VERSION)
EXCEL_MARKER = make_marker(CURRENT_MARKER_VERSION)


def parse_marker(text) -> str:
    """
    Извлечь версию метки из строки. Пустая строка, если метки нет.
    Возвращает, например, "V1".
    """
    if not text:
        return ""
    m = _MARKER_RE.search(str(text))
    return m.group(1) if m else ""


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

def mark_word_document(doc) -> None:
    """
    Проставить метку в Document. Перезаписывает существующую (обновляет
    до текущей версии).
    """
    cp = doc.core_properties
    cp.keywords = WORD_MARKER
    cp.comments = WORD_MARKER


def read_word_marker(filepath) -> dict:
    """
    Прочитать метку из .docx.

    Возвращает {"found": bool, "version": str, "raw": str}:
        found   — True, если метка найдена хотя бы в одном поле;
        version — версия метки ("V1", "V2", ...) или "";
        raw     — исходная строка метки (из первого найденного поля).
    """
    from docx import Document
    try:
        doc = Document(filepath)
    except Exception:  # noqa: BLE001
        return {"found": False, "version": "", "raw": ""}

    cp = doc.core_properties
    for raw in (cp.keywords, cp.comments):
        v = parse_marker(raw)
        if v:
            return {"found": True, "version": v, "raw": raw or ""}
    return {"found": False, "version": "", "raw": ""}


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def mark_excel_workbook(wb) -> None:
    """
    Проставить метку в Workbook (openpyxl). Перезаписывает существующую.
    """
    props = wb.properties
    props.keywords = EXCEL_MARKER
    props.description = EXCEL_MARKER


def read_excel_marker(filepath) -> dict:
    """
    Прочитать метку из .xlsx/.xlsm.

    Возвращает {"found": bool, "version": str, "raw": str}.
    """
    from openpyxl import load_workbook
    try:
        wb = load_workbook(filepath, read_only=True)
    except Exception:  # noqa: BLE001
        return {"found": False, "version": "", "raw": ""}

    try:
        props = wb.properties
        for raw in (props.keywords, props.description):
            v = parse_marker(raw)
            if v:
                return {"found": True, "version": v, "raw": raw or ""}
        return {"found": False, "version": "", "raw": ""}
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
