# -*- coding: utf-8 -*-
"""
Генерация фикстур модуля «Внутренняя опись» (этап 14f).

Создаёт:
    court_case_app/tests/data/inventory/source_demo.docx  — docx-опись;
    court_case_app/tests/data/inventory/template_demo.xlsx — шаблон «Таблица».

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tools\\make_inventory_fixtures.py
"""

import os

from docx import Document
from openpyxl import Workbook

DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "..", "tests", "data", "inventory")


def make_docx(path):
    doc = Document()
    doc.add_paragraph("Внутренняя опись дел")
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(("№ п/п", "Наименование", "Номера листов")):
        table.cell(0, i).text = h
    rows = [
        ("Решение по гражданскому делу № 2-14/2018", "1-5"),
        ("Заочное решение по гражданскому делу № 2-15/2018", "6"),
        ("Определение по гражданскому делу № 2-16/2018", "7-9"),
        ("Решение по гражданскому делу № 2-16/2018", "10"),
    ]
    for idx, (name, sheets) in enumerate(rows, start=1):
        row = table.add_row()
        row.cells[0].text = str(idx)
        row.cells[1].text = name
        row.cells[2].text = sheets
    doc.save(path)
    return path


def make_template(path, last_row=6):
    wb = Workbook()
    ws = wb.active
    ws.title = "Таблица"
    for row in range(4, last_row + 1):
        ws.cell(row=row, column=1).value = "=ROW()"
        ws.cell(row=row, column=2).value = f"=IF(F{row}=F{row-1},B{row-1},1)"
        ws.cell(row=row, column=4).value = "шаблон"
        ws.cell(row=row, column=6).value = f"=F{row-1}+E{row}"
        ws.cell(row=row, column=7).value = 2018
        ws.cell(row=row, column=9).value = f"=I{row-1}+1"
        ws.cell(row=row, column=11).value = "=ROW()"
        ws.cell(row=row, column=12).value = f"=A{row}"
        ws.cell(row=row, column=15).value = \
            f"=VLOOKUP(O{row}, $A$4:$F$10003, 6, 0)"
    wb.save(path)
    return path


def main() -> int:
    os.makedirs(DATA_DIR, exist_ok=True)
    docx_path = os.path.join(DATA_DIR, "source_demo.docx")
    tpl_path = os.path.join(DATA_DIR, "template_demo.xlsx")
    make_docx(docx_path)
    make_template(tpl_path)
    print(f"Создано: {docx_path}")
    print(f"Создано: {tpl_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())