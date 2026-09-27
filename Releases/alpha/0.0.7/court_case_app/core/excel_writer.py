# -*- coding: utf-8 -*-
"""
Формирование Excel-файла результата (.xlsx).

Аналог листа «Результат обработки» и FinalizeOutput() из Module1.bas.
Набор выводимых столбцов настраивается пользователем.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Столбцы результата: ключ записи -> заголовок
DEFAULT_COLUMNS = [
    ("sequential", "№ п/п"),
    ("title", "Заголовок дела"),
    ("dates", "Даты дела"),
    ("opis", "№ описи"),
    ("unit", "№ ед.хр."),
    ("count", "Кол-во ед.хр."),
    ("retention", "Срок хранения"),
    ("note", "Примечание"),
]

DEFAULT_KEYS = [key for key, _ in DEFAULT_COLUMNS]

# Ограничения ширины столбцов (аналог ColumnWidth=60 для заголовка)
COLUMN_WIDTH_LIMITS = {
    "sequential": 10,
    "title": 60,
    "dates": 22,
    "opis": 10,
    "unit": 12,
    "count": 12,
    "retention": 18,
    "note": 25,
}

HEADER_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
THIN_SIDE = Side(style="thin")
BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)


def write_excel_result(records, filepath: str, columns=None) -> str:
    """
    Создание .xlsx файла результата.

    Параметры:
        records  — список словарей записей (из case_processor);
        filepath — путь к создаваемому файлу;
        columns  — список ключей выводимых столбцов (по умолчанию
                   DEFAULT_KEYS — все 8, как в VBA).

    Возвращает filepath.
    """
    keys = columns if columns is not None else DEFAULT_KEYS
    header_titles = {key: title for key, title in DEFAULT_COLUMNS}

    wb = Workbook()
    ws = wb.active
    ws.title = "Результат обработки"

    # Заголовки
    for col_idx, key in enumerate(keys, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_titles.get(key, key))
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    # Данные
    # fix_10_v2: sanitize_text — нормализуем _x000D_/\r/\v в \n
    from .text_utils import sanitize_text
    for row_idx, record in enumerate(records, start=2):
        for col_idx, key in enumerate(keys, start=1):
            value = record.get(key, "")
            if isinstance(value, str):
                value = sanitize_text(value)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top")

    # Ширина столбцов и перенос текста
    for col_idx, key in enumerate(keys, start=1):
        letter = get_column_letter(col_idx)
        limit = COLUMN_WIDTH_LIMITS.get(key, 20)
        max_len = len(str(header_titles.get(key, key)))
        for record in records[:200]:  # оценка по первым 200 записям
            val = record.get(key, "")
            if val is None:
                continue
            line_len = max((len(line) for line in str(val).split("\n")), default=0)
            if line_len > max_len:
                max_len = line_len
        width = min(max(max_len + 2, len(str(header_titles.get(key, key))) + 2),
                    limit + 2 if key != "title" else limit)
        ws.column_dimensions[letter].width = width
        if key == "title":
            for row_idx in range(2, len(records) + 2):
                ws.cell(row=row_idx, column=col_idx).alignment = (
                    Alignment(vertical="top", wrap_text=True))

    # stage11b: невидимая метка «наш файл» (A2)
    from .converter.markers import mark_excel_workbook
    mark_excel_workbook(wb)

    wb.save(filepath)
    return filepath