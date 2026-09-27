# -*- coding: utf-8 -*-
"""
Тесты обработки и экспорта из БД (stage16d4b, D4b).

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tests\\test_db_pipeline_full.py
"""

import os
import sys
import tempfile
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import Workbook
from docx import Document

from core.db_pipeline import (
    export_year_result_to_file,
    process_year,
    save_year_source_from_xlsx,
)
from database import db

FAILED = 0
PASSED = 0

TEMPLATE = os.path.join(os.path.dirname(__file__), "..", "docs",
                        "АКТ уничтожения гражданских дел.docx")


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def year_data(s="1"):
    return {"судья": f"Судья{s}", "секретарь": f"Сек{s}",
            "дата_утверждения": "—", "дата_акта": "—",
            "номер_акта": s, "дата_подписи": "—",
            "протокол_эк_дата": "—", "протокол_эк_номер": s}


def make_source_xlsx(path):
    """xlsx с 3 делами за 2020 + 1 алиментное."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Дата", "№ дела", "Заявители", "Ответчики", "Категория",
               "№ описи", "№ ед.хр.", "Примечание", "Дата окончания"])
    ws.append([datetime(2020, 1, 10), "2-1/2020",
               "Тестов Тест Тестович", "Тестова Теста Тестовна",
               "Исковое заявление о расторжении брака",
               1, 1, "", datetime(2020, 2, 10)])
    ws.append([datetime(2020, 1, 11), "2-2/2020",
               "Иванова Ирина Ивановна", "Иванов Иван Иванович",
               "Исковое заявление о защите прав потребителей",
               1, 2, "", datetime(2020, 2, 11)])
    ws.append([datetime(2020, 1, 12), "2-3/2020",
               "Петров Пётр Петрович", "Петрова Полина Петровна",
               "Исковое заявление о взыскании долга",
               1, 3, "архив", datetime(2020, 2, 12)])
    ws.append([datetime(2020, 1, 13), "2-4/2020",
               "Алиментова А.А.", "Алиментов А.А.",
               "О взыскании алиментов на содержание ребенка",
               1, 4, "", datetime(2020, 2, 13)])
    wb.save(path)


def docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def main():
    tmp = tempfile.mkdtemp(prefix="db_pipe_b_")
    db_path = os.path.join(tmp, "app.db")
    out_dir = os.path.join(tmp, "out")
    os.makedirs(out_dir, exist_ok=True)

    conn = db.init_db(db_path)
    aid = db.add_court_area(conn, {"номер": "9"})
    yid = db.add_court_year(conn, aid, 2020, year_data("20"))

    xlsx = os.path.join(tmp, "источник.xlsx")
    make_source_xlsx(xlsx)
    save_year_source_from_xlsx(conn, yid, xlsx, "источник.xlsx",
                                sheet_name="Sheet", header_row=1)

    print("[1] process_year: без алиментов")
    result = process_year(conn, yid, case_type="civil",
                          process_alimony=False, court_area_id=aid)
    check("result_id > 0", result["result_id"] > 0, True)
    check("source_file_id > 0", result["source_file_id"] > 0, True)
    check("record_count = 3 (без алиментов, номер 2-4 исключен)",
          result["record_count"], 3)
    check("target_year = 2020", result["target_year"], 2020)
    check("stats есть", "case_prefix" in result["stats"], True)

    print("\n[2] Результат сохранён в БД")
    res = db.get_processing_result(conn, result["result_id"])
    check("case_type = civil", res["case_type"], "civil")
    check("record_count в шапке = 3", res["record_count"], 3)
    check("target_year = 2020", res["target_year"], 2020)
    check("court_area_id сохранён", res["court_area_id"], aid)
    check("output_format = word", res["output_format"], "word")

    print("\n[3] Строки результата (сквозная нумерация с пустыми)")
    rows = db.list_result_rows(conn, result["result_id"])
    # В БД сохраняются все строки, включая пустую для исключённого
    # алиментного дела 2-4 — это часть сквозной нумерации (см. VBA).
    check("4 строки (3 дела + пустая)", len(rows), 4)
    check("case_number в row[0]", rows[0]["case_number"], "2-1/2020")
    check("sequential row[0]", rows[0]["sequential"], "1")
    check("3-я строка — 2-3", rows[2]["case_number"], "2-3/2020")
    check("4-я строка пустая (алименты исключены)",
          rows[3]["title"], "")

    print("\n[4] process_year: с алиментами")
    result2 = process_year(conn, yid, case_type="civil",
                           process_alimony=True, court_area_id=aid)
    check("record_count = 4 (включая алиментное)",
          result2["record_count"], 4)
    check("Старый результат заменён",
          db.get_processing_result(conn, result["result_id"]), None)

    print("\n[5] Повторный process_year — перезапись")
    result3 = process_year(conn, yid, case_type="civil",
                           process_alimony=False, court_area_id=aid)
    check("Новый id", result3["result_id"] != result2["result_id"], True)
    check("Всего результатов года = 1",
          len(db.list_processing_results(conn, yid)), 1)

    print("\n[6] Экспорт в Word")
    if not os.path.exists(TEMPLATE):
        print("  SKIP шаблон не найден:", TEMPLATE)
    else:
        exp = export_year_result_to_file(
            conn, result3["result_id"],
            template_path=TEMPLATE, output_dir=out_dir,
            output_format="word")
        check("Файл создан", os.path.exists(exp["path"]), True)
        check("filename содержит '9'", "9" in exp["filename"], True)
        check("filename .docx", exp["filename"].endswith(".docx"), True)
        text = docx_text(exp["path"])
        check("Судья в акте", "Судья20" in text, True)
        check("Секретарь в акте", "Сек20" in text, True)
        check("Заголовок 2-1 в акте", "2-1/2020" in text, True)
        check("Заголовок 2-3 в акте", "2-3/2020" in text, True)

        from core.converter.markers import read_word_marker
        check("Метка Word есть",
              read_word_marker(exp["path"])["found"], True)

    print("\n[7] Экспорт в Excel")
    exp2 = export_year_result_to_file(
        conn, result3["result_id"],
        template_path=TEMPLATE, output_dir=out_dir,
        output_format="excel")
    check("Файл создан", os.path.exists(exp2["path"]), True)
    check("filename .xlsx", exp2["filename"].endswith(".xlsx"), True)
    from openpyxl import load_workbook
    wb = load_workbook(exp2["path"])
    ws = wb.active
    check("Лист 'Результат обработки'",
          ws.title, "Результат обработки")
    check("Заголовки 8 столбцов",
          ws.cell(row=1, column=1).value, "№ п/п")
    check("2-1 в Excel",
          "2-1/2020" in str(ws.cell(row=2, column=2).value or ""), True)
    wb.close()

    print("\n[8] Валидация")
    try:
        process_year(conn, yid, source_file_id=99999)
        check("Несуществующий source_file_id -> ValueError",
              "нет ошибки", "ValueError")
    except ValueError:
        check("Несуществующий source_file_id -> ValueError", True, True)

    try:
        process_year(conn, 99999)
        check("Несуществующий год -> ValueError", "нет ошибки", "ValueError")
    except ValueError:
        check("Несуществующий год -> ValueError", True, True)

    try:
        export_year_result_to_file(
            conn, 99999, template_path=TEMPLATE,
            output_dir=out_dir, output_format="word")
        check("Несуществующий result_id -> ValueError",
              "нет ошибки", "ValueError")
    except ValueError:
        check("Несуществующий result_id -> ValueError", True, True)

    try:
        export_year_result_to_file(
            conn, result3["result_id"], template_path=TEMPLATE,
            output_dir=out_dir, output_format="pdf")
        check("Неверный output_format -> ValueError",
              "нет ошибки", "ValueError")
    except ValueError:
        check("Неверный output_format -> ValueError", True, True)

    print("\n[9] is_stale при неактуальной версии файла")
    # Загружаем v2 файла — v1 больше не current
    xlsx2 = os.path.join(tmp, "источник_v2.xlsx")
    make_source_xlsx(xlsx2)
    save_year_source_from_xlsx(conn, yid, xlsx2, "источник_v2.xlsx",
                                sheet_name="Sheet", header_row=1)
    # Обрабатываем с конкретным v1 (не актуальным)
    v1_sf = [f for f in db.list_source_files(conn, yid, file_kind="source")
             if f["version"] == 1][0]
    r_stale = process_year(conn, yid, source_file_id=v1_sf["id"],
                            case_type="civil", court_area_id=aid)
    check("Результат помечен is_stale",
          db.get_processing_result(conn, r_stale["result_id"])["is_stale"],
          True)

    conn.close()
    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()
