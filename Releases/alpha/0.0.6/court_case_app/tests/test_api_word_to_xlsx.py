# -*- coding: utf-8 -*-
"""
Тесты API Word -> Excel (stage13b, под-этап C2).

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tests\\test_api_word_to_xlsx.py
"""

import io
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document
from openpyxl import load_workbook

from core.converter.markers import mark_word_document
import app as app_module

FAILED = 0
PASSED = 0


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


HEADERS = ["№ п/п", "Заголовок дела", "Даты дела", "№ описи",
           "№ ед.хр.", "Кол-во ед.хр.", "Срок хранения", "Примечание"]


def make_act(path, cases, num="9", year_hint="2020", marker=True):
    """Создаёт Word-акт с реквизитами и таблицей дел."""
    doc = Document()
    doc.add_paragraph("Акт о выделении к уничтожению документов")
    doc.add_paragraph(f"Судебный участок № {num}")
    doc.add_paragraph("Мировой судья: О. А. Левошина")
    doc.add_paragraph("Секретарь: А.М. Намаюшка")
    t = doc.add_table(rows=1, cols=len(HEADERS))
    for i, h in enumerate(HEADERS):
        t.cell(0, i).text = h
    for c in cases:
        row = t.add_row()
        for i, val in enumerate(c):
            row.cells[i].text = str(val)
    if marker:
        mark_word_document(doc)
    doc.save(path)


def main():
    tmp = tempfile.mkdtemp(prefix="api_w2x_")
    db_path = os.path.join(tmp, "app.db")
    out_dir = os.path.join(tmp, "out")
    up_dir = os.path.join(tmp, "up")
    os.makedirs(out_dir, exist_ok=True)
    os.makedirs(up_dir, exist_ok=True)

    app_module.DB_PATH = db_path
    app_module.OUTPUT_DIR = out_dir
    app_module.UPLOAD_DIR = up_dir

    client = app_module.app.test_client()

    docx_path = os.path.join(tmp, "акт.docx")
    make_act(docx_path, [
        (1, "Гражданское дело №2-1/2020 по заявлению Тестова Т. Т. "
            "к Тестову Т. Т. о расторжении брака",
         "01.01.2020\n01.02.2020", 1, 1, 1, "3 года ЭК Ст. 137", ""),
        (2, "Гражданское дело №2-2/2020 о защите прав потребителей",
         "02.01.2020\n02.02.2020", 1, 2, 1, "3 года ЭПК Ст. 208", ""),
        (3, "Гражданское дело №2-4/2020 о взыскании долга",
         "04.01.2020\n04.02.2020", 1, 4, 1, "3 года Ст. 222", "архив"),
    ])

    print("[1] Базовый POST -> JSON")
    with open(docx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "акт.docx")}
        r = client.post("/convert/word_to_xlsx", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("filename есть", bool(j.get("filename")), True)
    check("filename с '9СУ' и '2020'",
          "9СУ" in j["filename"] and "2020" in j["filename"], True)
    check("count = 3", j.get("count"), 3)
    check("total = 3", j.get("total"), 3)
    check("attention — список", isinstance(j.get("attention"), list), True)
    check("refs содержит 'судья'", "судья" in (j.get("refs") or {}), True)
    check("marker.found=True", j["marker"]["found"], True)
    check("preview — 3", len(j.get("preview") or []), 3)
    check("year = 2020", j.get("year"), "2020")

    print("\n[2] Файл создан, лист «Результат обработки»")
    out_file = os.path.join(out_dir, j["filename"])
    check("Файл существует", os.path.exists(out_file), True)
    wb = load_workbook(out_file)
    check("Лист 'Результат обработки'",
          "Результат обработки" in wb.sheetnames, True)
    ws = wb.active
    check("Шапка — 8 столбцов",
          [ws.cell(row=1, column=i+1).value for i in range(8)],
          HEADERS)
    check("Заголовок 2-1 на месте",
          "2-1/2020" in (ws.cell(row=2, column=2).value or ""), True)
    wb.close()

    print("\n[3] restore_gaps=1 восстанавливает пропуск 2-3")
    with open(docx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "акт.docx"),
                "restore_gaps": "1", "copy": "1"}
        r = client.post("/convert/word_to_xlsx", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j2 = r.get_json()
    check("total = 4 (с пустой строкой)", j2.get("total"), 4)
    check("count = 3 (только дела)", j2.get("count"), 3)
    wb = load_workbook(os.path.join(out_dir, j2["filename"]))
    ws = wb.active
    # Excel-строки: 1 — шапка, 2 — 2-1, 3 — 2-2, 4 — пустая, 5 — 2-4
    check("Строка 4 пустая (пропуск 2-3)",
          (ws.cell(row=4, column=2).value or ""), "")
    check("Строка 5 — 2-4",
          "2-4/2020" in (ws.cell(row=5, column=2).value or ""), True)
    wb.close()

    print("\n[4] Повторный POST без overwrite -> 409")
    with open(docx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "акт.docx")}
        r = client.post("/convert/word_to_xlsx", data=data,
                        content_type="multipart/form-data")
    check("HTTP 409", r.status_code, 409)
    check("conflict=true", r.get_json().get("conflict"), True)

    print("\n[5] overwrite=1 -> 200")
    with open(docx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "акт.docx"),
                "overwrite": "1"}
        r = client.post("/convert/word_to_xlsx", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)

    print("\n[6] download=1 -> send_file")
    with open(docx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "акт.docx"),
                "copy": "1", "download": "1"}
        r = client.post("/convert/word_to_xlsx", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    cd = r.headers.get("Content-Disposition", "")
    check("Content-Disposition: attachment", "attachment" in cd, True)
    check("Магические байты xlsx (PK)", r.data[:2] == b"PK", True)

    print("\n[7] Неверное расширение -> 400")
    with open(docx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "акт.xlsx")}
        r = client.post("/convert/word_to_xlsx", data=data,
                        content_type="multipart/form-data")
    check("HTTP 400", r.status_code, 400)
    check("Ошибка про .docx", ".docx" in (r.get_json() or {}).get("error", ""), True)

    print("\n[8] Файл без дел -> 400")
    empty_docx = os.path.join(tmp, "empty.docx")
    d = Document()
    d.add_paragraph("Просто текст")
    d.save(empty_docx)
    with open(empty_docx, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "empty.docx")}
        r = client.post("/convert/word_to_xlsx", data=data,
                        content_type="multipart/form-data")
    check("HTTP 400", r.status_code, 400)

    print("\n[9] Метка в созданном Excel")
    # Найдём последний (не-overwrite) файл в output/
    files = sorted(os.listdir(out_dir))
    latest = os.path.join(out_dir, files[-1])
    from core.converter.markers import read_excel_marker
    check("Метка Excel", read_excel_marker(latest)["found"], True)

    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()
