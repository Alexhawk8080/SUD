# -*- coding: utf-8 -*-
"""
Тесты невидимых меток Word/Excel (stage11b, под-этап A2).

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tests\\test_markers.py
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document
from openpyxl import Workbook, load_workbook

from core.converter.markers import (
    CURRENT_MARKER_VERSION,
    EXCEL_MARKER,
    WORD_MARKER,
    make_marker,
    mark_excel_workbook,
    mark_word_document,
    parse_marker,
    read_excel_marker,
    read_word_marker,
)
from core.excel_writer import write_excel_result
from core.word_writer import write_word_result

FAILED = 0
PASSED = 0

TEMPLATE = os.path.join(os.path.dirname(__file__), "..", "docs",
                        "АКТ уничтожения гражданских дел.docx")


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def make_records():
    return [
        {"sequential": 1,
         "title": "Гражданское дело №2-3/2019 по заявлению Тестова Т. Т. "
                  "к Тестову Т. Т. о расторжении брака",
         "dates": "26.11.2018\n26.12.2018",
         "opis": 1, "unit": 1, "count": 1,
         "retention": "3 года ЭК Ст. 137", "note": ""},
    ]


def make_common_values():
    return {
        "судебный_участок": "9",
        "судья": "О. А. Левошина",
        "дата_утверждения": "«___» ______________ 20__ года",
        "дата_акта": "«__» ____________ 20__ г.",
        "номер_акта": "1",
        "секретарь": "А.М. Намаюшка",
        "дата_подписи": "«__» ____________ 20__ г.",
        "протокол_эк_дата": "«__» ____________ 20__ г.",
        "протокол_эк_номер": "1",
        "год_дел": "2019",
    }


def main():
    tmp = tempfile.mkdtemp(prefix="markers_")

    print("[1] make_marker / parse_marker")
    check("Текущая версия", CURRENT_MARKER_VERSION, "V1")
    check("WORD_MARKER содержит маркер", "ACT-V1-COURT-CASE-APP" in WORD_MARKER, True)
    check("WORD_MARKER начинается с ZW", WORD_MARKER.startswith("\u200B"), True)
    check("WORD_MARKER заканчивается ZW", WORD_MARKER.endswith("\u200B"), True)
    check("EXCEL_MARKER == WORD_MARKER", EXCEL_MARKER, WORD_MARKER)
    check("parse_marker извлекает V1", parse_marker(WORD_MARKER), "V1")
    check("parse_marker V2", parse_marker(make_marker("V2")), "V2")
    check("parse_marker терпит 'v3'", parse_marker(make_marker("v3")), "V3")
    check("parse_marker на пустом -> ''", parse_marker(""), "")
    check("parse_marker на чужой строке -> ''", parse_marker("просто текст"), "")

    print("\n[2] Word: метка ставится и читается")
    doc = Document()
    doc.add_paragraph("тест")
    mark_word_document(doc)
    p1 = os.path.join(tmp, "marked.docx")
    doc.save(p1)
    res = read_word_marker(p1)
    check("Метка найдена", res["found"], True)
    check("Версия V1", res["version"], "V1")

    print("\n[3] Word: без метки -> не найдена")
    plain_doc = Document()
    plain_doc.add_paragraph("без метки")
    p2 = os.path.join(tmp, "plain.docx")
    plain_doc.save(p2)
    res2 = read_word_marker(p2)
    check("Метки нет", res2["found"], False)
    check("Версия пустая", res2["version"], "")

    print("\n[4] Excel: метка ставится и читается")
    wb = Workbook()
    ws = wb.active
    ws.append(["тест"])
    mark_excel_workbook(wb)
    x1 = os.path.join(tmp, "marked.xlsx")
    wb.save(x1)
    resx = read_excel_marker(x1)
    check("Метка найдена", resx["found"], True)
    check("Версия V1", resx["version"], "V1")

    print("\n[5] Excel: без метки -> не найдена")
    wb2 = Workbook()
    wb2.active.append(["без метки"])
    x2 = os.path.join(tmp, "plain.xlsx")
    wb2.save(x2)
    resx2 = read_excel_marker(x2)
    check("Метки нет", resx2["found"], False)

    print("\n[6] write_word_result ставит метку автоматически")
    out_docx = os.path.join(tmp, "акт.docx")
    if os.path.exists(TEMPLATE):
        write_word_result(make_records(), TEMPLATE, out_docx, make_common_values())
        resw = read_word_marker(out_docx)
        check("Метка в акте", resw["found"], True)
        check("Версия V1", resw["version"], "V1")
    else:
        print("  SKIP шаблон не найден:", TEMPLATE)

    print("\n[7] write_excel_result ставит метку автоматически")
    out_xlsx = os.path.join(tmp, "результат.xlsx")
    write_excel_result(make_records(), out_xlsx)
    reswx = read_excel_marker(out_xlsx)
    check("Метка в результате", reswx["found"], True)
    check("Версия V1", reswx["version"], "V1")

    print("\n[8] Метка не видна в тексте документа")
    d = Document(out_docx) if os.path.exists(out_docx) else Document(p1)
    text = "\n".join(p.text for p in d.paragraphs)
    check("Нет 'ACT-V1' в тексте абзацев", "ACT-V1" in text, False)
    check("Нет U+200B в тексте абзацев", "\u200B" in text, False)

    print("\n[9] Обновление версии при повторной записи")
    doc_v2 = Document(p1)  # был V1
    cp = doc_v2.core_properties
    cp.keywords = make_marker("V1")
    cp.comments = make_marker("V1")
    # перезаписываем V2
    from core.converter.markers import make_marker as _mk
    cp.keywords = _mk("V2")
    cp.comments = _mk("V2")
    p3 = os.path.join(tmp, "v2.docx")
    doc_v2.save(p3)
    res3 = read_word_marker(p3)
    check("Версия обновилась до V2", res3["version"], "V2")

    print("\n[10] Битый / несуществующий файл")
    check("Несуществующий .docx", read_word_marker(os.path.join(tmp, "no.docx"))["found"], False)
    check("Несуществующий .xlsx", read_excel_marker(os.path.join(tmp, "no.xlsx"))["found"], False)

    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()
