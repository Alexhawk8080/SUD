# -*- coding: utf-8 -*-
"""
Регрессия модуля «Внутренняя опись» на реальных фикстурах (этап 14f).

Сквозной сценарий: source_demo.docx -> поток -> запись в template_demo.xlsx.

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tests\\test_inventory_regression.py
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import load_workbook

from core.inventory.docx_reader import read_inventory_docx
from core.inventory.flow import build_flow
from core.inventory.normalizer import CASE_CIVIL
from core.inventory.writer import prepare_rows, write_result

FAILED = 0
PASSED = 0

DATA = os.path.join(os.path.dirname(__file__), "data", "inventory")
DOCX = os.path.join(DATA, "source_demo.docx")
TEMPLATE = os.path.join(DATA, "template_demo.xlsx")


def check(name, actual, expected):
    global FAILED, PASSED
    if actual == expected:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n"
              f"       получено:  {actual!r}")


def main():
    check("фикстура docx", os.path.exists(DOCX), True)
    check("фикстура xlsx", os.path.exists(TEMPLATE), True)

    print("[1] Чтение docx-описи")
    doc = read_inventory_docx(DOCX)
    check("строк", len(doc["rows"]), 4)
    check("томов", doc["tables"], 1)
    check("первые номера",
          [r["number"] for r in doc["rows"]], [1, 2, 3, 4])

    print("\n[2] Поток и подготовка строк")
    flow = build_flow([doc])
    check("тип", flow["type"], CASE_CIVIL)
    check("год", flow["year"], 2018)
    check("год-плана", flow["stream"]["year_plan"]["year"], 2018)
    prepared = prepare_rows(flow)
    check("строк подготовлено", len(prepared["rows"]), 4)
    check("f4", prepared["f4"], 14)
    check("номера дел", [r["F"] for r in prepared["rows"]],
          [14, 15, 16, 16])
    check("приращения", [r["E"] for r in prepared["rows"]],
          [None, 1, 1, 0])
    check("C первой строки", prepared["rows"][0]["C"],
          "Решение по гражданскому делу")
    check("C заочной", prepared["rows"][1]["C"],
          "Заочное решение по гражданскому делу")
    check("страницы первой (1-5)", prepared["rows"][0]["H"], 5)
    check("страницы второй (6)", prepared["rows"][1]["H"], 1)

    print("\n[3] Запись в шаблон «Таблица»")
    tmp = tempfile.mkdtemp(prefix="inv_reg_")
    out = os.path.join(tmp, "result.xlsx")
    res = write_result(TEMPLATE, flow, out)
    check("файл создан", os.path.exists(out), True)
    check("результат = подготовка", len(res["rows"]), 4)
    wb = load_workbook(out)
    ws = wb["Таблица"]
    check("C4", ws.cell(row=4, column=3).value,
          "Решение по гражданскому делу")
    check("D4 — префикс", ws.cell(row=4, column=4).value, "2")
    check("E4 пусто (не заполняется)", ws.cell(row=4, column=5).value, None)
    check("E5 = 1", ws.cell(row=5, column=5).value, 1)
    check("F4 из Word", ws.cell(row=4, column=6).value, 14)
    check("H4", ws.cell(row=4, column=8).value, 5)
    check("H7 = 1 (одна страница)", ws.cell(row=7, column=8).value, 1)
    check("лист внимания",
          "Требует внимания" in wb.sheetnames, True)

    print("\n" + "=" * 60)
    print(f"ИТОГО: PASSED = {PASSED}, FAILED = {FAILED}")
    return 1 if FAILED else 0


if __name__ == "__main__":
    sys.exit(main())
