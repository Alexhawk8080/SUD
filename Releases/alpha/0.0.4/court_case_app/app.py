# -*- coding: utf-8 -*-
"""
Веб-приложение «Акт уничтожения гражданских дел».

Flask-сервер: главная страница (обработка файлов), страница базы данных,
API для управления БД, скачивание результата, предпросмотр исходного файла.
Автозапуск браузера; остановка сервера при закрытии вкладки интерфейса
(heartbeat + сигнал при закрытии).
"""

import json
import os
import re
import threading
import time
import webbrowser

from flask import Flask, abort, jsonify, render_template, request, send_file

from core.pipeline import process_rows, process_upload
from core.excel_reader import (
    ROLES,
    ROLE_TITLES,
    auto_detect_roles,
    detect_header_row,
    read_source_table,
)
from core.excel_writer import DEFAULT_COLUMNS
from core.category_check import analyze
from core.legacy_reader import parse_legacy_file
from core.case_processor import contains_alimony, is_valid_case_number
from core.converter import (
    ConverterError,
    detect_year_from_numbers,
    inspect_sheets,
    make_warning_message,
    read_act_docx,
    read_result_xlsx,
)
from core.converter.xlsx_to_word import convert_xlsx_to_word
from core.inventory.api import register_routes as register_inventory_routes
from database import db

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "database", "app.db")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
TEMPLATE_DOCX = os.path.join(
    BASE_DIR, "docs", "АКТ уничтожения гражданских дел.docx")

# Heartbeat: если вкладка интерфейса закрыта и heartbeat пропал — остановить сервер.
# Web Worker шлёт heartbeat каждые 4 с; таймаут 20 с = запас на 4 пропуска
# (переходы между страницами интерфейса, кратковременные паузы браузера).
HEARTBEAT_TIMEOUT = 20.0  # секунд без heartbeat до остановки
_last_beat = time.time()


def _beat_watchdog():
    """Фоновый поток: останавливает сервер, если heartbeat пропал."""
    while True:
        time.sleep(2.0)
        if time.time() - _last_beat > HEARTBEAT_TIMEOUT:
            os._exit(0)


def start_beat_watchdog():
    threading.Thread(target=_beat_watchdog, daemon=True).start()


app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024  # 25 МБ

# Модуль «Внутренняя опись» (Word -> лист «Таблица»), этап 14e.
register_inventory_routes(app, lambda: (OUTPUT_DIR, UPLOAD_DIR))


def get_db():
    """Соединение с БД (создание/наполнение при первом обращении)."""
    return db.init_db(DB_PATH)


def _json_error(message, code=400):
    return jsonify({"error": message}), code


# ---------------------------------------------------------------------------
# Страницы
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    return render_template(
        "index.html",
        roles=ROLES,
        role_titles=ROLE_TITLES,
        columns=DEFAULT_COLUMNS,
    )


@app.route("/db_page")
def db_page():
    return render_template("db.html")


@app.route("/check_page")
def check_page():
    """Страница «Проверка категорий» (Доработка 6)."""
    return render_template("check.html")


@app.route("/heartbeat", methods=["GET"])
def heartbeat():
    """Сигнал «вкладка жива» от клиента."""
    global _last_beat
    _last_beat = time.time()
    return jsonify({"ok": True})


@app.route("/shutdown", methods=["POST"])
def shutdown():
    """Остановка сервера по сигналу клиента (закрытие вкладки)."""
    os._exit(0)


@app.route("/preview", methods=["POST"])
def preview():
    """
    Предпросмотр загруженного исходного файла.

    Возвращает первые строки листа, автоопределённые номер строки
    заголовков и маппинг ролей.
    """
    if "file" not in request.files:
        return _json_error("Файл не загружен")
    uploaded = request.files["file"]
    if not uploaded.filename:
        return _json_error("Файл не выбран")

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    path = os.path.join(UPLOAD_DIR, os.path.basename(uploaded.filename))
    uploaded.save(path)

    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        rows = []
        for values in ws.iter_rows(values_only=True):
            cells = []
            for v in values[:14]:
                if v is None:
                    cells.append("")
                else:
                    s = str(v)
                    cells.append(s[:80] + ("…" if len(s) > 80 else ""))
            rows.append(cells)
            if len(rows) >= 12:
                break
        wb.close()
    except Exception as exc:  # noqa: BLE001
        return _json_error(f"Не удалось прочитать файл: {exc}", 400)

    header_row = detect_header_row(rows)
    mapping = {}
    if header_row > 0 and header_row <= len(rows):
        mapping = auto_detect_roles(rows[header_row - 1])
    return jsonify({
        "rows": rows,
        "header_row": header_row,
        "mapping": mapping,
    })


# ---------------------------------------------------------------------------
# Проверка категорий дел (Доработка 6)
# ---------------------------------------------------------------------------

def _filter_rows_for_year(rows, year, alimony: bool) -> list:
    """
    Оставляет дела, попадающие в результат обработки:
    валидный номер за выбранный год + фильтр алиментных дел.
    """
    year = str(year)
    result = []
    for row in rows:
        if not is_valid_case_number(row.get("case_number"), year):
            continue
        if alimony or not contains_alimony(row.get("category")):
            result.append(row)
    return result


@app.route("/check_categories", methods=["POST"])
def check_categories():
    """
    Анализ загруженного файла: поиск дел с отсутствующей категорией,
    аналогов по истцу, групп «Без аналогов».

    Параметры формы:
        file       — файл (исходная таблица .xlsx/.xlsm или готовый результат);
        file_type  — 'source' (исходная таблица) или 'legacy' (готовый результат);
        year       — год дел;
        alimony    — '1' — включать алиментные дела;
        header_row — номер строки заголовков (для исходной таблицы).
    """
    if "file" not in request.files:
        return _json_error("Файл не загружен")
    uploaded = request.files["file"]
    if not uploaded.filename:
        return _json_error("Файл не выбран")

    year = request.form.get("year", "").strip()
    if not year.isdigit():
        return _json_error("Укажите год (например, 2023)")
    alimony = request.form.get("alimony") == "1"
    file_type = request.form.get("file_type", "source")

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    path = os.path.join(UPLOAD_DIR, os.path.basename(uploaded.filename))
    uploaded.save(path)

    try:
        if file_type == "legacy":
            # Готовый результат: записи извлекаются из заголовков
            raw_rows = parse_legacy_file(path)
        else:
            header_raw = request.form.get("header_row", "").strip()
            if header_raw == "":
                header_row = None
            else:
                try:
                    header_row = int(header_raw)
                except ValueError:
                    header_row = None
            _headers, raw_rows, _mapping = read_source_table(
                path, mapping=None, header_row=header_row)
    except Exception as exc:  # noqa: BLE001
        return _json_error(f"Не удалось прочитать файл: {exc}", 400)

    rows = _filter_rows_for_year(raw_rows, year, alimony)

    conn = get_db()
    try:
        org_names = db.organization_names(conn)
        categories_map = db.case_categories_map(conn)
        auto_fix_enabled = db.get_setting(conn, "auto_fix_categories", "1") == "1"
    finally:
        conn.close()

    result = analyze(rows, org_names, categories_map)
    result["cases_total"] = len(rows)
    result["auto_fix_enabled"] = auto_fix_enabled
    return jsonify(result)


@app.route("/apply_corrections", methods=["POST"])
def apply_corrections():
    """
    Сохранение исправлений в словарь «истец -> категория» (case_categories).

    Тело JSON: {"saves": [{"key": ..., "label": ..., "category": ...,
                           "is_organization": bool}, ...]}
    """
    data = request.get_json(force=True)
    saves = data.get("saves", []) or []
    conn = get_db()
    saved = 0
    try:
        for item in saves:
            key = str(item.get("key", "")).strip()
            label = str(item.get("label", "")).strip()
            category = str(item.get("category", "")).strip()
            if not key or not category:
                continue
            db.add_case_category(
                conn, key, label or key, category,
                bool(item.get("is_organization", False)))
            saved += 1
        return jsonify({"ok": True, "saved": saved})
    finally:
        conn.close()


@app.route("/download_no_analog_table", methods=["POST"])
def download_no_analog_table():
    """
    Выгрузка таблицы «Без аналогов» в Excel (.xlsx) для ручной работы.

    Тело JSON: {"groups": [...], "year": "2023"}
    """
    data = request.get_json(force=True)
    groups = data.get("groups", []) or []
    year = data.get("year", "")

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Без аналогов"
    ws.append(["Истец", "№ дела", "Категория (заполнить)"])
    for group in groups:
        label = str(group.get("label", ""))
        for case in group.get("cases", []):
            ws.append([label, str(case.get("case_number", "")), ""])
        # Подсказка из словаря БД
        opts = group.get("dictionary_options") or []
        if opts:
            ws.append(["", "", f"Варианты из словаря: {'; '.join(opts)}"])

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filename = f"без_аналогов_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = os.path.join(OUTPUT_DIR, filename)
    wb.save(path)
    return send_file(path, as_attachment=True, download_name=filename)


# ---------------------------------------------------------------------------
# Конвертация (stage12c, под-этап B3)
# ---------------------------------------------------------------------------

_CASE_NUM_IN_TITLE_RE = re.compile(r"№\s*(\d+-\d+/\d{4})")


@app.route("/convert_page")
def convert_page():
    """Страница «Конвертация» (заготовка B3, полный UI — B4)."""
    return render_template("convert.html")


def _check_xlsx_extension(filename: str) -> None:
    """Проверка расширения для Excel-ветки."""
    ext = os.path.splitext(filename)[1].lower()
    if ext not in (".xlsx", ".xlsm"):
        raise ConverterError("Ожидается файл .xlsx или .xlsm")


@app.route("/inspect_file", methods=["POST"])
def inspect_file():
    """
    Инспекция файла для страницы конвертации.

    Параметры формы:
        file  — .xlsx/.xlsm;
        sheet — (опционально) имя листа; если не передан — отдаём список
                листов и работаем с листом по умолчанию.

    Возвращает JSON:
        sheets       — список листов (только при первом вызове, без sheet)
        sheet        — имя использованного листа
        mapping      — role -> индекс столбца
        preview      — первые 12 записей
        total        — всего записей
        year         — определённый год (самый частый из номеров дел)
        years        — {год: количество}
        others       — [[год, количество], ...] «чужие» годы
        warning      — человекочитаемое предупреждение (может быть "")
    """
    if "file" not in request.files:
        return _json_error("Файл не загружен")
    uploaded = request.files["file"]
    if not uploaded.filename:
        return _json_error("Файл не выбран")

    try:
        _check_xlsx_extension(uploaded.filename)
    except ConverterError as exc:
        return _json_error(str(exc))

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    path = os.path.join(UPLOAD_DIR, os.path.basename(uploaded.filename))
    uploaded.save(path)

    sheet_param = request.form.get("sheet", "").strip() or None

    try:
        if sheet_param is None:
            info = inspect_sheets(path)
            sheets = info["sheets"]
            sheet = info["default"]
        else:
            sheets = None
            sheet = sheet_param

        records, mapping = read_result_xlsx(path, sheet_name=sheet)
    except ConverterError as exc:
        return _json_error(str(exc))
    except Exception as exc:  # noqa: BLE001
        return _json_error(f"Ошибка чтения файла: {exc}", 500)

    # Год из номеров дел (по заголовкам)
    nums = []
    for r in records:
        m = _CASE_NUM_IN_TITLE_RE.search(r.get("title") or "")
        if m:
            nums.append(m.group(1))

    try:
        yinfo = detect_year_from_numbers(nums)
    except ConverterError as exc:
        return _json_error(str(exc))

    preview = []
    for r in records[:12]:
        preview.append({
            "sequential": r.get("sequential"),
            "title": r.get("title") or "",
            "dates": (r.get("dates") or "").replace("\n", "<br>"),
            "opis": r.get("opis") or "",
            "unit": r.get("unit") or "",
            "count": r.get("count") or "",
            "retention": r.get("retention") or "",
            "note": r.get("note") or "",
        })

    return jsonify({
        "sheets": sheets,
        "sheet": sheet,
        "mapping": mapping,
        "preview": preview,
        "total": len(records),
        "year": yinfo["year"],
        "years": yinfo["years"],
        "others": yinfo["others"],
        "warning": make_warning_message(yinfo),
    })


@app.route("/convert/xlsx_to_word", methods=["POST"])
def convert_xlsx_to_word_route():
    """
    Конвертация Excel-результата в Word-акт.

    Параметры формы:
        file          — .xlsx/.xlsm;
        year          — год дел (обязателен);
        sheet         — имя листа (опционально);
        court_area_id — id участка (опционально);
        download      — '1' — сразу отдать файл;
        overwrite     — '1' — перезаписать существующий;
        copy          — '1' — сохранить как копию «(1).docx».

    Возвращает:
        download=1 -> файл;
        иначе      -> JSON {filename, count, total, preview};
        конфликт   -> 409 {error, filename, conflict: true}.
    """
    if "file" not in request.files:
        return _json_error("Файл не загружен")
    uploaded = request.files["file"]
    if not uploaded.filename:
        return _json_error("Файл не выбран")

    try:
        _check_xlsx_extension(uploaded.filename)
    except ConverterError as exc:
        return _json_error(str(exc))

    year = request.form.get("year", "").strip()
    if not year.isdigit():
        return _json_error("Укажите год (например, 2020)")

    sheet = request.form.get("sheet", "").strip() or None

    court_area_id_raw = request.form.get("court_area_id", "").strip()
    court_area_id = None
    if court_area_id_raw.isdigit():
        v = int(court_area_id_raw)
        if v > 0:
            court_area_id = v

    download = request.form.get("download") == "1"
    overwrite = request.form.get("overwrite") == "1"
    copy_mode = request.form.get("copy") == "1"

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    path = os.path.join(UPLOAD_DIR, os.path.basename(uploaded.filename))
    uploaded.save(path)

    # Реквизиты: приоритет — court_areas, fallback — settings
    conn = get_db()
    try:
        area_values = (db.court_area_settings_dict(conn, court_area_id)
                       if court_area_id else None)
        settings = db.get_settings(conn)
    finally:
        conn.close()

    if area_values:
        num = str(area_values.get("судебный_участок") or "").strip()
        common_values = {k: ("" if v is None else v)
                         for k, v in area_values.items()}
    else:
        num = str(settings.get("судебный_участок") or "").strip()
        common_values = {k: ("" if v is None else v)
                         for k, v in dict(settings).items()}

    base = "Акт уничтожения гражданских дел"
    if num:
        base += f" {num}СУ"
    base += f" {year}"
    filename = f"{base}.docx"
    out_path = os.path.join(OUTPUT_DIR, filename)

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    # Проверка конфликта имени (решения 42/42а/42б)
    if os.path.exists(out_path) and not (overwrite or copy_mode):
        return jsonify({
            "error": "Файл уже существует",
            "filename": filename,
            "conflict": True,
        }), 409

    if copy_mode and os.path.exists(out_path):
        i = 1
        while True:
            cand = f"{base} ({i}).docx"
            cand_path = os.path.join(OUTPUT_DIR, cand)
            if not os.path.exists(cand_path):
                filename = cand
                out_path = cand_path
                break
            i += 1

    try:
        result = convert_xlsx_to_word(
            path, TEMPLATE_DOCX, out_path,
            common_values=common_values, year=year, sheet_name=sheet)
    except ConverterError as exc:
        return _json_error(str(exc))
    except Exception as exc:  # noqa: BLE001
        return _json_error(f"Ошибка конвертации: {exc}", 500)

    if download:
        return send_file(out_path, as_attachment=True,
                         download_name=filename)

    preview = []
    for r in result["records"][:12]:
        preview.append({
            "sequential": r.get("sequential"),
            "title": r.get("title") or "",
            "dates": (r.get("dates") or "").replace("\n", "<br>"),
            "opis": r.get("opis") or "",
            "unit": r.get("unit") or "",
            "count": r.get("count") or "",
            "retention": r.get("retention") or "",
            "note": r.get("note") or "",
        })

    return jsonify({
        "filename": filename,
        "count": result["count"],
        "total": len(result["records"]),
        "preview": preview,
    })


@app.route("/convert/word_to_xlsx", methods=["POST"])
def convert_word_to_xlsx_route():
    """
    Конвертация Word-акта в Excel-результат.

    Параметры формы:
        file          — .docx;
        download      — '1' — сразу отдать файл;
        overwrite     — '1' — перезаписать существующий;
        copy          — '1' — сохранить как копию;
        restore_gaps  — '1' — восстановить пустые строки для пропусков;
        columns       — (опционально, повторяющийся) список ключей столбцов.

    Возвращает:
        download=1 -> файл;
        иначе      -> JSON {filename, count, total, attention, refs,
                            marker, preview};
        конфликт   -> 409 {error, filename, conflict: true}.
    """
    if "file" not in request.files:
        return _json_error("Файл не загружен")
    uploaded = request.files["file"]
    if not uploaded.filename:
        return _json_error("Файл не выбран")

    ext = os.path.splitext(uploaded.filename)[1].lower()
    if ext != ".docx":
        return _json_error("Ожидается файл .docx")

    download = request.form.get("download") == "1"
    overwrite = request.form.get("overwrite") == "1"
    copy_mode = request.form.get("copy") == "1"
    restore_gaps = request.form.get("restore_gaps") == "1"

    columns = request.form.getlist("columns") or None

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    path = os.path.join(UPLOAD_DIR, os.path.basename(uploaded.filename))
    uploaded.save(path)

    # Имя файла: «Акт уничтожения гражданских дел {N}СУ {YYYY}.xlsx»
    # Год берём из номеров дел акта, участок — из реквизитов шапки акта.
    try:
        parsed = read_act_docx(path)
    except ConverterError as exc:
        return _json_error(str(exc))
    except Exception as exc:  # noqa: BLE001
        return _json_error(f"Ошибка чтения акта: {exc}", 500)

    nums = []
    for r in parsed["records"]:
        m = _CASE_NUM_IN_TITLE_RE.search(str(r.get("title") or ""))
        if m:
            nums.append(m.group(1))
    try:
        yinfo = detect_year_from_numbers(nums)
    except ConverterError as exc:
        return _json_error(str(exc))

    # Год: приоритет — номерам дел (решение 6).
    # Участок: из реквизитов шапки акта, если удалось извлечь.
    refs = parsed.get("refs") or {}
    num = (refs.get("номер") or "").strip()

    base = "Акт уничтожения гражданских дел"
    if num:
        base += f" {num}СУ"
    base += f" {yinfo['year']}"
    filename = f"{base}.xlsx"
    out_path = os.path.join(OUTPUT_DIR, filename)

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    if os.path.exists(out_path) and not (overwrite or copy_mode):
        return jsonify({
            "error": "Файл уже существует",
            "filename": filename,
            "conflict": True,
        }), 409

    if copy_mode and os.path.exists(out_path):
        i = 1
        while True:
            cand = f"{base} ({i}).xlsx"
            cand_path = os.path.join(OUTPUT_DIR, cand)
            if not os.path.exists(cand_path):
                filename = cand
                out_path = cand_path
                break
            i += 1

    try:
        from core.converter.word_to_xlsx import convert_word_to_xlsx
        result = convert_word_to_xlsx(
            path, out_path, restore_gaps=restore_gaps, columns=columns)
    except ConverterError as exc:
        return _json_error(str(exc))
    except Exception as exc:  # noqa: BLE001
        return _json_error(f"Ошибка конвертации: {exc}", 500)

    if download:
        return send_file(out_path, as_attachment=True,
                         download_name=filename)

    # Превью первых 12 записей
    preview = []
    for r in result["records"][:12]:
        preview.append({
            "sequential": r.get("sequential"),
            "title": r.get("title") or "",
            "dates": (r.get("dates") or "").replace("\n", "<br>"),
            "opis": r.get("opis") or "",
            "unit": r.get("unit") or "",
            "count": r.get("count") or "",
            "retention": r.get("retention") or "",
            "note": r.get("note") or "",
        })

    return jsonify({
        "filename": filename,
        "count": result["count"],
        "total": len(result["records"]),
        "attention": result["attention"],
        "refs": result["refs"],
        "marker": result["marker"],
        "preview": preview,
        "year": yinfo["year"],
    })


# ---------------------------------------------------------------------------
# Обработка и скачивание
# ---------------------------------------------------------------------------

@app.route("/process", methods=["POST"])
def process():
    if "file" not in request.files:
        return _json_error("Файл не загружен")
    uploaded = request.files["file"]
    if not uploaded.filename:
        return _json_error("Файл не выбран")

    year = request.form.get("year", "").strip()
    if not year.isdigit():
        return _json_error("Укажите год (например, 2023)")

    alimony = request.form.get("alimony") == "1"
    output_format = request.form.get("format", "excel")
    if output_format not in ("excel", "word"):
        output_format = "excel"

    # stage11d: выбор судебного участка (id из court_areas)
    court_area_id_raw = request.form.get("court_area_id", "").strip()
    court_area_id = None
    if court_area_id_raw.isdigit():
        v = int(court_area_id_raw)
        if v > 0:
            court_area_id = v

    # Ручной маппинг ролей (если передан)
    mapping = {}
    for role in ROLES:
        value = request.form.get(f"map_{role}", "").strip()
        if value.isdigit() and int(value) > 0:
            mapping[role] = int(value)
    header_raw = request.form.get("header_row", "").strip()
    if header_raw == "":
        header_row = None  # автоопределение
    else:
        try:
            header_row = int(header_raw)
        except ValueError:
            header_row = None
        if header_row is not None and header_row < 0:
            header_row = 0

    # Выбранные столбцы для Excel (пусто = все)
    columns = request.form.getlist("columns")

    # Доработка 6: исправления категорий со страницы проверки + автоисправление
    fixes_raw = request.form.get("fixes", "").strip()
    category_fixes = {}
    if fixes_raw:
        try:
            raw = json.loads(fixes_raw) or {}
            # fix_02: ключи номеров дел — без хвостовых/неразрывных пробелов
            category_fixes = {
                str(k).strip(): v for k, v in raw.items() if str(k).strip()
            }
        except ValueError:
            category_fixes = {}
    auto_fix_enabled = request.form.get("auto_fix", "1") != "0"
    file_type = request.form.get("file_type", "source")

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    upload_path = os.path.join(UPLOAD_DIR, os.path.basename(uploaded.filename))
    uploaded.save(upload_path)

    try:
        if file_type == "legacy":
            # Готовый результат (Доработка 7): даты дела, № описи,
            # № ед.хр., кол-во и срок хранения переносятся «как есть».
            raw_rows = parse_legacy_file(upload_path)
            if not raw_rows:
                raise ValueError(
                    "В файле не найдено дел (проверьте, что это результат "
                    "обработки или акт уничтожения)")
            result = process_rows(
                raw_rows, int(year), alimony, output_format,
                DB_PATH, TEMPLATE_DOCX, OUTPUT_DIR,
                columns=columns or None,
                category_fixes=category_fixes,
                auto_fix_enabled=auto_fix_enabled,
                court_area_id=court_area_id)
        else:
            result = process_upload(
                upload_path, int(year), alimony, output_format,
                DB_PATH, TEMPLATE_DOCX, OUTPUT_DIR,
                columns=columns or None,
                mapping=mapping or None,
                header_row=header_row,
                category_fixes=category_fixes,
                auto_fix_enabled=auto_fix_enabled,
                court_area_id=court_area_id,
            )
    except ValueError as exc:
        return _json_error(str(exc))
    except Exception as exc:  # noqa: BLE001
        return _json_error(f"Ошибка обработки: {exc}", 500)

    # Доработка 7: прямое скачивание результата (кнопка «Скачать результат»)
    if request.form.get("download") == "1":
        return send_file(result["path"], as_attachment=True,
                         download_name=result["filename"])

    # Подготовка данных для предпросмотра
    preview = []
    for record in result["records"]:
        preview.append({
            "sequential": record.get("sequential"),
            "title": record.get("title") or "",
            "dates": (record.get("dates") or "").replace("\n", "<br>"),
            "opis": record.get("opis") or "",
            "unit": record.get("unit") or "",
            "count": record.get("count") or "",
            "retention": record.get("retention") or "",
            "note": record.get("note") or "",
        })

    stats = result["stats"]
    stats_message = (
        f"Всего номеров в диапазоне: {stats['total_in_range']}\n"
        f"Найдено подходящих дел: {stats['valid_count']}\n"
        f"Пропущенных номеров: {stats['skipped_numbers']}\n"
        f"Алиментных дел исключено: {stats['alimony_excluded']}\n"
        f"Категорий автоподставлено из словаря: {stats.get('auto_fixed', 0)}"
    )
    return jsonify({
        "records": preview,
        "stats": stats,
        "stats_message": stats_message,
        "filename": result["filename"],
    })


@app.route("/download/<path:filename>")
def download(filename):
    safe = os.path.basename(filename)
    path = os.path.join(OUTPUT_DIR, safe)
    if not os.path.isfile(path):
        abort(404)
    return send_file(path, as_attachment=True)


# ---------------------------------------------------------------------------
# API: организации
# ---------------------------------------------------------------------------

@app.route("/api/organizations", methods=["GET"])
def api_orgs_list():
    conn = get_db()
    try:
        return jsonify(db.list_organizations(conn))
    finally:
        conn.close()


@app.route("/api/organizations", methods=["POST"])
def api_orgs_add():
    data = request.get_json(force=True)
    name = str(data.get("name", "")).strip()
    if not name:
        return _json_error("Название организации не может быть пустым")
    claim_type = str(data.get("claim_type", "по иску")).strip()
    conn = get_db()
    try:
        db.add_organization(conn, name, claim_type)
        return jsonify({"ok": True})
    finally:
        conn.close()


@app.route("/api/organizations/<int:org_id>", methods=["PUT"])
def api_orgs_update(org_id):
    data = request.get_json(force=True)
    name = str(data.get("name", "")).strip()
    if not name:
        return _json_error("Название организации не может быть пустым")
    claim_type = str(data.get("claim_type", "по иску")).strip()
    conn = get_db()
    try:
        db.update_organization(conn, org_id, name, claim_type)
        return jsonify({"ok": True})
    finally:
        conn.close()


@app.route("/api/organizations/<int:org_id>", methods=["DELETE"])
def api_orgs_delete(org_id):
    conn = get_db()
    try:
        db.delete_organization(conn, org_id)
        return jsonify({"ok": True})
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# API: судебные участки (stage11d, под-этап A4)
# ---------------------------------------------------------------------------

@app.route("/api/court_areas", methods=["GET"])
def api_court_areas_list():
    conn = get_db()
    try:
        return jsonify(db.list_court_areas(conn))
    finally:
        conn.close()


@app.route("/api/court_areas", methods=["POST"])
def api_court_areas_add():
    data = request.get_json(force=True)
    conn = get_db()
    try:
        area_id = db.add_court_area(conn, data)
        return jsonify({"ok": True, "id": area_id})
    except ValueError as exc:
        return _json_error(str(exc))
    finally:
        conn.close()


@app.route("/api/court_areas/<int:area_id>", methods=["GET"])
def api_court_areas_get(area_id):
    conn = get_db()
    try:
        rec = db.get_court_area(conn, area_id)
        if not rec:
            return _json_error("Участок не найден", 404)
        return jsonify(rec)
    finally:
        conn.close()


@app.route("/api/court_areas/<int:area_id>", methods=["PUT"])
def api_court_areas_update(area_id):
    data = request.get_json(force=True)
    conn = get_db()
    try:
        if not db.get_court_area(conn, area_id):
            return _json_error("Участок не найден", 404)
        db.update_court_area(conn, area_id, data)
        return jsonify({"ok": True})
    except ValueError as exc:
        return _json_error(str(exc))
    finally:
        conn.close()


@app.route("/api/court_areas/<int:area_id>", methods=["DELETE"])
def api_court_areas_delete(area_id):
    conn = get_db()
    try:
        db.delete_court_area(conn, area_id)
        return jsonify({"ok": True})
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# API: импорт/экспорт базы данных
# ---------------------------------------------------------------------------

@app.route("/api/db/import", methods=["POST"])
def api_db_import():
    """Импорт организаций из Excel-файла (лист «База данных»)."""
    if "file" not in request.files:
        return _json_error("Файл не загружен")
    uploaded = request.files["file"]
    if not uploaded.filename:
        return _json_error("Файл не выбран")

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    path = os.path.join(UPLOAD_DIR, os.path.basename(uploaded.filename))
    uploaded.save(path)
    conn = get_db()
    try:
        added = db.import_database_from_excel(conn, path)
        return jsonify({"ok": True, "added": added})
    except Exception as exc:  # noqa: BLE001
        return _json_error(f"Ошибка импорта: {exc}", 400)
    finally:
        conn.close()


@app.route("/api/db/export")
def api_db_export():
    """Экспорт базы данных в Excel-файл."""
    conn = get_db()
    try:
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        filename = f"база_данных_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        path = os.path.join(OUTPUT_DIR, filename)
        db.export_database_to_excel(conn, path)
        return send_file(path, as_attachment=True,
                         download_name=filename)
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# API: правила сроков хранения
# ---------------------------------------------------------------------------

@app.route("/api/retention_rules", methods=["GET"])
def api_rules_list():
    conn = get_db()
    try:
        return jsonify(db.list_retention_rules(conn))
    finally:
        conn.close()


@app.route("/api/retention_rules", methods=["POST"])
def api_rules_add():
    data = request.get_json(force=True)
    keyword = str(data.get("keyword", "")).strip()
    code = str(data.get("code", "")).strip()
    if not keyword or not code:
        return _json_error("Укажите корень и код типа")
    conn = get_db()
    try:
        db.add_retention_rule(conn, keyword, code,
                              data.get("result_text"))
        return jsonify({"ok": True})
    finally:
        conn.close()


@app.route("/api/retention_rules/<int:rule_id>", methods=["PUT"])
def api_rules_update(rule_id):
    data = request.get_json(force=True)
    keyword = str(data.get("keyword", "")).strip()
    code = str(data.get("code", "")).strip()
    if not keyword or not code:
        return _json_error("Укажите корень и код типа")
    conn = get_db()
    try:
        db.update_retention_rule(conn, rule_id, keyword, code,
                                 data.get("result_text"))
        return jsonify({"ok": True})
    finally:
        conn.close()


@app.route("/api/retention_rules/<int:rule_id>", methods=["DELETE"])
def api_rules_delete(rule_id):
    conn = get_db()
    try:
        db.delete_retention_rule(conn, rule_id)
        return jsonify({"ok": True})
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# API: настройки
# ---------------------------------------------------------------------------

@app.route("/api/settings", methods=["GET"])
def api_settings_get():
    conn = get_db()
    try:
        return jsonify(db.get_settings(conn))
    finally:
        conn.close()


@app.route("/api/settings", methods=["POST"])
def api_settings_save():
    data = request.get_json(force=True)
    conn = get_db()
    try:
        for key, value in data.items():
            db.set_setting(conn, str(key), str(value))
        return jsonify({"ok": True})
    finally:
        conn.close()


if __name__ == "__main__":
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    # Инициализация БД при старте
    get_db().close()
    # Автозапуск браузера и контроль закрытия вкладки (остановка сервера)
    start_beat_watchdog()
    threading.Timer(1.0, lambda: webbrowser.open("http://127.0.0.1:5000")).start()
    app.run(host="127.0.0.1", port=5000, debug=False)