# -*- coding: utf-8 -*-
"""
Тесты чтения ранее обработанных файлов (Доработка 6): разбор заголовков
дел из результат_обработки_*.xlsx и акт_уничтожения_*.docx.

Запуск: .venv\\Scripts\\python.exe court_case_app\\tests\\test_legacy_reader.py
"""

import sys
import os
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from core.legacy_reader import (
    parse_title,
    parse_result_xlsx,
    parse_act_docx,
    parse_legacy_file,
)
from core.case_processor import build_case_record

FAILED = 0
PASSED = 0


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def main():
    tmp_dir = tempfile.mkdtemp(prefix="legacy_")

    print("[1] parse_title: корректные заголовки")
    r = parse_title(
        'Гражданское дело №2-3/2023 по иску ООО "Сбербанк" к Иванову Ивану '
        'Ивановичу о взыскании задолженности по кредитному договору')
    check("Номер дела", r["case_number"], "2-3/2023")
    check("Префикс", r["prefix"], "по иску")
    check("Истец", r["applicants"], 'ООО "Сбербанк"')
    check("Ответчик", r["respondents"], "Иванову Ивану Ивановичу")
    check("Категория", r["category"], "о взыскании задолженности по кредитному договору")

    r2 = parse_title(
        'Гражданское дело №2-11/2023 по заявлению Администрации города '
        'Оренбурга к Сидорову Петру о выдаче судебного приказа')
    check("Заявление: префикс", r2["prefix"], "по заявлению")
    check("Заявление: истец", r2["applicants"], "Администрации города Оренбурга")
    check("Заявление: ответчик", r2["respondents"], "Сидорову Петру")
    check("Заявление: категория", r2["category"], "о выдаче судебного приказа")

    r_nc = parse_title(
        'Гражданское дело №2-6/2023 по иску Петрова Петра к Сидоровой Анне')
    check("Без категории: ответчик = весь остаток",
          r_nc["respondents"], "Сидоровой Анне")

    print("\n[2] parse_title: проблемные/граничные случаи")
    r3 = parse_title(
        'Гражданское дело №2-4/2023 по иску Петровой Марии Ивановны '
        'к ООО "Ромашка"')
    check("Заголовок без категории -> категория пустая", r3["category"], "")
    # Истец в готовом заголовке уже склонён (родительный падеж)
    check("Без категории: истец", r3["applicants"], "Петровой Марии Ивановны")
    check("Без категории: ответчик в rest",
          r3["title"].split(" к ")[-1], 'ООО "Ромашка"')
    check("Не заголовок -> None", parse_title("Просто текст"), None)
    check("Пусто -> None", parse_title(""), None)
    check("Двойные пробелы схлопываются",
          parse_title('Гражданское дело №2-5/2023  по заявлению   А  к  Б '
                      ' о взыскании')["category"], "о взыскании")

    print("\n[3] parse_result_xlsx")
    from openpyxl import Workbook
    xlsx_path = os.path.join(tmp_dir, "результат_обработки.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Результат обработки"
    ws.append(["№ п/п", "Заголовок дела", "Даты дела", "№ описи", "№ ед.хр.",
               "Кол-во ед.хр.", "Срок хранения", "Примечание"])
    ws.append([1,
               'Гражданское дело №2-1/2023 по иску ООО "Сбербанк" к Иванову '
               'Ивану Ивановичу о взыскании задолженности по кредитному договору',
               "01.01.2023\n05.01.2023", "1", "1", "1", "3 года Ст. 223", ""])
    ws.append([2, 'Гражданское дело №2-2/2023 по заявлению Ивановой Марии '
                  'к Сидорову Петру', "", "", "", "", "", ""])
    ws.append([3, "", "", "", "", "", "", ""])  # пустая строка пропуска
    wb.save(xlsx_path)

    recs = parse_result_xlsx(xlsx_path)
    check("Из xlsx извлечено 2 записи", len(recs), 2)
    check("xlsx: номер 2-1", recs[0]["case_number"], "2-1/2023")
    check("xlsx: категория 2-1",
          recs[0]["category"], "о взыскании задолженности по кредитному договору")
    check("xlsx: дело 2-2 без категории", recs[1]["category"], "")

    print("\n[3а] Перенос данных «как есть» (Доработка 7)")
    check("маркер legacy_raw", recs[0].get("legacy_raw"), True)
    check("даты как есть", recs[0]["dates_raw"], "01.01.2023\n05.01.2023")
    check("№ описи как есть", recs[0]["opis_raw"], "1")
    check("№ ед.хр. как есть", recs[0]["unit_raw"], "1")
    check("кол-во как есть", recs[0]["count_raw"], "1")
    check("срок хранения как есть", recs[0]["retention_raw"], "3 года Ст. 223")
    check("пустое остаётся пустым (даты)", recs[1]["dates_raw"], "")
    check("пустое остаётся пустым (срок)", recs[1]["retention_raw"], "")

    print("\n[4] parse_act_docx")
    from docx import Document
    docx_path = os.path.join(tmp_dir, "акт_уничтожения.docx")
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    table.cell(0, 0).text = "1"
    table.cell(0, 1).text = ('Гражданское дело №2-7/2023 по иску ПАО Сбербанк '
                             'к Иванову Ивану Ивановичу о взыскании процентов')
    table.cell(1, 0).text = "2"
    table.cell(1, 1).text = ('Гражданское дело №2-8/2023 по заявлению '
                             'Петрова Петра Петровича к ООО "Ромашка"')
    doc.save(docx_path)

    recs_docx = parse_act_docx(docx_path)
    check("Из docx извлечено 2 записи", len(recs_docx), 2)
    check("docx: номер 2-7", recs_docx[0]["case_number"], "2-7/2023")
    check("docx: категория 2-7", recs_docx[0]["category"], "о взыскании процентов")
    check("docx: дело 2-8 без категории", recs_docx[1]["category"], "")

    print("\n[4а] parse_act_docx: перенос данных по шапке (Доработка 7)")
    docx2_path = os.path.join(tmp_dir, "акт_с_данными.docx")
    doc2 = Document()
    hdr = ["№ п/п", "Заголовок дела", "Даты дела", "№ описи",
           "№ ед.хр.", "Кол-во ед.хр.", "Срок хранения", "Примечание"]
    t = doc2.add_table(rows=2, cols=len(hdr))
    for i, h in enumerate(hdr):
        t.cell(0, i).text = h
    t.cell(1, 0).text = "1"
    t.cell(1, 1).text = ('Гражданское дело №2-7/2023 по иску ПАО Сбербанк '
                         'к Иванову Ивану Ивановичу о взыскании процентов')
    t.cell(1, 2).text = "08.12.2017\n08.01.2018"
    t.cell(1, 3).text = "4"
    t.cell(1, 4).text = "9"
    t.cell(1, 5).text = "1"
    t.cell(1, 6).text = "3 года Ст. 223"
    doc2.save(docx2_path)
    recs_d2 = parse_act_docx(docx2_path)
    check("docx2: одна запись", len(recs_d2), 1)
    check("docx2: даты как есть", recs_d2[0]["dates_raw"],
          "08.12.2017\n08.01.2018")
    check("docx2: № описи", recs_d2[0]["opis_raw"], "4")
    check("docx2: № ед.хр.", recs_d2[0]["unit_raw"], "9")
    check("docx2: срок хранения", recs_d2[0]["retention_raw"], "3 года Ст. 223")

    print("\n[5] parse_legacy_file (по расширению)")
    check("xlsx распознан", len(parse_legacy_file(xlsx_path)), 2)
    check("docx распознан", len(parse_legacy_file(docx_path)), 2)
    check("Неизвестное расширение -> пусто",
          parse_legacy_file(os.path.join(tmp_dir, "data.txt")), [])

    print("\n[5а] build_case_record: перенос «как есть» (Доработка 7)")
    legacy_row = recs[0]
    lr = build_case_record(legacy_row, legacy_row["prefix"], [], {},
                           category=legacy_row["category"])
    check("legacy: даты как есть", lr["dates"], "01.01.2023\n05.01.2023")
    check("legacy: № описи", lr["opis"], "1")
    check("legacy: № ед.хр.", lr["unit"], "1")
    check("legacy: кол-во (не 1 по умолчанию, а из файла)", lr["count"], "1")
    check("legacy: срок хранения как есть", lr["retention"], "3 года Ст. 223")
    check("legacy: заголовок строится", lr["title"].startswith(
        "Гражданское дело №2-1/2023"), True)

    legacy_empty = recs[1]
    lr2 = build_case_record(legacy_empty, legacy_empty["prefix"], [], {},
                            category=legacy_empty["category"])
    check("legacy: пустые даты остаются пустыми", lr2["dates"], "")
    check("legacy: пустой срок остаётся пустым", lr2["retention"], "")

    print("\n[5б] build_case_record: обычная обработка не изменилась")
    src = {"case_number": "2-9/2023", "date": "01.01.2023",
           "end_date": "05.02.2023", "applicants": "Петров Пётр",
           "respondents": "Сидоров Сидор", "category": "о взыскании долга",
           "opis_number": "5", "unit_number": "7"}
    nr = build_case_record(src, "по иску", [], {})
    check("обычная: count = 1", nr["count"], 1)
    check("обычная: № описи из исходника", nr["opis"], "5")
    check("обычная: № ед.хр. из исходника", nr["unit"], "7")
    check("обычная: срок пересчитан", bool(nr["retention"]), True)


    print("\n[6] Регрессия fix_05: объединённые ячейки не дублируют дела")
    doc2 = Document()
    # Вертикальный merge: одна ячейка на 2 строки
    t2 = doc2.add_table(rows=2, cols=3)
    t2.cell(0, 0).text = "1"
    t2.cell(0, 1).text = ('Гражданское дело №2-1/2023 по иску ООО "Сбербанк" '
                          'к Иванову И. И. о взыскании долга')
    # t2.cell(1,1) объединяем с t2.cell(0,1) по вертикали
    a = t2.cell(0, 1)
    b = t2.cell(1, 1)
    a.merge(b)
    doc2.save(os.path.join(tmp_dir, "merge_v.docx"))

    recs_v = parse_act_docx(os.path.join(tmp_dir, "merge_v.docx"))
    check("fix_05: вертикальный merge — одна запись", len(recs_v), 1)

    # Горизонтальный merge: ячейка на 2 столбца
    doc3 = Document()
    t3 = doc3.add_table(rows=1, cols=3)
    t3.cell(0, 0).text = "1"
    a = t3.cell(0, 1)
    b = t3.cell(0, 2)
    a.merge(b)
    t3.cell(0, 1).text = ('Гражданское дело №2-2/2023 по заявлению '
                          'Петрова П. П. к ООО "Ромашка" о защите прав')
    doc3.save(os.path.join(tmp_dir, "merge_h.docx"))

    recs_h = parse_act_docx(os.path.join(tmp_dir, "merge_h.docx"))
    check("fix_05: горизонтальный merge — одна запись", len(recs_h), 1)

    # Две разные таблицы с одним заголовком — дубликаты НЕ убираются
    doc4 = Document()
    for _ in range(2):
        t = doc4.add_table(rows=1, cols=2)
        t.cell(0, 1).text = ('Гражданское дело №2-3/2023 по иску '
                             'Сидорова С. С. к Петрову П. П. о взыскании')
    doc4.save(os.path.join(tmp_dir, "two_tables.docx"))
    recs_tt = parse_act_docx(os.path.join(tmp_dir, "two_tables.docx"))
    check("fix_05: разные таблицы не схлопываются", len(recs_tt), 2)

    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()