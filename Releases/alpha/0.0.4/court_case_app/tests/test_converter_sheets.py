# -*- coding: utf-8 -*-
"""
Тесты core.converter.sheets (stage12a, под-этап B1).

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tests\\test_converter_sheets.py
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import Workbook
from docx import Document

from core.converter.sheets import (
    RESULT_SHEET_NAME,
    inspect_sheets,
    sheet_exists,
    verify_docx_content,
    verify_xlsx_content,
)
from core.converter.exceptions import BadFileError, NoSheetsFound

FAILED = 0
PASSED = 0


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def expect_error(name, fn, exc_type):
    global FAILED, PASSED
    try:
        fn()
    except exc_type:
        PASSED += 1
        print(f"  OK   {name}")
        return
    except Exception as exc:  # noqa: BLE001
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось {exc_type.__name__}, "
              f"получено {type(exc).__name__}: {exc}")
        return
    FAILED += 1
    print(f"  FAIL {name}\n       исключение не поднято")


def make_xlsx(path, sheet_names):
    wb = Workbook()
    # Первый лист — он же active по умолчанию
    wb.active.title = sheet_names[0]
    for name in sheet_names[1:]:
        wb.create_sheet(name)
    wb.save(path)


def make_docx(path):
    doc = Document()
    doc.add_paragraph("тест")
    doc.save(path)


def main():
    tmp = tempfile.mkdtemp(prefix="conv_sheets_")

    print("[1] verify_xlsx_content: валидный xlsx")
    p_ok = os.path.join(tmp, "ok.xlsx")
    make_xlsx(p_ok, ["Лист1"])
    # Проверяем, что валидный файл не вызывает исключений
    try:
        verify_xlsx_content(p_ok)
        check("Валидный xlsx — без исключения", True, True)
    except Exception as exc:  # noqa: BLE001
        check("Валидный xlsx — без исключения", False, f"упало: {exc}")

    print("\n[2] verify_xlsx_content: не-xlsx")
    p_txt = os.path.join(tmp, "fake.xlsx")
    with open(p_txt, "w", encoding="utf-8") as f:
        f.write("это не excel")
    expect_error("Текстовый файл -> BadFileError",
                 lambda: verify_xlsx_content(p_txt), BadFileError)

    print("\n[3] verify_xlsx_content: docx как xlsx")
    p_docx = os.path.join(tmp, "real.docx")
    make_docx(p_docx)
    expect_error("docx под именем .xlsx -> BadFileError",
                 lambda: verify_xlsx_content(p_docx), BadFileError)

    print("\n[4] verify_xlsx_content: несуществующий")
    expect_error("Нет файла -> BadFileError",
                 lambda: verify_xlsx_content(os.path.join(tmp, "no.xlsx")),
                 BadFileError)

    print("\n[5] inspect_sheets: активный первый")
    p1 = os.path.join(tmp, "first.xlsx")
    make_xlsx(p1, ["Лист1", "Лист2"])
    res = inspect_sheets(p1)
    check("sheets содержит оба листа", res["sheets"], ["Лист1", "Лист2"])
    check("active = Лист1", res["active"], "Лист1")
    check("default = active (нет 'Результат обработки')", res["default"], "Лист1")

    print("\n[6] inspect_sheets: 'Результат обработки' приоритетнее active")
    p2 = os.path.join(tmp, "with_result.xlsx")
    make_xlsx(p2, ["Лист1", RESULT_SHEET_NAME, "Служебный"])
    res = inspect_sheets(p2)
    check("sheets все три", res["sheets"],
          ["Лист1", RESULT_SHEET_NAME, "Служебный"])
    check("active = Лист1 (первый)", res["active"], "Лист1")
    check("default = 'Результат обработки'", res["default"], RESULT_SHEET_NAME)

    print("\n[7] sheet_exists")
    check("Лист1 есть", sheet_exists(p2, "Лист1"), True)
    check("Нет такого листа", sheet_exists(p2, "НетТакого"), False)
    check("'Результат обработки' есть", sheet_exists(p2, RESULT_SHEET_NAME), True)

    print("\n[8] verify_docx_content")
    p_docx_ok = os.path.join(tmp, "ok.docx")
    make_docx(p_docx_ok)
    try:
        verify_docx_content(p_docx_ok)
        check("Валидный docx — без исключения", True, True)
    except Exception as exc:  # noqa: BLE001
        check("Валидный docx — без исключения", False, f"упало: {exc}")
    expect_error("Текстовый файл как .docx -> BadFileError",
                 lambda: verify_docx_content(p_txt), BadFileError)
    expect_error("xlsx как docx -> BadFileError",
                 lambda: verify_docx_content(p_ok), BadFileError)

    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()
