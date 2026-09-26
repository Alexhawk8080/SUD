# -*- coding: utf-8 -*-
"""
Тесты API судебных участков и /process с court_area_id (stage11d, A4).

Использует app.test_client() с подменой DB_PATH, OUTPUT_DIR, UPLOAD_DIR
на временные — чтобы не задеть рабочую БД.

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tests\\test_api_court_areas.py
"""

import io
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import Workbook
from docx import Document

# Импортируем модуль app (сам сервер не запустится — только __main__)
import app as app_module


FAILED = 0
PASSED = 0


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def make_area(номер="5", судья="Пятый П. П."):
    return {
        "номер": номер, "судья": судья,
        "дата_утверждения": "«___» ______________ 20__ года",
        "дата_акта": "«__» ____________ 20__ г.",
        "номер_акта": "1",
        "секретарь": "Сек5 С. С.",
        "дата_подписи": "«__» ____________ 20__ г.",
        "протокол_эк_дата": "«__» ____________ 20__ г.",
        "протокол_эк_номер": "1",
    }


def make_source_xlsx(path):
    """Минимальный xlsx: заголовки + 2 дела за 2020."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Дата", "№ дела", "Заявители", "Ответчики", "Категория",
               "№ описи", "№ ед.хр.", "Примечание", "Дата окончания"])
    ws.append([None, "2-1/2020", "Тестов Тест Тестович",
               "Тестова Теста Тестовна",
               "Исковое заявление о расторжении брака",
               1, 1, "", None])
    ws.append([None, "2-2/2020", "Другой Друг Другович",
               "Другая Друга Друговна",
               "Исковое заявление о расторжении брака",
               1, 2, "", None])
    wb.save(path)


def docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def main():
    tmp = tempfile.mkdtemp(prefix="api_areas_")
    db_path = os.path.join(tmp, "app.db")
    out_dir = os.path.join(tmp, "out")
    up_dir = os.path.join(tmp, "up")
    os.makedirs(out_dir, exist_ok=True)
    os.makedirs(up_dir, exist_ok=True)

    # Подмена глобальных путей в модуле app
    app_module.DB_PATH = db_path
    app_module.OUTPUT_DIR = out_dir
    app_module.UPLOAD_DIR = up_dir

    client = app_module.app.test_client()

    print("[1] GET /api/court_areas: миграция даёт 1 запись")
    r = client.get("/api/court_areas")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("Список — список", isinstance(j, list), True)
    check("Одна запись", len(j), 1)
    check("Номер = '3' (из settings)", j[0]["номер"], "3")

    print("\n[2] POST /api/court_areas: добавление")
    r = client.post("/api/court_areas", json=make_area("5", "Пятый П. П."))
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("ok=true", j.get("ok"), True)
    id5 = j.get("id")
    check("id получен", isinstance(id5, int) and id5 > 0, True)

    print("\n[3] POST с пустым номером -> 400")
    bad = make_area("", "Кто-то")
    r = client.post("/api/court_areas", json=bad)
    check("HTTP 400", r.status_code, 400)
    check("Ошибка про номер", "Номер" in (r.get_json() or {}).get("error", ""), True)

    print("\n[4] GET /api/court_areas/<id>")
    r = client.get(f"/api/court_areas/{id5}")
    check("HTTP 200", r.status_code, 200)
    rec = r.get_json()
    check("Номер = 5", rec["номер"], "5")
    check("Судья Пятый", rec["судья"], "Пятый П. П.")

    print("\n[5] GET несуществующего -> 404")
    r = client.get("/api/court_areas/99999")
    check("HTTP 404", r.status_code, 404)

    print("\n[6] PUT /api/court_areas/<id>")
    upd = make_area("5", "Обновлённый О. О.")
    upd["секретарь"] = "СекОбн С. С."
    r = client.put(f"/api/court_areas/{id5}", json=upd)
    check("HTTP 200", r.status_code, 200)
    r = client.get(f"/api/court_areas/{id5}")
    check("Судья обновлён", r.get_json()["судья"], "Обновлённый О. О.")

    print("\n[7] PUT несуществующего -> 404")
    r = client.put("/api/court_areas/99999", json=make_area("99", "X"))
    check("HTTP 404", r.status_code, 404)

    print("\n[8] DELETE /api/court_areas/<id>")
    # Добавим временный и удалим
    r = client.post("/api/court_areas", json=make_area("7", "Семёрка"))
    tmp_id = r.get_json()["id"]
    r = client.delete(f"/api/court_areas/{tmp_id}")
    check("HTTP 200", r.status_code, 200)
    r = client.get(f"/api/court_areas/{tmp_id}")
    check("После удаления 404", r.status_code, 404)

    print("\n[9] /process с court_area_id='5'")
    xlsx_path = os.path.join(tmp, "источник.xlsx")
    make_source_xlsx(xlsx_path)
    with open(xlsx_path, "rb") as f:
        data = {
            "file": (io.BytesIO(f.read()), "источник.xlsx"),
            "year": "2020",
            "format": "word",
            "alimony": "0",
            "court_area_id": str(id5),
        }
        r = client.post("/process", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("filename содержит '5СУ'", "5СУ" in (j or {}).get("filename", ""), True)
    # Проверим, что в файле реквизиты участка
    out_path = os.path.join(out_dir, j["filename"])
    text = docx_text(out_path)
    check("Судья участка 5 в акте", "Обновлённый О. О." in text, True)
    check("Секретарь участка 5 в акте", "СекОбн С. С." in text, True)

    print("\n[10] /process без court_area_id -> fallback на settings ('3СУ')")
    with open(xlsx_path, "rb") as f:
        data = {
            "file": (io.BytesIO(f.read()), "источник.xlsx"),
            "year": "2020",
            "format": "word",
            "alimony": "0",
        }
        r = client.post("/process", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("filename содержит '3СУ'", "3СУ" in (j or {}).get("filename", ""), True)

    print("\n[11] /process с несуществующим id -> fallback ('3СУ')")
    with open(xlsx_path, "rb") as f:
        data = {
            "file": (io.BytesIO(f.read()), "источник.xlsx"),
            "year": "2020",
            "format": "word",
            "alimony": "0",
            "court_area_id": "99999",
        }
        r = client.post("/process", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("Fallback '3СУ'", "3СУ" in (j or {}).get("filename", ""), True)

    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()
