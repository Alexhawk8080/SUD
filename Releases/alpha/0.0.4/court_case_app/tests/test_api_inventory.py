# -*- coding: utf-8 -*-
"""
Тесты API модуля «Внутренняя опись» (этап 14e).

Использует app.test_client() с подменой DB_PATH/OUTPUT_DIR/UPLOAD_DIR.

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tests\\test_api_inventory.py
"""

import io
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document
from openpyxl import Workbook, load_workbook

import app as app_module

FAILED = 0
PASSED = 0


def check(name, actual, expected):
    global FAILED, PASSED
    if actual == expected:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n"
              f"       получено:  {actual!r}")


HEADERS = ("№ п/п", "Наименование", "Номера листов")


def make_inventory_docx(path, numbers, kind="civil"):
    """Реальная docx-опись: таблица 3 столбца."""
    word = ("гражданскому" if kind == "civil"
            else "об административном правонарушении")
    doc = Document()
    doc.add_paragraph("Внутренняя опись дел")
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(HEADERS):
        table.cell(0, i).text = h
    for idx, n in enumerate(numbers, start=1):
        row = table.add_row()
        row.cells[0].text = str(idx)
        row.cells[1].text = f"Решение по {word} делу № 2-{n}/2018"
        row.cells[2].text = "1"
    doc.save(path)
    return path


def make_template(path, last_row=4):
    """Шаблон с листом «Таблица» и формулами."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Таблица"
    for row in range(4, last_row + 1):
        ws.cell(row=row, column=1).value = "=ROW()"
        ws.cell(row=row, column=2).value = f"=IF(F{row}=F{row-1},B{row-1},1)"
        ws.cell(row=row, column=6).value = f"=F{row-1}+E{row}"
        ws.cell(row=row, column=7).value = 2018
        ws.cell(row=row, column=9).value = f"=I{row-1}+1"
        ws.cell(row=row, column=11).value = "=ROW()"
        ws.cell(row=row, column=12).value = f"=A{row}"
    wb.save(path)
    return path


def main():
    tmp = tempfile.mkdtemp(prefix="api_inv_")
    os.makedirs(os.path.join(tmp, "out"), exist_ok=True)
    os.makedirs(os.path.join(tmp, "up"), exist_ok=True)
    app_module.DB_PATH = os.path.join(tmp, "app.db")
    app_module.OUTPUT_DIR = os.path.join(tmp, "out")
    app_module.UPLOAD_DIR = os.path.join(tmp, "up")
    client = app_module.app.test_client()

    doc1 = make_inventory_docx(os.path.join(tmp, "vol1.docx"), [14, 15, 16])
    doc2 = make_inventory_docx(os.path.join(tmp, "vol2.docx"), [17, 18])
    tpl = make_template(os.path.join(tmp, "tpl.xlsx"), last_row=4)

    print("[1] GET /inventory_page")
    r = client.get("/inventory_page")
    check("HTTP 200", r.status_code, 200)
    check("есть id='inventory-card'",
          'id="inventory-card"' in r.get_data(as_text=True), True)

    print("\n[2] POST /inventory/inspect: один docx")
    with open(doc1, "rb") as f:
        data = {"files": (io.BytesIO(f.read()), "vol1.docx")}
        r = client.post("/inventory/inspect", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("тип — гражданских", j["stats"]["type_word"], "гражданских")
    check("год потока", j["stats"]["year"], 2018)
    check("дел записано", j["stats"]["count"], 3)
    check("первое дело", j["stats"]["first"], "2-14/2018")
    check("последнее дело", j["stats"]["last"], "2-16/2018")
    check("preview непусто", len(j["preview"]), 3)

    print("\n[3] POST /inventory/inspect: два docx (стыковка)")
    with open(doc1, "rb") as f1, open(doc2, "rb") as f2:
        data = {"files": [(io.BytesIO(f1.read()), "vol1.docx"),
                          (io.BytesIO(f2.read()), "vol2.docx")]}
        r = client.post("/inventory/inspect", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("файлов обработано", j["stats"]["files"], 2)
    check("дел всего", j["stats"]["count"], 5)
    check("последнее дело", j["stats"]["last"], "2-18/2018")

    print("\n[4] POST /inventory/inspect: без файлов -> 400")
    r = client.post("/inventory/inspect", data={},
                    content_type="multipart/form-data")
    check("HTTP 400", r.status_code, 400)

    print("\n[5] POST /inventory/convert: валидация")
    with open(doc1, "rb") as f:
        data = {"files": (io.BytesIO(f.read()), "vol1.docx")}
        r = client.post("/inventory/convert", data=data,
                        content_type="multipart/form-data")
    check("без шаблона -> 400", r.status_code, 400)

    with open(doc1, "rb") as f:
        data = {"template": (io.BytesIO(b"x"), "bad.docx"),
                "files": (io.BytesIO(f.read()), "vol1.docx")}
        r = client.post("/inventory/convert", data=data,
                        content_type="multipart/form-data")
    check("шаблон не .xlsx -> 400", r.status_code, 400)

    print("\n[6] POST /inventory/convert: успешная конвертация")
    with open(tpl, "rb") as tf, open(doc1, "rb") as f:
        data = {"template": (io.BytesIO(tf.read()), "tpl.xlsx"),
                "files": (io.BytesIO(f.read()), "vol1.docx"),
                "court_area": "9СУ"}
        r = client.post("/inventory/convert", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("имя с участком", "9СУ" in (j.get("filename") or ""), True)
    check("файл создан",
          os.path.exists(os.path.join(app_module.OUTPUT_DIR,
                                      j["filename"])), True)
    wb_out = load_workbook(os.path.join(app_module.OUTPUT_DIR,
                                        j["filename"]))
    check("лист «Требует внимания»",
          "Требует внимания" in wb_out.sheetnames, True)
    ws_out = wb_out["Таблица"]
    check("C4 — решение", ws_out.cell(row=4, column=3).value,
          "Решение по гражданскому делу")
    check("F4", ws_out.cell(row=4, column=6).value, 14)

    print("\n[7] POST /inventory/convert: конфликт и перезапись")
    with open(tpl, "rb") as tf, open(doc1, "rb") as f:
        data = {"template": (io.BytesIO(tf.read()), "tpl.xlsx"),
                "files": (io.BytesIO(f.read()), "vol1.docx"),
                "court_area": "9СУ"}
        r = client.post("/inventory/convert", data=data,
                        content_type="multipart/form-data")
    check("повторно -> 409", r.status_code, 409)
    check("conflict=true", r.get_json().get("conflict"), True)

    with open(tpl, "rb") as tf, open(doc1, "rb") as f:
        data = {"template": (io.BytesIO(tf.read()), "tpl.xlsx"),
                "files": (io.BytesIO(f.read()), "vol1.docx"),
                "court_area": "9СУ", "overwrite": "1"}
        r = client.post("/inventory/convert", data=data,
                        content_type="multipart/form-data")
    check("overwrite -> 200", r.status_code, 200)

    print("\n[8] Скачивание результата")
    with open(tpl, "rb") as tf, open(doc1, "rb") as f:
        data = {"template": (io.BytesIO(tf.read()), "tpl.xlsx"),
                "files": (io.BytesIO(f.read()), "vol1.docx"),
                "court_area": "9СУ", "overwrite": "1", "download": "1"}
        r = client.post("/inventory/convert", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    check("attachment",
          "attachment" in r.headers.get("Content-Disposition", ""), True)

    print("\n" + "=" * 60)
    print(f"ИТОГО: PASSED = {PASSED}, FAILED = {FAILED}")
    return 1 if FAILED else 0


if __name__ == "__main__":
    sys.exit(main())