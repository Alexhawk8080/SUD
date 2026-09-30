# stage_46
# fix_48_01
# stage_48
# stage_47
# fix_46_01
# -*- coding: utf-8 -*-
"""
Перенос данных из Excel-результата в Word-акт.

Порт VBA Sub ПереносДанныхСНумерациейWord() с отличиями:
  * строки данных в целевой таблице предварительно очищаются
    (кроме заголовка) — повторный прогон идемпотентен;
  * исходный Word-файл не перезаписывается, результат сохраняется
    рядом с суффиксом `_filled`;
  * автонумерация столбца 1 сохраняется путём копирования
    XML-структуры строки-образца (deepcopy <w:tr>).

Использование:
    from core.converter.excel_to_act import import_excel_to_act
    info = import_excel_to_act("result.xlsx", "act.docx")
"""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook


SHEET_NAME = "Результат обработки"
HEADER_LABELS = ["1", "2", "3", "4", "5", "6", "7", "8"]
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


# --- Поиск целевой таблицы --------------------------------------------------

def _cell_text(cell) -> str:
    """Текст ячейки без завершающего \\r\\x07, который добавляет Word."""
    return cell.text.replace("\r", "").replace("\x07", "").strip()


def _find_target_table(doc):
    """Ищет таблицу, у которой первая строка содержит «1»..«8».
    Возвращает (index, table) или (None, None)."""
    for idx, table in enumerate(doc.tables):
        if len(table.rows) < 1 or len(table.columns) < 8:
            continue
        try:
            if all(_cell_text(table.cell(0, c)) == HEADER_LABELS[c]
                   for c in range(8)):
                return idx, table
        except Exception:
            continue
    return None, None


# --- Чтение Excel -----------------------------------------------------------

def _read_excel_rows(excel_path: Path):
    wb = load_workbook(excel_path, data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError(
                f"В файле Excel нет листа «{SHEET_NAME}». "
                f"Доступные листы: {', '.join(wb.sheetnames)}"
            )
        ws = wb[SHEET_NAME]

        # Последняя заполненная строка по столбцу A (аналог xlUp).
        last_row = 0
        for r in range(ws.max_row, 0, -1):
            v = ws.cell(row=r, column=1).value
            if v is not None and str(v).strip():
                last_row = r
                break
        if last_row < 2:
            raise ValueError(
                f"На листе «{SHEET_NAME}» нет данных (только заголовок)."
            )

        rows = []
        for r in range(2, last_row + 1):
            row_vals = []
            for c in range(2, 9):  # B..H
                v = ws.cell(row=r, column=c).value
                row_vals.append("" if v is None else str(v))
            rows.append(row_vals)
        return rows
    finally:
        wb.close()


# --- Работа с Word ----------------------------------------------------------

def _clear_data_rows(table):
    """Удаляет все строки кроме первой (заголовка).
    Возвращает deepcopy первой строки данных (образец для копирования).
    Если строк данных не было — возвращает None."""
    trs = table._tbl.findall(f"{W_NS}tr")
    if len(trs) < 2:
        return None
    sample = deepcopy(trs[1])
    for tr in trs[1:]:
        table._tbl.remove(tr)
    return sample


def _set_cell_text(cell, text: str) -> None:
    """Записывает текст, сохраняя форматирование первого run.
    Удаляет лишние абзацы в ячейке."""
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.add_paragraph(text)
        return
    first = paragraphs[0]
    if first.runs:
        first.runs[0].text = text
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.add_run(text)
    for p in paragraphs[1:]:
        p._element.getparent().remove(p._element)


def _set_tc_text(tc, text: str) -> None:
    """Установить текст в <w:tc> напрямую через XML.

    Сохраняет <w:rPr> первого run (форматирование), удаляет лишние
    абзацы и runs. Работает быстро на больших N — не пересчитывает
    grid таблицы, в отличие от python-docx table.cell().
    """
    ps = tc.findall(qn("w:p"))
    if not ps:
        p = OxmlElement("w:p")
        tc.append(p)
    else:
        for extra in ps[1:]:
            tc.remove(extra)
        p = ps[0]

    runs = p.findall(qn("w:r"))
    if not runs:
        r = OxmlElement("w:r")
        p.append(r)
    else:
        for extra in runs[1:]:
            p.remove(extra)
        r = runs[0]

    # Оставляем только <w:rPr> внутри run, остальное удаляем.
    for child in list(r):
        if child.tag != qn("w:rPr"):
            r.remove(child)

    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)


# --- Публичный API ----------------------------------------------------------

def import_excel_to_act(excel_path, word_path, output_path=None,
                        on_progress=None) -> dict:
    """Переносит данные из Excel-листа «Результат обработки» в Word-таблицу
    с заголовками «1»..«8».

    Параметры:
        excel_path  — путь к .xlsx с листом «Результат обработки».
        word_path   — путь к .docx с целевой таблицей.
        output_path — куда сохранить (по умолчанию <word>_filled.docx).
        on_progress — опциональный колбэк (message, current, total).
                      Вызывается на ключевых этапах, чтобы UI мог
                      показать прогресс. total=0 — «этап без счётчика».

    Возвращает:
        {
          "rows_written": int,   # сколько строк добавлено
          "table_index":  int,   # индекс целевой таблицы (0-based)
          "output_path":  str,   # путь к сохранённому файлу
          "sheet_used":   str,   # имя использованного листа
        }

    Бросает ValueError при проблемах с исходными файлами.
    """
    def _emit(message, current=0, total=0):
        if on_progress is not None:
            try:
                on_progress(message, current, total)
            except Exception:
                pass

    excel_path = Path(excel_path)
    word_path = Path(word_path)
    if not excel_path.exists():
        raise ValueError(f"Excel-файл не найден: {excel_path}")
    if not word_path.exists():
        raise ValueError(f"Word-файл не найден: {word_path}")

    _emit("Чтение Excel…")
    rows = _read_excel_rows(excel_path)
    _emit(f"Прочитано строк Excel: {len(rows)}", len(rows), len(rows))

    _emit("Поиск таблицы в Word…")
    doc = Document(str(word_path))
    table_idx, table = _find_target_table(doc)
    if table is None:
        raise ValueError(
            "В документе не найдена таблица с заголовками '1'..'8' "
            "и минимум 8 столбцами."
        )
    _emit("Таблица найдена")

    _emit("Очистка старых строк…")
    sample_tr = _clear_data_rows(table)
    if sample_tr is None:
        # Строк данных нет — образцом служит заголовок.
        sample_tr = deepcopy(table._tbl.findall(f"{W_NS}tr")[0])

    total = len(rows)
    _emit("Запись строк в Word", 0, total)
    tbl = table._tbl
    sample_tcs = sample_tr.findall(qn("w:tc"))
    n_cols = len(sample_tcs)

    # fix_48_01: если в образце столбца 1 есть <w:numPr> — автонумерация
    # Word активна, оставляем столбец пустым. Если numPr нет — пишем
    # номер текстом (1, 2, 3, ...), как это было в исходном акте.
    has_numpr = False
    if n_cols >= 1:
        p0 = sample_tcs[0].find(qn("w:p"))
        if p0 is not None:
            ppr0 = p0.find(qn("w:pPr"))
            if ppr0 is not None and ppr0.find(qn("w:numPr")) is not None:
                has_numpr = True

    # stage_48: работаем с XML напрямую. table.cell() из python-docx
    # пересчитывает grid таблицы на каждый вызов — при 5000+ строк это
    # даёт квадратичное замедление.
    for idx, row_vals in enumerate(rows, start=1):
        new_tr = deepcopy(sample_tr)
        tcs = new_tr.findall(qn("w:tc"))
        # Столбец 1: либо пусто (numPr сам отрисует номер), либо текст.
        if n_cols >= 1:
            _set_tc_text(tcs[0], "" if has_numpr else str(idx))
        # Столбцы 2..N: значения из строки Excel.
        for j, val in enumerate(row_vals):
            col = j + 1
            if col < n_cols:
                _set_tc_text(tcs[col], val)
        tbl.append(new_tr)
        if idx % 100 == 0 or idx == total:
            _emit(f"Записано {idx} из {total}", idx, total)

    if output_path is None:
        output_path = word_path.with_name(
            f"{word_path.stem}_filled{word_path.suffix}")
    output_path = Path(output_path)
    _emit("Сохранение…")
    doc.save(str(output_path))
    _emit("Готово", total, total)

    return {
        "rows_written": len(rows),
        "table_index": table_idx,
        "output_path": str(output_path),
        "sheet_used": SHEET_NAME,
    }
