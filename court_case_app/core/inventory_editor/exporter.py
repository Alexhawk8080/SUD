# stage_49b
"""Замена C/D/E/G/H в готовом .xlsx. Остальные ячейки не трогаем.

Правила:
  * строки 4 .. 4+N-1   — пишем C/D/E/G/H из docs (E — производное);
  * строки 4+N .. max   — зануляем только G, C/D/E/H не трогаем;
  * листы «База» и «На пичать» — не трогаем;
  * файл-шаблон не перезаписываем — сохраняем как output_path.
"""
from __future__ import annotations
from typing import List
from .model import Doc
from .calculator import compute_derived

DATA_START_ROW = 4


def export_to_excel(docs: List[Doc], template_path: str, output_path: str,
                    sheet_name: str = "Таблица") -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(template_path, data_only=False)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Лист '{sheet_name}' не найден в {template_path}")
    ws = wb[sheet_name]

    max_row = ws.max_row or DATA_START_ROW
    n = len(docs)
    last_data_row = DATA_START_ROW + n - 1
    if last_data_row > max_row:
        raise ValueError(
            f"В шаблоне только {max_row - DATA_START_ROW + 1} строк данных, "
            f"а нужно {n}. Расширьте шаблон."
        )

    derived = compute_derived(docs)
    for i, (d, der) in enumerate(zip(docs, derived)):
        r = DATA_START_ROW + i
        ws.cell(r, 3).value = d.title
        ws.cell(r, 4).value = d.prefix
        ws.cell(r, 5).value = der["E"]
        ws.cell(r, 7).value = d.year
        ws.cell(r, 8).value = d.pages_count

    # Хвост: зануляем G. C/D/E/H не трогаем.
    tail_zeroed = 0
    for r in range(last_data_row + 1, max_row + 1):
        ws.cell(r, 7).value = 0
        tail_zeroed += 1

    wb.save(output_path)
    return {"rows": n, "tail_zeroed": tail_zeroed, "output": output_path,
            "last_data_row": last_data_row}
