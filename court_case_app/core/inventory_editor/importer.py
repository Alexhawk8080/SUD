# stage_49b
"""Импорт листа «Таблица» из .xlsx в список Doc.

Читаем колонки:
  C (3) — Наименование      (обязательна; конец данных = первая пустая C)
  D (4) — Префикс           (2, 2а, 5, …)
  E (5) — Увеличение        (накопительно: case_number = prev + E)
  G (7) — Год
  H (8) — Кол-во стр.
Данные начинаются со строки 4 (строки 1-3 — служебные: шапка).
"""
from __future__ import annotations
from typing import List, Optional
from .model import Doc

DATA_START_ROW = 4


def _read_cell(ws, r: int, c: int):
    v = ws.cell(r, c).value
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if not v or v.startswith("="):
            return None
    return v


def _to_int(v, default: int = 0) -> int:
    if v is None or v == "":
        return default
    try:
        return int(v)
    except (TypeError, ValueError):
        try:
            return int(float(v))
        except (TypeError, ValueError):
            return default


def import_from_excel(path: str, sheet_name: str = "Таблица") -> List[Doc]:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=False)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Лист '{sheet_name}' не найден в {path}")
    ws = wb[sheet_name]

    docs: List[Doc] = []
    prev_case = 0
    r = DATA_START_ROW
    max_r = ws.max_row or DATA_START_ROW
    while r <= max_r:
        title = _read_cell(ws, r, 3)
        if not title:
            break
        prefix = _read_cell(ws, r, 4) or ""
        e_val = _read_cell(ws, r, 5)
        year = _read_cell(ws, r, 7)
        pages = _read_cell(ws, r, 8)

        case_number = prev_case + _to_int(e_val, 0)
        docs.append(Doc(
            title=str(title),
            prefix=str(prefix),
            case_number=case_number,
            year=_to_int(year, 0),
            pages_count=_to_int(pages, 1) or 1,
        ))
        prev_case = case_number
        r += 1
    return docs
