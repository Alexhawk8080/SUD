# -*- coding: utf-8 -*-
"""
Чтение исходной таблицы административных дел из Excel (.xlsx).

Структура файла «2020 адм.xlsx» (шапка в строке 3):
    A  Вх. №                     -> in_number
    B  Дата поступления          -> date
    C  № Дела в производстве     -> case_number (5-N/YYYY)
    D  ФИО/наименование лица     -> person
    E  Статья                    -> article
    H  Текущая задача            -> task
    K  Окончание производства    -> end_date
    M  Состояние акта            -> act_state

Отдельный модуль (не трогаем стабильный excel_reader.py для гражданских):
роли и алиасы заголовков принципиально другие.
"""

from openpyxl import load_workbook

# Роли столбцов административной таблицы
ADMIN_ROLES = [
    "in_number",     # входящий номер
    "date",          # дата поступления
    "case_number",   # № дела в производстве
    "person",        # ФИО/наименование лица
    "article",       # статья
    "task",          # текущая задача
    "end_date",      # окончание производства
    "act_state",     # состояние акта
]

ADMIN_ROLE_TITLES = {
    "in_number": "Вх. №",
    "date": "Дата поступления",
    "case_number": "№ дела",
    "person": "ФИО/наименование лица",
    "article": "Статья",
    "task": "Текущая задача",
    "end_date": "Окончание производства",
    "act_state": "Состояние акта",
}

# Варианты заголовков для автоопределения (нормализованные)
ADMIN_HEADER_ALIASES = {
    "in_number": ["вх №", "вх номер", "входящий номер", "вх"],
    "date": ["дата поступления", "дата"],
    "case_number": ["№ дела в производстве", "№ дела", "номер дела",
                    "№ дела в производстве"],
    "person": ["фио наименование лица", "фио", "наименование лица", "лицо"],
    "article": ["статья"],
    "task": ["текущая задача", "задача"],
    "end_date": ["окончание производства", "дата окончания",
                 "окончание"],
    "act_state": ["состояние акта"],
}


class AdminMappingError(Exception):
    """Ошибка маппинга столбцов административной таблицы."""


class AdminRow:
    """Строка административной таблицы: значения по ролям (None для пустых)."""

    def __init__(self, data: dict):
        self._data = data

    def get(self, role: str):
        return self._data.get(role)

    @property
    def data(self) -> dict:
        return dict(self._data)


def _normalize_header(value) -> str:
    """
    Нормализация заголовка: строка, нижний регистр, схлопнутые пробелы,
    удаление дефисов и слэшей (Word переносит длинные слова через «-»,
    а «ФИО/наименование» содержит «/»).
    """
    if value is None:
        return ""
    text = " ".join(str(value).strip().split())
    return text.lower().replace("-", "").replace("/", " ").replace("  ", " ")


def auto_detect_admin_roles(headers) -> dict:
    """
    Автоопределение ролей по заголовкам (двухпроходный алгоритм).

    Возвращает словарь role -> индекс столбца (1-based).
    """
    mapping = {}
    used_cols = set()

    def _assign(candidates):
        candidates.sort(key=lambda c: (-c[0], c[2]))
        for _ln, role, idx in candidates:
            if role in mapping or idx in used_cols:
                continue
            mapping[role] = idx
            used_cols.add(idx)

    # 1-й проход: точные совпадения
    exact = []
    for idx, header in enumerate(headers, start=1):
        h = _normalize_header(header)
        if not h:
            continue
        for role, aliases in ADMIN_HEADER_ALIASES.items():
            for alias in aliases:
                if h == alias:
                    exact.append((len(alias), role, idx))
    _assign(exact)

    # 2-й проход: по началу строки
    prefix = []
    for idx, header in enumerate(headers, start=1):
        if idx in used_cols:
            continue
        h = _normalize_header(header)
        if not h:
            continue
        for role, aliases in ADMIN_HEADER_ALIASES.items():
            if role in mapping:
                continue
            for alias in aliases:
                if h.startswith(alias):
                    prefix.append((len(alias), role, idx))
    _assign(prefix)

    return mapping


def detect_admin_header_row(rows) -> int:
    """
    Автоопределение номера строки с заголовками (1-based).

    Перебирает первые строки листа и возвращает индекс строки, где
    распознано максимальное число ролей. Если совпадений нет — 1.
    """
    best_row = 1
    best_score = 0
    for idx, row in enumerate(rows[:20], start=1):
        mapping = auto_detect_admin_roles(row)
        score = len(mapping)
        if score > best_score:
            best_score = score
            best_row = idx
    return best_row


def read_admin_table(filepath: str, mapping=None, sheet_name: str = None,
                     header_row: int = None):
    """
    Чтение административной таблицы.

    Параметры:
        filepath    — путь к .xlsx файлу;
        mapping     — ручной маппинг role -> индекс столбца (1-based);
        sheet_name  — имя листа (по умолчанию первый активный);
        header_row  — номер строки с заголовками (1-based).
                      None = автоопределение; 0 = заголовков нет.

    Возвращает кортеж (headers, rows, used_mapping):
        headers      — список заголовков строки header_row ([] при 0);
        rows         — список AdminRow (данные со следующей строки);
        used_mapping — словарь role -> индекс столбца.
    """
    wb = load_workbook(filepath, data_only=True, read_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        all_rows = list(ws.iter_rows(values_only=True))

        if header_row is None:
            header_row = detect_admin_header_row(all_rows)

        headers = []
        if header_row > 0 and header_row <= len(all_rows):
            headers = list(all_rows[header_row - 1])

        if mapping is None:
            mapping = auto_detect_admin_roles(headers)

        rows = []
        start = header_row if header_row > 0 else 0
        for values in all_rows[start:]:
            data = {}
            for role, col_idx in mapping.items():
                if 1 <= col_idx <= len(values):
                    data[role] = values[col_idx - 1]
                else:
                    data[role] = None
            rows.append(AdminRow(data))
        return headers, rows, mapping
    finally:
        wb.close()
