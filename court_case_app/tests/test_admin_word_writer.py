# -*- coding: utf-8 -*-
"""
Тесты формирования Word-акта по admin-шаблону (D8d).

Проверяет: выбор admin-шаблона, формулировку {{итого}} для admin-дел,
подстановку {{количество_прописью}}, отсутствие оставшихся плейсхолдеров.

Запуск: .venv\\Scripts\\python.exe court_case_app\\tests\\test_admin_word_writer.py
"""

import os
import re
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document
from core.word_writer import write_word_result, number_to_words

FAILED = 0
PASSED = 0

TEMPLATE_ADMIN = os.path.join(
    os.path.dirname(__file__), "..", "docs",
    "АКТ уничтожения административных дел.docx")


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def make_records():
    return [
        {"sequential": 1,
         "title": "Дело об административном правонарушении №5-1/2020 "
                  "в отношении Суркова Эдуарда Владимировича "
                  "по ст. 15.6 ч. 1 КоАП РФ",
         "dates": "10.01.2020\n10.03.2020",
         "opis": "", "unit": "", "count": 1,
         "retention": "2 года Ст. 369", "note": "в работе"},
        {"sequential": 2, "title": "", "dates": "", "opis": "", "unit": "",
         "count": "", "retention": "", "note": ""},
        {"sequential": 3,
         "title": "Дело об административном правонарушении №5-3/2020 "
                  "в отношении ООО Ромашка по 1.1 ч.1 ЗСО №104",
         "dates": "20.02.2020\n20.04.2020",
         "opis": "", "unit": "", "count": 1,
         "retention": "2 года Ст. 369", "note": ""},
    ]


def make_common_values():
    return {
        "судебный_участок": "9",
        "судья": "О. А. Левошина",
        "дата_утверждения": "«___» ______________ 2026 года",
        "дата_акта": "«__» ____________ 2026 г.",
        "номер_акта": "1",
        "секретарь": "А.М. Намаюшка",
        "дата_подписи": "«__» ____________ 2026 г.",
        "протокол_эк_дата": "«__» ____________ 2026 г.",
        "протокол_эк_номер": "1",
        "год_дел": "2020",
    }


def doc_full_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    parts.append(p.text)
    return "\n".join(parts)


def main():
    check("admin-шаблон существует", os.path.exists(TEMPLATE_ADMIN), True)

    print("[1] Формирование admin-акта")
    tmp_dir = tempfile.mkdtemp(prefix="admin_word_")
    out_path = os.path.join(tmp_dir, "акт_адм.docx")
    write_word_result(make_records(), TEMPLATE_ADMIN, out_path,
                      make_common_values(), case_type="admin")
    check("Файл создан", os.path.exists(out_path), True)

    doc = Document(out_path)
    text = doc_full_text(doc)

    remaining = re.findall(r"\{\{[^}]+\}\}", text)
    check("Нет оставшихся плейсхолдеров {{...}}", remaining, [])

    check("Итоговая строка (admin)",
          "Итого 2 (два) административных дел за 2020 год." in text, True)
    check("{{количество_прописью}} подставлено",
          "{{количество_прописью}}" not in text, True)
    check("Число прописью присутствует", "два" in text, True)
    check("Судья подставлен", "О. А. Левошина" in text, True)
    check("Судебный участок", "Судебный участок № 9" in text, True)

    # Таблица данных: нумерация + 2 дела (пустая запись исключена)
    tables_with_data = [
        t for t in doc.tables
        if any("Дело об административном правонарушении" in (c.text or "")
               for r in t.rows for c in r.cells)
    ]
    check("Найдена таблица со строками дел", len(tables_with_data), 1)
    table = tables_with_data[0]
    row_texts = [" ".join(c.text for c in r.cells) for r in table.rows]
    check("Строка №5-1 присутствует",
          any("№5-1/2020" in t for t in row_texts), True)
    check("Строка №5-3 присутствует",
          any("№5-3/2020" in t for t in row_texts), True)
    check("Даты дела в таблице",
          any("10.01.2020" in t for t in row_texts), True)

    print("[2] Формулировка {{итого}} для civil (регрессия)")
    tmp_dir2 = tempfile.mkdtemp(prefix="civil_word_")
    out2 = os.path.join(tmp_dir2, "акт_гр.docx")
    civil_tpl = os.path.join(os.path.dirname(__file__), "..", "docs",
                             "АКТ уничтожения гражданских дел.docx")
    civil_records = [{
        "sequential": 1,
        "title": "Гражданское дело №2-1/2020 о защите прав",
        "dates": "08.12.2017\n08.01.2018",
        "opis": 1, "unit": 1, "count": 1,
        "retention": "3 года Ст. 208", "note": "",
    }]
    write_word_result(civil_records, civil_tpl, out2, make_common_values(),
                      case_type="civil")
    text2 = doc_full_text(Document(out2))
    check("Итоговая строка (civil)",
          "Итого 1 (один) гражданских дел за 2020 год." in text2, True)

    print()
    print(f"Итого: {PASSED} OK, {FAILED} FAIL")
    return 1 if FAILED else 0


if __name__ == "__main__":
    sys.exit(main())
