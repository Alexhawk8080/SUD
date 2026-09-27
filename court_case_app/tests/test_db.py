# -*- coding: utf-8 -*-
"""
Тесты базы данных SQLite (Этап 8 + доработки: объединение организаций
и исключений, импорт/экспорт из Excel).

Запуск: .venv\\Scripts\\python.exe court_case_app\\tests\\test_db.py
"""

import sys
import os
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import Workbook
from database import db
from core.retention import get_retention_info
from core.case_processor import get_case_type_prefix, process_names_list

FAILED = 0
PASSED = 0


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def make_import_file(path):
    """Excel-файл в формате листа «База данных» пользователя."""
    wb = Workbook()
    ws = wb.active
    ws.title = "База данных"
    ws.append(['"Связной', "по заявлению"])
    ws.append(["АО", "по заявлению"])
    ws.append(["СБЕРБАНК", "по иску"])
    ws.append([None, None])
    ws.append(["ПАО", "по иску"])
    wb.save(path)


def main():
    tmp_dir = tempfile.mkdtemp(prefix="db_")
    db_path = os.path.join(tmp_dir, "app.db")

    print("[1] Создание БД и начальное наполнение")
    conn = db.init_db(db_path)
    orgs = db.list_organizations(conn)
    check("Организации заполнены (>10)", len(orgs) > 10, True)
    check("ООО -> по иску",
          any(o["name"] == "ООО" and o["claim_type"] == "по иску" for o in orgs),
          True)
    names = db.organization_names(conn)
    check("Список слов организаций содержит ООО", "ООО" in names, True)
    check("Нет таблицы exclusions",
          conn.execute("SELECT name FROM sqlite_master WHERE name='exclusions'"
                       ).fetchone() is None, True)
    rules = db.list_retention_rules(conn)
    check("Правила сроков заполнены (>20)", len(rules) > 20, True)
    keywords, texts = db.load_retention_rules(conn)
    check("Порядок ключей сохранён",
          [k for k, _ in keywords].index("пенсион") <
          [k for k, _ in keywords].index("пенси"), True)
    settings = db.get_settings(conn)
    # stage16a1a_v2: после перехода на схему v2 реквизитные ключи
    # (судья, секретарь, ...) в settings не хранятся — они уехали
    # в court_years. Проверяем служебный ключ.
    check("Настройки заполнены (auto_fix_categories)",
          settings.get("auto_fix_categories"), "1")
    check("schema_version = v3",
          settings.get("schema_version"), "v3")

    print("\n[2] CRUD: организации")
    new_id = db.add_organization(conn, "РАЙПО", "по заявлению")
    check("Добавлена организация",
          any(o["name"] == "РАЙПО" for o in db.list_organizations(conn)), True)
    db.update_organization(conn, new_id, "РАЙПО СОЮЗ", "по заявлению")
    check("Организация изменена",
          any(o["name"] == "РАЙПО СОЮЗ" for o in db.list_organizations(conn)),
          True)
    db.delete_organization(conn, new_id)
    check("Организация удалена",
          all(o["name"] != "РАЙПО СОЮЗ" for o in db.list_organizations(conn)),
          True)

    print("\n[3] CRUD: правила сроков")
    rid = db.add_retention_rule(conn, "тесткорень", "test_code", "3 года Ст. 227")
    check("Правило добавлено",
          any(r["keyword"] == "тесткорень" for r in db.list_retention_rules(conn)),
          True)
    db.delete_retention_rule(conn, rid)
    check("Правило удалено",
          all(r["keyword"] != "тесткорень" for r in db.list_retention_rules(conn)),
          True)

    print("\n[4] Сохранение после перезапуска и миграция")
    db.add_organization(conn, "ХРАНИТЬ", "по иску")
    conn.close()
    conn2 = db.init_db(db_path)
    check("Организация сохранилась",
          any(o["name"] == "ХРАНИТЬ" for o in db.list_organizations(conn2)), True)
    check("Наполнение не задваивается",
          sum(1 for o in db.list_organizations(conn2) if o["name"] == "ООО") == 1,
          True)

    print("\n[5] Влияние БД на обработку")
    orgs_dict = db.organizations_dict(conn2)
    excl_list = db.organization_names(conn2)
    check("Префикс: физлицо -> по заявлению",
          get_case_type_prefix("Иванов Иван Иванович", orgs_dict, excl_list),
          "по заявлению")
    check("Префикс: АО -> по иску (в списке)",
          get_case_type_prefix('АО "МАКС"', orgs_dict, excl_list),
          "по иску")
    check("Организация не склоняется (единый список)",
          process_names_list("ООО Ромашка; Иванов Иван", "Dative", excl_list),
          "ООО Ромашка; Иванову Ивану")
    keywords_bd, texts_bd = db.load_retention_rules(conn2)
    check("Retention по правилам БД (развод)",
          get_retention_info("Исковое заявление о расторжении брака супругов",
                             keywords_bd, texts_bd),
          "3 года ЭК Ст. 137")

    print("\n[6] Импорт из Excel (формат файла пользователя)")
    import_path = os.path.join(tmp_dir, "import.xlsx")
    make_import_file(import_path)
    added = db.import_database_from_excel(conn2, import_path)
    check("Импортировано новых: 2 (дубликаты АО/ПАО пропущены)", added, 2)
    orgs3 = db.organizations_dict(conn2)
    check("СБЕРБАНК импортирован (по иску)", orgs3.get("СБЕРБАНК"), "по иску")
    check('"Связной импортирован (по заявлению)',
          orgs3.get('"Связной'), "по заявлению")
    check("Дубликат АО не перезаписан (остался по иску)",
          orgs3.get("АО"), "по иску")

    print("\n[7] Экспорт базы в Excel")
    export_path = os.path.join(tmp_dir, "export.xlsx")
    db.export_database_to_excel(conn2, export_path)
    check("Файл экспорта создан", os.path.exists(export_path), True)
    from openpyxl import load_workbook
    wb = load_workbook(export_path)
    check("Лист «База данных» есть", "База данных" in wb.sheetnames, True)
    check("Лист «Правила сроков» есть", "Правила сроков" in wb.sheetnames, True)
    check("Лист «Настройки» есть", "Настройки" in wb.sheetnames, True)
    check("Лист «Категории дел» есть", "Категории дел" in wb.sheetnames, True)
    wb.close()

    print("\n[8] Настройки")
    db.set_setting(conn2, "номер_акта", "7")
    check("Настройка сохранена", db.get_settings(conn2).get("номер_акта"), "7")
    check("Автоисправление включено по умолчанию",
          db.get_setting(conn2, "auto_fix_categories", "0"), "1")
    db.set_setting(conn2, "auto_fix_categories", "0")
    check("Автоисправление отключается",
          db.get_setting(conn2, "auto_fix_categories", "1"), "0")

    print("\n[9] Словарь «истец -> категория» (case_categories)")
    cid = db.add_case_category(
        conn2, "сбербанк", 'ООО "Сбербанк"', "о взыскании задолженности по кредитному договору", True)
    check("Запись добавлена",
          len(db.list_case_categories(conn2)) == 1, True)
    dup_id = db.add_case_category(
        conn2, "сбербанк", 'ООО "Сбербанк"', "о взыскании задолженности по кредитному договору", True)
    check("Дубликат пары (ключ, категория) игнорируется",
          (dup_id == cid and len(db.list_case_categories(conn2)) == 1), True)
    db.add_case_category(
        conn2, "сбербанк", 'ООО "Сбербанк"', "о взыскании процентов по кредиту", True)
    db.add_case_category(
        conn2, "иванов иван иванович", "Иванов Иван Иванович",
        "о взыскании сумм по договору займа", False)
    check("Несколько категорий на истца",
          len(db.list_case_categories(conn2)) == 3, True)
    cmap = db.case_categories_map(conn2)
    check("Карта: сбербанк -> 2 категории",
          len(cmap.get("сбербанк", [])), 2)
    check("Карта: физлицо -> 1 категория",
          len(cmap.get("иванов иван иванович", [])), 1)
    labels = db.case_categories_labels(conn2)
    check("Метка истца сохраняется",
          labels.get("сбербанк"), 'ООО "Сбербанк"')
    rec = db.list_case_categories(conn2)[0]
    db.update_case_category(conn2, rec["id"], "о взыскании суммы долга")
    check("Категория изменена",
          db.list_case_categories(conn2)[0]["category"], "о взыскании суммы долга")
    db.delete_case_category(conn2, rec["id"])
    check("Запись удалена", len(db.list_case_categories(conn2)), 2)
    conn2.close()

    # Перезапуск: данные словаря сохраняются
    conn3 = db.init_db(db_path)
    check("Словарь сохранился после перезапуска",
          len(db.list_case_categories(conn3)), 2)
    conn3.close()

    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()