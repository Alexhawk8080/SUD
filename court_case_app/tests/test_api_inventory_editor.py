# stage_49c
"""Smoke-тесты blueprint ``inventory_editor``."""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent.parent
APP_DIR = ROOT / "court_case_app"
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))


@pytest.fixture
def client(tmp_path, monkeypatch):
    import app as appmod
    db_file = str(tmp_path / "test.db")
    monkeypatch.setattr(appmod, "DB_PATH", db_file)
    appmod.app.config["DB_PATH"] = db_file
    appmod.app.config["TESTING"] = True
    with appmod.app.test_client() as c:
        yield c, appmod


def _make_inventory(appmod) -> int:
    """Создаёт area+year+result+inventory, возвращает inventory_id."""
    from database import db as db_mod
    conn = db_mod.init_db(appmod.app.config["DB_PATH"])
    try:
        area_id = db_mod.add_court_area(conn, {"номер": "9"})
        year_id = db_mod.add_court_year(conn, area_id, 2019, {
            "судья": "Иванов", "секретарь": "Петрова",
            "дата_утверждения": "2020-01-15", "дата_акта": "2020-01-10",
            "номер_акта": "1", "дата_подписи": "2020-01-16",
            "протокол_эк_дата": "2020-01-09",
            "протокол_эк_номер": "5",
        })
        conn.execute(
            "INSERT INTO processing_results "
            "(court_year_id, case_type, record_count) VALUES (?, ?, ?)",
            (year_id, "civil", 0))
        result_id = conn.execute(
            "SELECT last_insert_rowid()").fetchone()[0]
        conn.execute(
            "INSERT INTO processing_inventories (result_id) VALUES (?)",
            (result_id,))
        inv_id = conn.execute(
            "SELECT last_insert_rowid()").fetchone()[0]
        conn.commit()
        return inv_id
    finally:
        conn.close()


def test_page_200(client):
    c, appmod = client
    inv_id = _make_inventory(appmod)
    r = c.get("/inventory_editor/%d" % inv_id)
    assert r.status_code == 200
    assert "Внутренняя опись".encode("utf-8") in r.data


def test_doc_types(client):
    c, _appmod = client
    r = c.get("/api/inventory_editor/doc-types")
    assert r.status_code == 200
    data = r.get_json()
    assert len(data) == 7
    assert {d["kind"] for d in data} == {"civil", "admin"}


def test_add_case_and_list(client):
    c, appmod = client
    inv_id = _make_inventory(appmod)
    r = c.post("/api/inventory_editor/%d/rows" % inv_id, json={
        "prefix": "2", "case_number": 100, "year": 2019,
        "docs": [{"title": "Решение по гражданскому делу",
                  "pages_count": 5}],
    })
    assert r.status_code == 200, r.get_json()
    data = r.get_json()
    assert data["total"] == 1
    row = data["rows"][0]
    assert row["case_number"] == 100
    assert row["E"] == 100        # первая строка: prev_case = 0
    assert row["I"] == 1
    assert row["J"] == 5
    assert row["A"] == 1 and row["B"] == 1

    r2 = c.get("/api/inventory_editor/%d/rows" % inv_id)
    assert r2.get_json()["total"] == 1


def test_delete_case(client):
    c, appmod = client
    inv_id = _make_inventory(appmod)
    c.post("/api/inventory_editor/%d/rows" % inv_id, json={
        "prefix": "2", "case_number": 100, "year": 2019,
        "docs": [{"title": "Решение по гражданскому делу"}],
    })
    r = c.delete("/api/inventory_editor/%d/cases/2019/100" % inv_id)
    assert r.status_code == 200
    payload = r.get_json()
    assert payload["removed"] == 1
    assert payload["total"] == 0


def test_import_xlsx(client):
    from openpyxl import Workbook
    c, appmod = client
    inv_id = _make_inventory(appmod)

    wb = Workbook()
    ws = wb.active
    ws.title = "Таблица"
    ws.cell(4, 3).value = "Решение по гражданскому делу"
    ws.cell(4, 4).value = "2"
    ws.cell(4, 5).value = 55
    ws.cell(4, 7).value = 2019
    ws.cell(4, 8).value = 3
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    r = c.post(
        "/api/inventory_editor/%d/import" % inv_id,
        data={"file": (buf, "test.xlsx")},
        content_type="multipart/form-data",
    )
    assert r.status_code == 200, r.get_json()
    payload = r.get_json()
    assert payload["total"] == 1
    assert payload["rows"][0]["case_number"] == 55
    assert payload["rows"][0]["J"] == 3


def test_export_xlsx(client):
    from openpyxl import Workbook
    c, appmod = client
    inv_id = _make_inventory(appmod)
    c.post("/api/inventory_editor/%d/rows" % inv_id, json={
        "prefix": "2", "case_number": 100, "year": 2019,
        "docs": [{"title": "Решение по гражданскому делу",
                  "pages_count": 1}],
    })

    wb = Workbook()
    ws = wb.active
    ws.title = "Таблица"
    ws.cell(100, 7).value = 0        # растягиваем шаблон до строки 100
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    r = c.post(
        "/api/inventory_editor/%d/export" % inv_id,
        data={"template": (buf, "tmpl.xlsx")},
        content_type="multipart/form-data",
    )
    assert r.status_code == 200, (r.status_code, r.data[:200])
    assert r.data[:2] == b"PK"       # xlsx = zip
