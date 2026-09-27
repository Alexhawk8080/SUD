# -*- coding: utf-8 -*-
"""
Сквозной тест обработки административных дел (D8e).

Проверяет полный цикл: admin-файл -> БД (cases) -> process_year(admin) ->
экспорт в Word по admin-шаблону, а также миграцию схемы v2 -> v3
(колонка court_years.admin_retention).

Запуск: .venv\\Scripts\\python.exe court_case_app\\tests\\test_admin_pipeline.py
"""

import os
import sqlite3
import sys
import tempfile
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import Workbook
from docx import Document

from core.db_pipeline import (
    export_year_result_to_file,
    process_year,
    save_year_source_from_xlsx,
)
from database import db

FAILED = 0
PASSED = 0

TEMPLATE_CIVIL = os.path.join(os.path.dirname(__file__), "..", "docs",
                              "АКТ уничтожения гражданских дел.docx")
TEMPLATE_ADMIN = os.path.join(os.path.dirname(__file__), "..", "docs",
                              "АКТ уничтожения административных дел.docx")


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def year_data(s="1"):
    return {"судья": f"Судья{s}", "секретарь": f"Сек{s}",
            "дата_утверждения": "—", "дата_акта": "—",
            "номер_акта": s, "дата_подписи": "—",
            "протокол_эк_дата": "—", "протокол_эк_номер": s}


def make_admin_xlsx(path):
    """admin-файл: шапка в строке 3, 3 дела за 2020 + 1 за 2021."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Адм"
    ws.append(["Отчёт", None, None, None, None])
    ws.append([None, None, None, None, None])
    ws.append(["Вх. №", "Дата поступления", "№ Дела в производстве",
               "ФИО/наименование лица", "Статья", None, None,
               "Текущая задача", None, None, "Окончание производства",
               None, "Состояние акта"])
    ws.append([101, datetime(2020, 1, 10), "5-1/2020",
               "Сурков Эдуард Владимирович", "ст. 15.6 ч. 1", None, None,
               "в работе", None, None, None, None, "не сформирован"])
    ws.append([102, datetime(2020, 2, 20), "5-2/2020",
               "ООО Ромашка", "1.1 ч.1 ЗСО №104", None, None,
               "завершено", None, None, None, None, "сформирован"])
    ws.append([103, datetime(2020, 3, 15), "5-3/2020",
               "Петров Пётр Петрович", "ст. 12.8 ч. 1", None, None,
               "", None, None, None, None, ""])
    ws.append([104, datetime(2021, 1, 5), "5-1/2021",
               "Иванов Иван Иванович", "ст. 1", None, None,
               "", None, None, None, None, ""])
    wb.save(path)


def docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def main():
    tmp = tempfile.mkdtemp(prefix="admin_pipe_")
    db_path = os.path.join(tmp, "app.db")
    out_dir = os.path.join(tmp, "out")
    os.makedirs(out_dir, exist_ok=True)

    conn = db.init_db(db_path)
    aid = db.add_court_area(conn, {"номер": "9"})
    yid = db.add_court_year(conn, aid, 2020, year_data("20"))

    print("[1] Миграция схемы v3: admin_retention")
    check("SCHEMA_VERSION = v3", db.SCHEMA_VERSION, "v3")
    cy = db.get_court_year(conn, yid)
    check("admin_retention по умолчанию",
          cy["admin_retention"], db.DEFAULT_ADMIN_RETENTION)
    cols = [r[1] for r in conn.execute(
        "PRAGMA table_info(court_years)").fetchall()]
    check("колонка admin_retention есть", "admin_retention" in cols, True)

    print("[2] save_year_source_from_xlsx (admin)")
    xlsx = os.path.join(tmp, "2020 адм.xlsx")
    make_admin_xlsx(xlsx)
    res = save_year_source_from_xlsx(
        conn, yid, xlsx, "2020 адм.xlsx", case_type="admin")
    check("source_file_id > 0", res["source_file_id"] > 0, True)
    check("row_count = 4", res["row_count"], 4)
    check("valid_count = 3 (год 2020)", res["valid_count"], 3)
    check("alimony_count = 0", res["alimony_count"], 0)

    print("[3] Дела в БД (переиспользование колонок)")
    cases = db.list_cases(conn, source_file_id=res["source_file_id"],
                          case_type="admin")
    check("4 дела сохранено", len(cases), 4)
    c0 = cases[0]
    check("case_type = admin", c0["case_type"], "admin")
    check("applicants = person", c0["applicants"], "Сурков Эдуард Владимирович")
    check("category = article (нормализована)",
          c0["category"], "ст. 15.6 ч. 1 КоАП РФ")
    check("opis_number = in_number", c0["opis_number"], "101")
    check("note = task", c0["note"], "в работе")
    check("skip_reason = act_state", c0["skip_reason"], "не сформирован")
    # дело 2021 невалидно для года 2020
    c3 = cases[3]
    check("дело 2021 невалидно", c3["is_valid"], 0)

    print("[4] process_year (admin)")
    result = process_year(conn, yid, case_type="admin", court_area_id=aid)
    check("result_id > 0", result["result_id"] > 0, True)
    check("record_count = 3", result["record_count"], 3)
    check("target_year = 2020", result["target_year"], 2020)
    check("case_prefix = 5", result["stats"]["case_prefix"], "5")

    res_rec = db.get_processing_result(conn, result["result_id"])
    check("case_type = admin", res_rec["case_type"], "admin")

    print("[5] Строки результата")
    rows = db.list_result_rows(conn, result["result_id"])
    check("3 строки", len(rows), 3)
    check("заголовок admin",
          rows[0]["title"].startswith(
              "Дело об административном правонарушении №5-1/2020"), True)
    check("склонение физлица в заголовке",
          "Суркова Эдуарда Владимировича" in rows[0]["title"], True)
    check("организация как есть",
          "ООО Ромашка" in rows[1]["title"], True)
    check("retention из admin_retention",
          rows[0]["retention"], db.DEFAULT_ADMIN_RETENTION)
    check("даты +2 месяца",
          rows[0]["dates"], "10.01.2020\n10.03.2020")

    print("[6] Экспорт в Word (admin-шаблон)")
    if not os.path.exists(TEMPLATE_ADMIN):
        print("  SKIP admin-шаблон не найден:", TEMPLATE_ADMIN)
    else:
        exp = export_year_result_to_file(
            conn, result["result_id"],
            template_path=TEMPLATE_CIVIL,
            template_path_admin=TEMPLATE_ADMIN,
            output_dir=out_dir, output_format="word")
        check("Файл создан", os.path.exists(exp["path"]), True)
        check("filename .docx", exp["filename"].endswith(".docx"), True)
        text = docx_text(exp["path"])
        check("Заголовок 5-1 в акте", "№5-1/2020" in text, True)
        check("Итог admin",
              "Итого 3 (три) административных дел за 2020 год." in text, True)
        check("Нет оставшихся плейсхолдеров",
              "{{" not in text, True)

    print("[7] Изменение admin_retention влияет на результат")
    db.update_court_year(conn, yid, dict(year_data("20"),
                                         admin_retention="5 лет Ст. 999"))
    result2 = process_year(conn, yid, case_type="admin", court_area_id=aid)
    rows2 = db.list_result_rows(conn, result2["result_id"])
    check("новый retention применён",
          rows2[0]["retention"], "5 лет Ст. 999")

    print("[8] Миграция v2 -> v3 на «старой» БД")
    old_path = os.path.join(tmp, "old.db")
    _make_v2_db(old_path)
    conn2 = db.init_db(old_path)
    cols2 = [r[1] for r in conn2.execute(
        "PRAGMA table_info(court_years)").fetchall()]
    check("admin_retention добавлена миграцией",
          "admin_retention" in cols2, True)
    ver = conn2.execute(
        "SELECT value FROM settings WHERE key='schema_version'").fetchone()
    check("версия стала v3", ver[0] if ver else None, "v3")

    print()
    print(f"Итого: {PASSED} OK, {FAILED} FAIL")
    return 1 if FAILED else 0


def _make_v2_db(path):
    """Создаёт минимальную БД со схемой v2 (без admin_retention)."""
    conn = sqlite3.connect(path)
    conn.executescript("""
        CREATE TABLE settings (key TEXT PRIMARY KEY, value TEXT);
        INSERT INTO settings (key, value) VALUES ('schema_version', 'v2');
        CREATE TABLE court_areas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            номер TEXT, судебный_участок TEXT, судья TEXT, секретарь TEXT,
            дата_утверждения TEXT, дата_акта TEXT, номер_акта TEXT,
            дата_подписи TEXT, протокол_эк_дата TEXT, протокол_эк_номер TEXT
        );
        CREATE TABLE court_years (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            court_area_id INTEGER, year INTEGER,
            судья TEXT, секретарь TEXT, дата_утверждения TEXT,
            дата_акта TEXT, номер_акта TEXT, дата_подписи TEXT,
            протокол_эк_дата TEXT, протокол_эк_номер TEXT,
            note TEXT, is_incomplete INTEGER DEFAULT 1,
            is_closed INTEGER DEFAULT 0,
            created_at TEXT, updated_at TEXT
        );
    """)
    conn.commit()
    conn.close()


if __name__ == "__main__":
    sys.exit(main())
