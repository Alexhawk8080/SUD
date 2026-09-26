# -*- coding: utf-8 -*-
"""
Тесты чтения docx-описи «Внутренняя опись» (этап 14a).

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tests\\test_inventory_docx_reader.py
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document

from core.inventory.docx_reader import (
    clean_name,
    extract_rows,
    find_inventory_tables,
    make_record,
    parse_number,
    read_inventory_docx,
)
from core.inventory.exceptions import (
    BadInventoryFileError,
    NoInventoryTableError,
)

FAILED = 0
PASSED = 0


def check(name, actual, expected):
    global FAILED, PASSED
    if actual == expected:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n"
              f"       получено:  {actual!r}")


def expect_error(name, fn, exc_type):
    global FAILED, PASSED
    try:
        fn()
    except exc_type:
        PASSED += 1
        print(f"  OK   {name}")
        return
    except Exception as exc:  # noqa: BLE001
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось {exc_type.__name__}, "
              f"получено {type(exc).__name__}: {exc}")
        return
    FAILED += 1
    print(f"  FAIL {name}\n       исключение не поднято")


HEADERS = ["№ п/п", "Наименование", "Номера листов"]


def make_docx(path, volumes, footer_total=None):
    """
    volumes — список списков строк: каждая строка = (номер, наименование,
    номера листов). Между томами добавляется заголовок «Том N».
    """
    doc = Document()
    doc.add_paragraph("Внутренняя опись дел")
    for v_no, rows in enumerate(volumes, start=1):
        doc.add_paragraph(f"Том {v_no}")
        table = doc.add_table(rows=1, cols=3)
        for i, h in enumerate(HEADERS):
            table.cell(0, i).text = h
        for num, name, sheets in rows:
            row = table.add_row()
            row.cells[0].text = str(num)
            row.cells[1].text = str(name)
            row.cells[2].text = str(sheets)
    if footer_total is not None:
        doc.add_paragraph(f"Итого наряд на {footer_total} листах")
    doc.save(path)
    return path


def main():
    tmp = tempfile.mkdtemp(prefix="inventory_")

    print("[1] clean_name: артефакты Word")
    check("обычный текст", clean_name("Решение по делу"),
          "Решение по делу")
    check("{.mark}", clean_name("Решение {.mark} по делу"),
          "Решение по делу")
    check("<u> и </u>", clean_name("<u>Определение</u> о возврате"),
          "Определение о возврате")
    check("перенос строки", clean_name("Решение\nпо делу"),
          "Решение по делу")
    check("лишние пробелы", clean_name("  Решение   по   делу  "),
          "Решение по делу")
    check("None", clean_name(None), "")

    print("\n[2] parse_number")
    check("171", parse_number("171"), (171, "171"))
    check("231.", parse_number("231."), (231, "231."))
    check(" 12 ", parse_number(" 12 "), (12, "12"))
    check("пусто", parse_number(""), (None, ""))
    check("None", parse_number(None), (None, ""))
    check("текст", parse_number("прим."), (None, "прим."))

    print("\n[3] find_inventory_tables")
    p1 = make_docx(os.path.join(tmp, "one.docx"),
                   [[(1, "Решение по делу", "8")]])
    doc1 = Document(p1)
    check("одна таблица описи", len(find_inventory_tables(doc1)), 1)
    check("индекс таблицы", find_inventory_tables(doc1)[0][0], 0)

    print("\n[4] extract_rows: 3 столбца")
    table = Document(p1).tables[0]
    recs, att, info = extract_rows(table, file_name="one.docx")
    check("строк данных", len(recs), 1)
    check("номер", recs[0]["number"], 1)
    check("наименование", recs[0]["name"], "Решение по делу")
    check("номера листов", recs[0]["sheets_raw"], "8")
    check("том", recs[0]["volume"], 1)
    check("без attention", att, [])
    check("start_row", info["start_row"], 1)

    print("\n[5] make_record: структура")
    r = make_record("f.docx", 7, 2, 5, "5", "имя", "имя", "3-4")
    check("set ключей", set(r),
          {"file", "word_row", "volume", "number", "number_raw",
           "name", "name_raw", "sheets_raw"})
    check("word_row", r["word_row"], 7)
    check("volume", r["volume"], 2)

    print("\n[6] read_inventory_docx: один файл, футер")
    res = read_inventory_docx(p1)
    check("файл", res["file"], "one.docx")
    check("tables", res["tables"], 1)
    check("footer_total", res["footer_total"], None)
    check("строк", len(res["rows"]), 1)
    check("volume-инфо", len(res["volumes"]), 1)
    check("title тома", res["volumes"][0]["titles"], ["Том 1"])

    print("\n[7] несколько томов в одном файле")
    p2 = make_docx(
        os.path.join(tmp, "two.docx"),
        [[(1, "A", "1-2"), (2, "B", "3")],
         [(3, "C", "4")]],
        footer_total=25)
    res2 = read_inventory_docx(p2)
    check("2 тома", res2["tables"], 2)
    check("3 строки", len(res2["rows"]), 3)
    check("footer_total", res2["footer_total"], 25)
    check("volume строки 1", res2["rows"][0]["volume"], 1)
    check("volume строки 3", res2["rows"][2]["volume"], 2)
    check("сквозной word_row (растёт)",
          res2["rows"][2]["word_row"] > res2["rows"][1]["word_row"], True)

    print("\n[8] пропуск и дубликат № п/п -> attention")
    p3 = make_docx(
        os.path.join(tmp, "gaps.docx"),
        [[(1, "A", "1"), (3, "B", "2"), (3, "C", "3")]])
    res3 = read_inventory_docx(p3)
    reasons = [a["reason"] for a in res3["attention"]]
    check("пропуск", any("пропуск" in x for x in reasons), True)
    check("дубликат", any("дубликат" in x for x in reasons), True)
    check("строки сохранены", len(res3["rows"]), 3)

    print("\n[9] ошибки: битый файл и отсутствие таблицы")
    bad = os.path.join(tmp, "bad.docx")
    with open(bad, "wb") as f:
        f.write(b"not a docx")
    expect_error("битый файл", lambda: read_inventory_docx(bad),
                 BadInventoryFileError)

    empty = Document()
    empty.add_paragraph("Просто текст без таблиц")
    p4 = os.path.join(tmp, "empty.docx")
    empty.save(p4)
    expect_error("нет таблицы", lambda: read_inventory_docx(p4),
                 NoInventoryTableError)

    print("\n" + "=" * 60)
    print(f"ИТОГО: PASSED = {PASSED}, FAILED = {FAILED}")
    return 1 if FAILED else 0


if __name__ == "__main__":
    sys.exit(main())
