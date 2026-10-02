# stage_49b
import openpyxl
from court_case_app.core.inventory_editor.importer import import_from_excel


def _mk_sheet(tmp_path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Таблица"
    # Служебные строки 1-3 — пустые (или шапка).
    for i, row in enumerate(rows):
        r = 4 + i
        ws.cell(r, 3).value = row.get("title")
        ws.cell(r, 4).value = row.get("prefix")
        ws.cell(r, 5).value = row.get("E")
        ws.cell(r, 7).value = row.get("year")
        ws.cell(r, 8).value = row.get("pages")
    p = tmp_path / "in.xlsx"
    wb.save(p)
    return str(p)


def test_basic_three_rows(tmp_path):
    path = _mk_sheet(tmp_path, [
        {"title": "Приказ", "prefix": "2", "E": 105, "year": 2019, "pages": 1},
        {"title": "Определение", "prefix": "2", "E": 0, "year": 2019, "pages": 2},
        {"title": "Решение", "prefix": "2", "E": 1, "year": 2019, "pages": 1},
    ])
    docs = import_from_excel(path)
    assert [d.case_number for d in docs] == [105, 105, 106]
    assert docs[1].pages_count == 2
    assert docs[2].title == "Решение"


def test_stops_at_empty_C(tmp_path):
    path = _mk_sheet(tmp_path, [
        {"title": "A", "prefix": "2", "E": 1, "year": 2019, "pages": 1},
        {"title": "B", "prefix": "2", "E": 1, "year": 2019, "pages": 1},
        {"title": None, "prefix": "2", "E": 1, "year": 2019, "pages": 1},
        {"title": "C", "prefix": "2", "E": 1, "year": 2019, "pages": 1},
    ])
    docs = import_from_excel(path)
    assert len(docs) == 2


def test_skips_formulas_in_C(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Таблица"
    ws.cell(4, 3).value = "=IF(1,2,3)"   # формула — значит, конец данных
    p = tmp_path / "f.xlsx"
    wb.save(p)
    docs = import_from_excel(str(p))
    assert docs == []
