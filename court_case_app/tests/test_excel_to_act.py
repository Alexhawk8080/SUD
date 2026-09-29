# stage_46
# -*- coding: utf-8 -*-
"""Тесты переноса Excel → Word-акт (stage_46)."""
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document
from openpyxl import Workbook

from core.converter.excel_to_act import import_excel_to_act

FAILED = 0
PASSED = 0


def check(name, actual, expected):
    global FAILED, PASSED
    if actual == expected:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def make_excel(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Результат обработки"
    ws.append(["№ п/п", "Заголовок дела", "Даты дела", "№ описи",
               "№ ед.хр.", "Кол-во", "Срок хранения", "Примечание"])
    for r in rows:
        ws.append(r)
    wb.save(path)


def make_word(path, data_rows):
    doc = Document()
    doc.add_paragraph("Шапка акта.")
    table = doc.add_table(rows=1, cols=8)
    for i in range(8):
        table.cell(0, i).text = str(i + 1)
    for row_vals in data_rows:
        row = table.add_row()
        for i, v in enumerate(row_vals):
            row.cells[i].text = v
    doc.save(path)


def read_word_table(path):
    doc = Document(path)
    for tbl in doc.tables:
        if len(tbl.columns) >= 8 and tbl.cell(0, 0).text.strip() == "1":
            return tbl
    return None


def main():
    tmp = Path(tempfile.mkdtemp(prefix="st46_"))
    xlsx = tmp / "res.xlsx"
    docx = tmp / "act.docx"

    rows = [
        ["1", "Гражданское дело №2-1/2019 по заявлению Иванова к Петрову о взыскании",
         "01.01.2019\n02.01.2019", "1", "1", "1", "3 года Ст. 227", ""],
        ["2", "Гражданское дело №2-2/2019 по заявлению Сидорова к Кузнецову о расторжении",
         "03.01.2019\n04.01.2019", "1", "2", "1", "3 года ЭК Ст. 137", ""],
        ["3", "Гражданское дело №2-3/2019 по иску ООО «Ромашка» к Федорову",
         "05.01.2019\n06.01.2019", "1", "3", "1", "3 года Ст. 227", ""],
    ]
    make_excel(xlsx, rows)

    existing = [
        ["50", "СТАРАЯ ЗАПИСЬ 1", "старые даты 1", "x", "x", "x", "x", ""],
        ["51", "СТАРАЯ ЗАПИСЬ 2", "старые даты 2", "x", "x", "x", "x", ""],
    ]
    make_word(docx, existing)

    print("[1] Базовый вызов")
    info = import_excel_to_act(str(xlsx), str(docx))
    check("rows_written = 3", info["rows_written"], 3)
    check("sheet_used", info["sheet_used"], "Результат обработки")
    check("файл создан", Path(info["output_path"]).exists(), True)
    check("суффикс _filled", Path(info["output_path"]).name, "act_filled.docx")

    print("\n[2] Исходный файл не изменён")
    orig = read_word_table(str(docx))
    check("в исходном всё ещё 3 строки (заголовок + 2 старых)",
          len(orig.rows), 3)
    check("в исходном старая строка на месте",
          "СТАРАЯ ЗАПИСЬ 1" in orig.cell(1, 1).text, True)

    print("\n[3] Выходной файл: 1 заголовок + 3 новые строки")
    out = read_word_table(info["output_path"])
    check("строк = 4", len(out.rows), 4)

    print("\n[4] Данные в столбцах 2..8")
    check("строка 1, столбец 2 = заголовок 1",
          "Гражданское дело №2-1/2019" in out.cell(1, 1).text, True)
    check("строка 2, столбец 2 = заголовок 2",
          "Гражданское дело №2-2/2019" in out.cell(2, 1).text, True)
    check("строка 3, столбец 2 = заголовок 3",
          "ООО «Ромашка»" in out.cell(3, 1).text, True)
    check("строка 1, столбец 7 = срок",
          out.cell(1, 6).text.strip(), "3 года Ст. 227")
    check("строка 2, столбец 7 = срок",
          out.cell(2, 6).text.strip(), "3 года ЭК Ст. 137")

    print("\n[5] Столбец 1 не заполнен текстом (для автонумерации)")
    check("строка 1, столбец 1 пуст",
          out.cell(1, 0).text.strip(), "")
    check("строка 2, столбец 1 пуст",
          out.cell(2, 0).text.strip(), "")
    check("строка 3, столбец 1 пуст",
          out.cell(3, 0).text.strip(), "")

    print("\n[6] Идемпотентность: повторный вызов даёт те же 4 строки")
    info2 = import_excel_to_act(str(xlsx), str(docx),
                                output_path=info["output_path"])
    out2 = read_word_table(info2["output_path"])
    check("строк = 4", len(out2.rows), 4)
    check("данные не задвоились",
          out2.cell(1, 1).text.count("№2-1/2019"), 1)

    print("\n[7] Ошибки при отсутствии листа")
    xlsx_bad = tmp / "bad.xlsx"
    wb = Workbook()
    wb.active.title = "Лист1"
    wb.save(xlsx_bad)
    try:
        import_excel_to_act(str(xlsx_bad), str(docx))
        check("ValueError при отсутствии листа", "нет ошибки", "ValueError")
    except ValueError as e:
        check("ValueError при отсутствии листа",
              "нет листа" in str(e), True)

    print("\n[8] Ошибки при отсутствии таблицы в Word")
    docx_bad = tmp / "bad.docx"
    bad = Document()
    bad.add_paragraph("пусто")
    bad.save(docx_bad)
    try:
        import_excel_to_act(str(xlsx), str(docx_bad))
        check("ValueError при отсутствии таблицы", "нет ошибки", "ValueError")
    except ValueError as e:
        check("ValueError при отсутствии таблицы",
              "не найдена таблица" in str(e), True)

    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()
