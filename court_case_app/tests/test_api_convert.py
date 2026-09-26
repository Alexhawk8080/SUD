# -*- coding: utf-8 -*-
"""
Тесты API конвертации Excel -> Word (stage12c, под-этап B3).

Использует app.test_client() с подменой DB_PATH/OUTPUT_DIR/UPLOAD_DIR.

Запуск:
    .venv\\Scripts\\python.exe court_case_app\\tests\\test_api_convert.py
"""

import io
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from openpyxl import Workbook
from docx import Document

import app as app_module

FAILED = 0
PASSED = 0


def check(name: str, actual, expected):
    global FAILED, PASSED
    ok = actual == expected
    if ok:
        PASSED += 1
        print(f"  OK   {name}")
    else:
        FAILED += 1
        print(f"  FAIL {name}\n       ожидалось: {expected!r}\n       получено:  {actual!r}")


def make_result_xlsx(path, sheet_name="Результат обработки", years=("2020",)):
    """Синтетический Excel-результат."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(["№ п/п", "Заголовок дела", "Даты дела", "№ описи",
               "№ ед.хр.", "Кол-во ед.хр.", "Срок хранения", "Примечание"])
    n = 1
    for y in years:
        for i in range(1, 3):
            ws.append([
                n,
                f"Гражданское дело №2-{i}/{y} по заявлению Тестова Т. Т. "
                f"к Тестову Т. Т. о расторжении брака",
                f"01.01.{y}\n01.02.{y}", 1, i, 1,
                "3 года ЭК Ст. 137", ""])
            n += 1
    wb.save(path)


def make_xlsx_without_numbers(path):
    """Excel-результат без номеров дел в заголовках."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Результат обработки"
    ws.append(["№ п/п", "Заголовок дела", "Даты дела", "№ описи",
               "№ ед.хр.", "Кол-во ед.хр.", "Срок хранения", "Примечание"])
    ws.append([1, "Просто заголовок без номера", "01.01.2020", 1, 1, 1, "", ""])
    ws.append([2, "Ещё один заголовок", "02.01.2020", 1, 2, 1, "", ""])
    wb.save(path)


def main():
    tmp = tempfile.mkdtemp(prefix="api_conv_")
    db_path = os.path.join(tmp, "app.db")
    out_dir = os.path.join(tmp, "out")
    up_dir = os.path.join(tmp, "up")
    os.makedirs(out_dir, exist_ok=True)
    os.makedirs(up_dir, exist_ok=True)

    app_module.DB_PATH = db_path
    app_module.OUTPUT_DIR = out_dir
    app_module.UPLOAD_DIR = up_dir

    client = app_module.app.test_client()

    xlsx_path = os.path.join(tmp, "результат.xlsx")
    make_result_xlsx(xlsx_path, years=("2019", "2020", "2020"))

    nonum_path = os.path.join(tmp, "без_номеров.xlsx")
    make_xlsx_without_numbers(nonum_path)

    print("[1] GET /convert_page")
    r = client.get("/convert_page")
    check("HTTP 200", r.status_code, 200)
    check("Есть id='convert-card'", 'id="convert-card"' in r.get_data(as_text=True), True)

    print("\n[2] POST /inspect_file: базовый сценарий")
    with open(xlsx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "результат.xlsx")}
        r = client.post("/inspect_file", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("sheets — список", isinstance(j.get("sheets"), list), True)
    check("Один лист", j.get("sheets"), ["Результат обработки"])
    check("sheet = 'Результат обработки'", j.get("sheet"), "Результат обработки")
    check("mapping содержит 'title'", "title" in (j.get("mapping") or {}), True)
    check("total = 6", j.get("total"), 6)
    check("year = 2020 (самый частый)", j.get("year"), "2020")
    check("years содержит 2019", "2019" in (j.get("years") or {}), True)
    check("others содержит 2019", len(j.get("others") or []), 1)
    check("warning непустой", bool(j.get("warning")), True)
    check("preview — 6 записей (все, <12)", len(j.get("preview") or []), 6)

    print("\n[3] POST /inspect_file с sheet='Результат обработки' -> sheets=None")
    with open(xlsx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "результат.xlsx"),
                "sheet": "Результат обработки"}
        r = client.post("/inspect_file", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("sheets = None", j.get("sheets"), None)
    check("sheet = 'Результат обработки'", j.get("sheet"), "Результат обработки")

    print("\n[4] POST /inspect_file с чужим расширением")
    with open(xlsx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "результат.docx")}
        r = client.post("/inspect_file", data=data,
                        content_type="multipart/form-data")
    check("HTTP 400", r.status_code, 400)
    check("Ошибка про .xlsx", ".xlsx" in (r.get_json() or {}).get("error", ""), True)

    print("\n[5] POST /inspect_file без номеров дел -> 400 (решение 27)")
    with open(nonum_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "без_номеров.xlsx")}
        r = client.post("/inspect_file", data=data,
                        content_type="multipart/form-data")
    check("HTTP 400", r.status_code, 400)
    err = (r.get_json() or {}).get("error", "")
    check("Сообщение про номера дел", "номер" in err.lower(), True)

    print("\n[6] POST /convert/xlsx_to_word: базовый JSON-ответ")
    with open(xlsx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "результат.xlsx"),
                "year": "2020"}
        r = client.post("/convert/xlsx_to_word", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("filename есть", bool(j.get("filename")), True)
    check("filename с '3СУ' и '2020'",
          "3СУ" in j["filename"] and "2020" in j["filename"], True)
    check("count = 6", j.get("count"), 6)
    check("preview = 6", len(j.get("preview") or []), 6)

    print("\n[7] Файл создан в output/")
    out_file = os.path.join(out_dir, j["filename"])
    check("Файл существует", os.path.exists(out_file), True)
    from core.converter.markers import read_word_marker
    check("Метка в акте", read_word_marker(out_file)["found"], True)

    print("\n[8] Повторный POST без overwrite -> 409")
    with open(xlsx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "результат.xlsx"),
                "year": "2020"}
        r = client.post("/convert/xlsx_to_word", data=data,
                        content_type="multipart/form-data")
    check("HTTP 409", r.status_code, 409)
    j = r.get_json()
    check("conflict=true", j.get("conflict"), True)
    check("filename в ответе", bool(j.get("filename")), True)

    print("\n[9] overwrite=1 -> 200 и перезапись")
    with open(xlsx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "результат.xlsx"),
                "year": "2020", "overwrite": "1"}
        r = client.post("/convert/xlsx_to_word", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    check("filename без скобок", "(" in r.get_json()["filename"], False)

    print("\n[10] copy=1 -> 200 и файл с (1)")
    with open(xlsx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "результат.xlsx"),
                "year": "2020", "copy": "1"}
        r = client.post("/convert/xlsx_to_word", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("filename содержит '(1)'", "(1)" in j.get("filename", ""), True)
    check("Файл (1) создан", os.path.exists(os.path.join(out_dir, j["filename"])), True)

    print("\n[11] download=1 -> send_file")
    with open(xlsx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "результат.xlsx"),
                "year": "2020", "copy": "1", "download": "1"}
        r = client.post("/convert/xlsx_to_word", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    cd = r.headers.get("Content-Disposition", "")
    check("Content-Disposition: attachment", "attachment" in cd, True)
    check("Есть магические байты docx (PK)",
          r.data[:2] == b"PK", True)

    print("\n[12] Валидация года")
    with open(xlsx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "результат.xlsx"),
                "year": "не год"}
        r = client.post("/convert/xlsx_to_word", data=data,
                        content_type="multipart/form-data")
    check("HTTP 400", r.status_code, 400)

    print("\n[13] Валидация расширения в /convert")
    with open(xlsx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "результат.doc"),
                "year": "2020"}
        r = client.post("/convert/xlsx_to_word", data=data,
                        content_type="multipart/form-data")
    check("HTTP 400", r.status_code, 400)

    print("\n[14] court_area_id: создаём участок, конвертируем с ним")
    r = client.post("/api/court_areas", json={
        "номер": "5", "судья": "Пятый П. П.",
        "дата_утверждения": "«___» ______________ 20__ года",
        "дата_акта": "«__» ____________ 20__ г.", "номер_акта": "1",
        "секретарь": "Сек5 С. С.",
        "дата_подписи": "«__» ____________ 20__ г.",
        "протокол_эк_дата": "«__» ____________ 20__ г.",
        "протокол_эк_номер": "1",
    })
    check("Участок создан", r.status_code, 200)
    area_id = r.get_json()["id"]
    with open(xlsx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "результат.xlsx"),
                "year": "2020", "court_area_id": str(area_id),
                "copy": "1"}
        r = client.post("/convert/xlsx_to_word", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("filename с '5СУ'", "5СУ" in j.get("filename", ""), True)
    out_file = os.path.join(out_dir, j["filename"])
    doc = Document(out_file)
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                text += "\n" + c.text
    check("В акте судья участка 5", "Пятый П. П." in text, True)
    check("В акте секретарь участка 5", "Сек5 С. С." in text, True)

    print("\n[15] /process file_type=legacy: перенос данных «как есть»")
    with open(xlsx_path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), "результат.xlsx"),
                "year": "2020", "format": "excel", "file_type": "legacy"}
        r = client.post("/process", data=data,
                        content_type="multipart/form-data")
    check("HTTP 200", r.status_code, 200)
    j = r.get_json()
    check("Предупреждения о неполноте нет",
          "не восстановлены" not in (j.get("stats_message") or ""), True)
    from openpyxl import load_workbook
    wb_res = load_workbook(os.path.join(out_dir, j["filename"]),
                           data_only=True)
    ws_res = wb_res["Результат обработки"]
    headers = [c.value for c in next(ws_res.iter_rows(min_row=1, max_row=1))]
    first = [c.value for c in next(ws_res.iter_rows(min_row=2, max_row=2))]
    col = {h: i for i, h in enumerate(headers)}
    check("Даты дела перенесены",
          first[col["Даты дела"]], "01.01.2020\n01.02.2020")
    check("№ описи перенесён", str(first[col["№ описи"]]), "1")
    check("№ ед.хр. перенесён", str(first[col["№ ед.хр."]]), "1")
    check("Срок хранения перенесён",
          first[col["Срок хранения"]], "3 года ЭК Ст. 137")
    check("Имена не пересклонены (ответчик)",
          "Тестову Т. Т." in str(first[col["Заголовок дела"]]), True)

    print(f"\nИТОГО: пройдено {PASSED}, ошибок {FAILED}")
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()
